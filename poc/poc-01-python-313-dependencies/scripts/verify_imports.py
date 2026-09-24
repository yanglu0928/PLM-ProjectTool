from __future__ import annotations

import importlib
import io
import json
import platform
import sys
import tempfile
from pathlib import Path
from typing import Callable


RESULTS: list[dict[str, str]] = []


def check(name: str, operation: Callable[[], None]) -> None:
    try:
        operation()
    except Exception as exc:  # PoC harness must record third-party failures.
        RESULTS.append({"check": name, "status": "FAIL", "detail": f"{type(exc).__name__}: {exc}"})
    else:
        RESULTS.append({"check": name, "status": "PASS", "detail": ""})


def import_module(name: str) -> None:
    importlib.import_module(name)


def verify_fastapi() -> None:
    from fastapi import FastAPI

    app = FastAPI()
    assert app is not None


def verify_sqlalchemy() -> None:
    from sqlalchemy import create_engine, text

    engine = create_engine("sqlite+pysqlite:///:memory:")
    with engine.connect() as connection:
        assert connection.execute(text("select 1")).scalar_one() == 1


def verify_pdf_stack() -> None:
    import pymupdf as fitz
    import pdfplumber

    document = fitz.open()
    page = document.new_page()
    page.insert_text((72, 72), "PLM POC-01")
    pdf_bytes = document.tobytes()
    document.close()
    with pdfplumber.open(io.BytesIO(pdf_bytes)) as parsed:
        assert "PLM POC-01" in (parsed.pages[0].extract_text() or "")


def verify_docx() -> None:
    from docx import Document

    stream = io.BytesIO()
    document = Document()
    document.add_paragraph("PLM POC-01")
    document.save(stream)
    stream.seek(0)
    assert Document(stream).paragraphs[0].text == "PLM POC-01"


def verify_pptx() -> None:
    from pptx import Presentation

    stream = io.BytesIO()
    presentation = Presentation()
    presentation.slides.add_slide(presentation.slide_layouts[6])
    presentation.save(stream)
    stream.seek(0)
    assert len(Presentation(stream).slides) == 1


def verify_openpyxl() -> None:
    from openpyxl import Workbook, load_workbook

    with tempfile.TemporaryDirectory() as temp_dir:
        path = Path(temp_dir) / "poc-01.xlsx"
        workbook = Workbook()
        workbook.active["A1"] = "PLM POC-01"
        workbook.save(path)
        assert load_workbook(path).active["A1"].value == "PLM POC-01"


def main() -> int:
    check("python_3_13", lambda: (_ for _ in ()).throw(RuntimeError(platform.python_version())) if sys.version_info[:2] != (3, 13) else None)
    check("fastapi_minimal", verify_fastapi)
    check("uvicorn_import", lambda: import_module("uvicorn"))
    check("sqlalchemy_minimal", verify_sqlalchemy)
    check("alembic_import", lambda: import_module("alembic"))
    check("psycopg_import", lambda: import_module("psycopg"))
    check("pgvector_import", lambda: import_module("pgvector"))
    check("pdf_stack_minimal", verify_pdf_stack)
    check("docx_minimal", verify_docx)
    check("pptx_minimal", verify_pptx)
    check("openpyxl_minimal", verify_openpyxl)
    check("paddleocr_import", lambda: import_module("paddleocr"))
    check("paddle_import", lambda: import_module("paddle"))
    check("pytesseract_import", lambda: import_module("pytesseract"))
    check("ocrmypdf_import", lambda: import_module("ocrmypdf"))

    report = {
        "python_version": platform.python_version(),
        "status": "PASS" if all(item["status"] == "PASS" for item in RESULTS) else "FAIL",
        "checks": RESULTS,
    }
    print(json.dumps(report, ensure_ascii=False, indent=2))
    return 0 if report["status"] == "PASS" else 1


if __name__ == "__main__":
    raise SystemExit(main())
