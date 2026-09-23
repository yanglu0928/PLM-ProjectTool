from __future__ import annotations

import argparse
import hashlib
import json
import zipfile
from pathlib import Path

from docx import Document
from pptx import Presentation


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def validate_docx(path: Path) -> dict[str, object]:
    with zipfile.ZipFile(path) as archive:
        broken = archive.testzip()
        media_count = len([name for name in archive.namelist() if name.startswith("word/media/")])
    document = Document(path)
    text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    headings = [paragraph.style.name for paragraph in document.paragraphs if paragraph.style.name.startswith("Heading")]
    explicit_breaks = sum(1 for paragraph in document.paragraphs if 'w:type="page"' in paragraph._p.xml)
    checks = {
        "zip_integrity": broken is None,
        "explicit_page_breaks": explicit_breaks == 99,
        "heading_level_1": "Heading 1" in headings,
        "heading_level_2": "Heading 2" in headings,
        "heading_level_3": "Heading 3" in headings,
        "chinese_text": "项目实施" in text,
        "table_count": len(document.tables) >= 9,
        "image_count": len(document.inline_shapes) >= 4,
        "media_count": media_count >= 1,
    }
    return {
        "status": "PASS" if all(checks.values()) else "FAIL",
        "size_bytes": path.stat().st_size,
        "sha256": sha256(path),
        "paragraph_count": len(document.paragraphs),
        "table_count": len(document.tables),
        "inline_shape_count": len(document.inline_shapes),
        "checks": checks,
    }


def validate_pptx(path: Path) -> dict[str, object]:
    with zipfile.ZipFile(path) as archive:
        broken = archive.testzip()
    presentation = Presentation(path)
    texts = []
    table_count = 0
    image_count = 0
    connector_count = 0
    for slide in presentation.slides:
        for shape in slide.shapes:
            if getattr(shape, "has_text_frame", False):
                texts.append(shape.text)
            if getattr(shape, "has_table", False):
                table_count += 1
            if shape.shape_type == 13:
                image_count += 1
            if shape.shape_type == 9:
                connector_count += 1
    checks = {
        "zip_integrity": broken is None,
        "slide_count": len(presentation.slides) == 50,
        "chinese_text": "项目实施" in "\n".join(texts),
        "native_table": table_count >= 1,
        "editable_data_visualization": "所有数据条、标签和数值" in "\n".join(texts),
        "image_count": image_count >= 1,
        "connector_count": connector_count >= 3,
    }
    return {
        "status": "PASS" if all(checks.values()) else "FAIL",
        "size_bytes": path.stat().st_size,
        "sha256": sha256(path),
        "slide_count": len(presentation.slides),
        "table_count": table_count,
        "image_count": image_count,
        "connector_count": connector_count,
        "checks": checks,
    }


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--docx", required=True, type=Path)
    parser.add_argument("--pptx", required=True, type=Path)
    parser.add_argument("--output", required=True, type=Path)
    args = parser.parse_args()
    result = {"docx": validate_docx(args.docx), "pptx": validate_pptx(args.pptx)}
    result["status"] = "PASS" if result["docx"]["status"] == result["pptx"]["status"] == "PASS" else "FAIL"
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(result, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps({"status": result["status"], "docx": result["docx"]["checks"], "pptx": result["pptx"]["checks"]}, ensure_ascii=False))
    return 0 if result["status"] == "PASS" else 1


if __name__ == "__main__":
    raise SystemExit(main())
