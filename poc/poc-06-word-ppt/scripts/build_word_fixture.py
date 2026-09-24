from __future__ import annotations

import argparse
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


PAGE_COUNT = 100
FONT_NAME = "Microsoft YaHei"


def set_run_font(run, size: float, *, bold: bool = False, color: str = "000000") -> None:
    run.font.name = FONT_NAME
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), FONT_NAME)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def shade_cell(cell, fill: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shading = properties.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        properties.append(shading)
    shading.set(qn("w:fill"), fill)


def set_cell_margins(cell, top: int = 90, start: int = 120, bottom: int = 90, end: int = 120) -> None:
    properties = cell._tc.get_or_add_tcPr()
    margins = properties.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        properties.append(margins)
    for tag, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = margins.find(qn(f"w:{tag}"))
        if node is None:
            node = OxmlElement(f"w:{tag}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run("第 ")
    set_run_font(run, 9, color="666666")
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    paragraph._p.append(field)
    run = paragraph.add_run(" 页")
    set_run_font(run, 9, color="666666")


def add_table(document: Document, page_number: int) -> None:
    table = document.add_table(rows=4, cols=4)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    values = [
        ["对象", "版本", "状态", "责任域"],
        [f"产品结构 {page_number:03d}", "A.1", "受控", "设计"],
        [f"工艺路线 {page_number:03d}", "B.2", "评审", "工艺"],
        [f"变更记录 {page_number:03d}", "C.3", "归档", "项目"],
    ]
    for row_index, row in enumerate(table.rows):
        for column_index, cell in enumerate(row.cells):
            cell.text = values[row_index][column_index]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            if row_index == 0:
                shade_cell(cell, "17365D")
            elif row_index % 2 == 0:
                shade_cell(cell, "EAF2F8")
            for paragraph in cell.paragraphs:
                paragraph.alignment = (
                    WD_ALIGN_PARAGRAPH.LEFT if column_index == 0 else WD_ALIGN_PARAGRAPH.CENTER
                )
                for run in paragraph.runs:
                    set_run_font(run, 9.5, bold=row_index == 0, color="FFFFFF" if row_index == 0 else "000000")


def add_flow(document: Document) -> None:
    table = document.add_table(rows=1, cols=7)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    labels = ["需求确认", "→", "方案设计", "→", "实施验证", "→", "验收归档"]
    for index, cell in enumerate(table.rows[0].cells):
        cell.text = labels[index]
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_margins(cell, top=140, bottom=140)
        if index % 2 == 0:
            shade_cell(cell, "DCE6F1")
        for paragraph in cell.paragraphs:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in paragraph.runs:
                set_run_font(run, 10, bold=index % 2 == 0, color="17365D")


def add_body_paragraph(document: Document, text: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(7)
    paragraph.paragraph_format.line_spacing = 1.25
    run = paragraph.add_run(text)
    set_run_font(run, 11)


def configure_styles(document: Document) -> None:
    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = FONT_NAME
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_NAME)
    normal.font.size = Pt(11)
    for name, size in (("Title", 28), ("Heading 1", 18), ("Heading 2", 14), ("Heading 3", 12)):
        style = styles[name]
        style.font.name = FONT_NAME
        style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_NAME)
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.font.bold = True


def build(output_path: Path, image_path: Path) -> None:
    document = Document()
    section = document.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)
    configure_styles(document)
    document.core_properties.title = "PLM 项目实施辅助工具 Word 兼容性测试"
    document.core_properties.subject = "POC-06 synthetic fixture"
    document.core_properties.author = "PLM Project Tool"
    add_page_number(section.footer.paragraphs[0])

    title = document.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_before = Pt(120)
    set_run_font(title.add_run("PLM 项目实施辅助工具"), 28, bold=True)
    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(subtitle.add_run("100 页 Word 兼容性与版式验证样例"), 16, color="17365D")
    note = document.add_paragraph()
    note.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(note.add_run("纯合成内容  版本 POC-06.1"), 10, color="666666")
    document.add_picture(str(image_path), width=Inches(5.8))
    document.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER

    for page_number in range(2, PAGE_COUNT + 1):
        document.add_page_break()
        chapter = ((page_number - 2) // 10) + 1
        topic_index = ((page_number - 2) % 10) + 1
        document.add_heading(f"第 {chapter} 章  项目实施验证域", level=1)
        document.add_heading(f"{chapter}.{topic_index}  合成业务场景 {page_number:03d}", level=2)
        if page_number % 3 == 0:
            document.add_heading(f"{chapter}.{topic_index}.1  处理规则", level=3)
        add_body_paragraph(
            document,
            f"本页用于验证第 {page_number:03d} 个中文内容页面。样例描述产品结构、文档版本、变更记录和交付证据之间的关系，不包含真实客户、合同或项目数据。",
        )
        add_body_paragraph(
            document,
            "系统按照项目范围保存来源定位和版本信息。所有条目均为确定性测试内容，可用于检查中文字体、分页、段落间距和多级标题。",
        )
        if page_number in {5, 25, 45, 65, 85}:
            add_table(document, page_number)
        elif page_number in {20, 40, 60, 80}:
            add_flow(document)
        elif page_number in {10, 50, 90}:
            document.add_picture(str(image_path), width=Inches(5.4))
            document.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
            caption = document.add_paragraph()
            caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
            set_run_font(caption.add_run("图：合成工业研发协同场景"), 9.5, color="666666")
        else:
            bullets = [
                "保留章节层级和项目内来源定位。",
                "表格与流程对象保持可编辑。",
                "输出文件不得包含外部链接或客户标识。",
            ]
            for item in bullets:
                paragraph = document.add_paragraph(style="List Bullet")
                paragraph.paragraph_format.space_after = Pt(3)
                set_run_font(paragraph.add_run(item), 10.5)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(output_path)


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--output", required=True, type=Path)
    parser.add_argument("--image", required=True, type=Path)
    args = parser.parse_args()
    build(args.output, args.image)
    print(f"output={args.output}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
