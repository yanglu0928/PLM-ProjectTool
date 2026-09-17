from __future__ import annotations

import csv
import hashlib
import mimetypes
import os
import tempfile
from pathlib import Path
from typing import Any, Iterable

import pymupdf
from docx import Document
from docx.oxml.table import CT_Tbl
from docx.oxml.text.paragraph import CT_P
from docx.table import Table
from docx.text.paragraph import Paragraph
from openpyxl import load_workbook
from PIL import Image
from pptx import Presentation

from .models import Block, Page, ParsedDocument, Source


SUPPORTED_SUFFIXES = {".docx", ".pptx", ".xlsx", ".csv", ".pdf"}


class BlockBuilder:
    def __init__(self) -> None:
        self.blocks: list[Block] = []

    def add(
        self,
        block_type: str,
        text: str,
        *,
        page: int | None,
        section: str | None,
        source_locator: str,
        table: dict[str, Any] | None = None,
        metadata: dict[str, Any] | None = None,
    ) -> None:
        clean = " ".join(str(text).split())
        if not clean:
            return
        self.blocks.append(
            Block(
                id=f"b{len(self.blocks) + 1:04d}",
                type=block_type,
                text=clean,
                page=page,
                section=section,
                table=table,
                source_locator=source_locator,
                metadata=metadata or {},
            )
        )


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _source(path: Path) -> Source:
    media_type = mimetypes.guess_type(path.name)[0] or "application/octet-stream"
    return Source(path.name, _sha256(path), media_type, path.stat().st_size)


def _docx(path: Path) -> ParsedDocument:
    document = Document(path)
    builder = BlockBuilder()
    pages: list[Page] = [Page(1, "Page 1", "word/page/1")]
    page_number = 1
    section: str | None = None
    title: str | None = None
    paragraph_index = 0
    table_index = 0

    for child in document.element.body.iterchildren():
        if isinstance(child, CT_P):
            paragraph_index += 1
            paragraph = Paragraph(child, document)
            style = paragraph.style.name if paragraph.style else ""
            has_page_break = bool(child.xpath(".//w:br[@w:type='page']"))
            text = paragraph.text.strip()
            if text:
                block_type = "paragraph"
                if style == "Title":
                    block_type = "title"
                    title = title or text
                elif style.startswith("Heading"):
                    block_type = "heading"
                    section = text
                builder.add(
                    block_type,
                    text,
                    page=page_number,
                    section=section,
                    source_locator=f"word/paragraph/{paragraph_index}",
                    metadata={"style": style},
                )
            if has_page_break:
                page_number += 1
                pages.append(Page(page_number, f"Page {page_number}", f"word/page/{page_number}"))
        elif isinstance(child, CT_Tbl):
            table_index += 1
            table = Table(child, document)
            for row_index, row in enumerate(table.rows, start=1):
                cells = [" ".join(cell.text.split()) for cell in row.cells]
                builder.add(
                    "table_row",
                    " | ".join(cells),
                    page=page_number,
                    section=section,
                    source_locator=f"word/table/{table_index}/row/{row_index}",
                    table={"index": table_index, "row": row_index, "cells": cells},
                )

    return ParsedDocument(
        source=_source(path),
        title=title,
        pages=pages,
        blocks=builder.blocks,
        metadata={"format": "docx", "explicit_page_count": len(pages)},
        warnings=["DOCX page numbers reflect explicit page breaks, not automatic layout pagination."],
    )


def _pptx(path: Path) -> ParsedDocument:
    presentation = Presentation(path)
    builder = BlockBuilder()
    pages: list[Page] = []
    title: str | None = None

    for slide_number, slide in enumerate(presentation.slides, start=1):
        pages.append(Page(slide_number, f"Slide {slide_number}", f"ppt/slides/{slide_number}"))
        slide_title = None
        if slide.shapes.title is not None:
            slide_title = slide.shapes.title.text.strip() or None
        if slide_title is None:
            for candidate in slide.shapes:
                if (
                    getattr(candidate, "has_text_frame", False)
                    and candidate.name.lower().startswith("title")
                ):
                    slide_title = candidate.text.strip() or None
                    break
        title = title or slide_title
        section = slide_title
        for shape_index, shape in enumerate(slide.shapes, start=1):
            locator = f"ppt/slides/{slide_number}/shapes/{shape_index}"
            if getattr(shape, "has_table", False):
                for row_index, row in enumerate(shape.table.rows, start=1):
                    cells = [" ".join(cell.text.split()) for cell in row.cells]
                    builder.add(
                        "table_row",
                        " | ".join(cells),
                        page=slide_number,
                        section=section,
                        source_locator=f"{locator}/table/1/row/{row_index}",
                        table={"index": 1, "row": row_index, "cells": cells},
                    )
            elif getattr(shape, "has_text_frame", False):
                text = shape.text.strip()
                if not text:
                    continue
                is_title = shape is slide.shapes.title or shape.name.lower().startswith("title")
                builder.add(
                    "heading" if is_title else "paragraph",
                    text,
                    page=slide_number,
                    section=section,
                    source_locator=locator,
                    metadata={"shape_name": shape.name},
                )

    return ParsedDocument(
        source=_source(path),
        title=title,
        pages=pages,
        blocks=builder.blocks,
        metadata={"format": "pptx", "slide_count": len(pages)},
    )


def _xlsx(path: Path) -> ParsedDocument:
    workbook = load_workbook(path, data_only=False, read_only=True)
    builder = BlockBuilder()
    pages: list[Page] = []

    for sheet_index, sheet in enumerate(workbook.worksheets, start=1):
        pages.append(Page(None, sheet.title, f"excel/sheets/{sheet.title}"))
        for row_index, row in enumerate(sheet.iter_rows(values_only=True), start=1):
            values = ["" if value is None else str(value) for value in row]
            if not any(value.strip() for value in values):
                continue
            last_col = max(index for index, value in enumerate(values, start=1) if value.strip())
            cells = values[:last_col]
            builder.add(
                "table_row",
                " | ".join(cells),
                page=None,
                section=sheet.title,
                source_locator=f"excel/sheets/{sheet.title}/rows/{row_index}",
                table={
                    "sheet": sheet.title,
                    "row": row_index,
                    "range": f"A{row_index}:{sheet.cell(row=row_index, column=last_col).coordinate}",
                    "cells": cells,
                },
            )

    return ParsedDocument(
        source=_source(path),
        title=workbook.properties.title or path.stem,
        pages=pages,
        blocks=builder.blocks,
        metadata={"format": "xlsx", "sheet_count": len(pages)},
        warnings=["XLSX has no stable physical page number; sheet and cell range are retained."],
    )


def _csv(path: Path) -> ParsedDocument:
    builder = BlockBuilder()
    with path.open("r", encoding="utf-8-sig", newline="") as stream:
        rows = list(csv.reader(stream))
    for row_index, cells in enumerate(rows, start=1):
        if not any(cell.strip() for cell in cells):
            continue
        builder.add(
            "table_row",
            " | ".join(cells),
            page=None,
            section=path.stem,
            source_locator=f"csv/rows/{row_index}",
            table={"row": row_index, "range": f"row:{row_index}", "cells": cells},
        )
    return ParsedDocument(
        source=_source(path),
        title=path.stem,
        pages=[Page(None, path.stem, "csv")],
        blocks=builder.blocks,
        metadata={"format": "csv", "row_count": len(rows)},
        warnings=["CSV has no physical page number; row number is retained."],
    )


def _pdf_text(path: Path) -> ParsedDocument:
    builder = BlockBuilder()
    pages: list[Page] = []
    title: str | None = None
    current_section: str | None = None
    page_sections: dict[int, str | None] = {}
    with pymupdf.open(path) as document:
        for page_index, page in enumerate(document, start=1):
            pages.append(Page(page_index, f"Page {page_index}", f"pdf/pages/{page_index}"))
            raw = page.get_text("dict")
            for block_index, raw_block in enumerate(raw.get("blocks", []), start=1):
                if "lines" not in raw_block:
                    continue
                spans = [span for line in raw_block["lines"] for span in line.get("spans", [])]
                text = " ".join(span.get("text", "") for span in spans).strip()
                if not text:
                    continue
                max_size = max((float(span.get("size", 0)) for span in spans), default=0)
                block_type = "heading" if max_size >= 15 else "paragraph"
                if block_type == "heading":
                    current_section = " ".join(text.split())
                    title = title or current_section
                builder.add(
                    block_type,
                    text,
                    page=page_index,
                    section=current_section,
                    source_locator=f"pdf/pages/{page_index}/text-blocks/{block_index}",
                    metadata={"bbox": raw_block.get("bbox"), "max_font_size": max_size},
                )
            page_sections[page_index] = current_section

    try:
        import pdfplumber

        with pdfplumber.open(path) as document:
            for page_index, page in enumerate(document.pages, start=1):
                for table_index, table in enumerate(page.extract_tables(), start=1):
                    for row_index, row in enumerate(table, start=1):
                        cells = ["" if cell is None else " ".join(cell.split()) for cell in row]
                        builder.add(
                            "table_row",
                            " | ".join(cells),
                            page=page_index,
                            section=page_sections.get(page_index),
                            source_locator=(
                                f"pdf/pages/{page_index}/tables/{table_index}/rows/{row_index}"
                            ),
                            table={"index": table_index, "row": row_index, "cells": cells},
                        )
    except Exception as exc:  # pragma: no cover - evidence captures parser-specific failure
        warning = f"pdfplumber table extraction failed: {type(exc).__name__}: {exc}"
    else:
        warning = None

    warnings = [warning] if warning else []
    return ParsedDocument(
        source=_source(path),
        title=title or path.stem,
        pages=pages,
        blocks=builder.blocks,
        metadata={"format": "pdf", "ocr_engine": None},
        warnings=warnings,
    )


def _extract_paddle_records(result_items: Iterable[Any]) -> list[tuple[str, float | None, Any]]:
    records: list[tuple[str, float | None, Any]] = []
    for item in result_items:
        payload = getattr(item, "json", item)
        if callable(payload):
            payload = payload()
        if isinstance(payload, str):
            import json

            payload = json.loads(payload)
        if isinstance(payload, dict) and isinstance(payload.get("res"), dict):
            payload = payload["res"]
        if not isinstance(payload, dict):
            continue
        texts = payload.get("rec_texts") or payload.get("texts") or []
        scores = payload.get("rec_scores") or payload.get("scores") or []
        boxes = payload.get("rec_polys") or payload.get("dt_polys") or []
        for index, text in enumerate(texts):
            score = float(scores[index]) if index < len(scores) else None
            box = boxes[index].tolist() if index < len(boxes) and hasattr(boxes[index], "tolist") else (
                boxes[index] if index < len(boxes) else None
            )
            records.append((str(text), score, box))
    return records


def _ocr_pdf(path: Path, engine: str) -> ParsedDocument:
    builder = BlockBuilder()
    pages: list[Page] = []
    warnings: list[str] = []
    title: str | None = None

    paddle = None
    if engine == "paddle":
        from paddleocr import PaddleOCR

        detection_model_dir = os.environ.get("PADDLE_DET_MODEL_DIR")
        recognition_model_dir = os.environ.get("PADDLE_REC_MODEL_DIR")
        for configured_dir in (detection_model_dir, recognition_model_dir):
            if configured_dir and not Path(configured_dir).is_dir():
                raise FileNotFoundError(f"Configured PaddleOCR model directory not found: {configured_dir}")
        paddle = PaddleOCR(
            text_detection_model_name="PP-OCRv5_mobile_det",
            text_detection_model_dir=detection_model_dir,
            text_recognition_model_name="PP-OCRv5_mobile_rec",
            text_recognition_model_dir=recognition_model_dir,
            use_doc_orientation_classify=False,
            use_doc_unwarping=False,
            use_textline_orientation=False,
            device="cpu",
            enable_mkldnn=False,
        )
    elif engine != "tesseract":
        raise ValueError(f"Unsupported OCR engine: {engine}")

    with tempfile.TemporaryDirectory(prefix="poc05-ocr-") as temp_dir, pymupdf.open(path) as document:
        for page_index, page in enumerate(document, start=1):
            pages.append(Page(page_index, f"Page {page_index}", f"pdf/pages/{page_index}"))
            pixmap = page.get_pixmap(matrix=pymupdf.Matrix(2.5, 2.5), alpha=False)
            image_path = Path(temp_dir) / f"page-{page_index}.png"
            pixmap.save(image_path)
            if engine == "paddle":
                records = _extract_paddle_records(paddle.predict(str(image_path)))
            else:
                import pytesseract
                from pytesseract import Output

                configured = os.environ.get("TESSERACT_EXE")
                if configured:
                    pytesseract.pytesseract.tesseract_cmd = configured
                data = pytesseract.image_to_data(
                    Image.open(image_path), lang="chi_sim+eng", output_type=Output.DICT
                )
                records = []
                for index, text in enumerate(data["text"]):
                    clean = text.strip()
                    if not clean:
                        continue
                    try:
                        score = float(data["conf"][index]) / 100.0
                    except (TypeError, ValueError):
                        score = None
                    records.append(
                        (
                            clean,
                            score,
                            [
                                data["left"][index],
                                data["top"][index],
                                data["width"][index],
                                data["height"][index],
                            ],
                        )
                    )

            for line_index, (text, score, box) in enumerate(records, start=1):
                title = title or text
                builder.add(
                    "ocr_line",
                    text,
                    page=page_index,
                    section=title,
                    source_locator=f"pdf/pages/{page_index}/ocr/{engine}/lines/{line_index}",
                    metadata={"engine": engine, "confidence": score, "bbox": box},
                )
            if not records:
                warnings.append(f"No OCR records returned for page {page_index} by {engine}.")

    return ParsedDocument(
        source=_source(path),
        title=title or path.stem,
        pages=pages,
        blocks=builder.blocks,
        metadata={
            "format": "pdf",
            "ocr_engine": engine,
            "ocr_model": (
                {
                    "detection": "PP-OCRv5_mobile_det",
                    "recognition": "PP-OCRv5_mobile_rec",
                    "source": "explicit_local_directories"
                    if detection_model_dir and recognition_model_dir
                    else "paddlex_managed_cache",
                }
                if engine == "paddle"
                else {"languages": ["chi_sim", "eng"]}
            ),
        },
        warnings=warnings,
    )


def parse_document(path: str | Path, *, ocr_engine: str | None = None) -> ParsedDocument:
    source_path = Path(path).resolve()
    if not source_path.is_file():
        raise FileNotFoundError(source_path)
    suffix = source_path.suffix.lower()
    if suffix not in SUPPORTED_SUFFIXES:
        raise ValueError(f"Unsupported document type: {suffix}")
    if suffix == ".docx":
        return _docx(source_path)
    if suffix == ".pptx":
        return _pptx(source_path)
    if suffix == ".xlsx":
        return _xlsx(source_path)
    if suffix == ".csv":
        return _csv(source_path)
    if ocr_engine:
        return _ocr_pdf(source_path, ocr_engine)
    return _pdf_text(source_path)
