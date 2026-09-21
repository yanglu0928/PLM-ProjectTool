from __future__ import annotations

import argparse
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


FONT_NAME = "Microsoft YaHei"


def set_cell_shading(cell, fill: str) -> None:
    props = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    props.append(shading)


def set_cell_margins(cell, top: int = 100, start: int = 120, bottom: int = 100, end: int = 120) -> None:
    props = cell._tc.get_or_add_tcPr()
    margins = props.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        props.append(margins)
    for tag, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = margins.find(qn(f"w:{tag}"))
        if node is None:
            node = OxmlElement(f"w:{tag}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def style_run(run, size: float, *, bold: bool = False, color: str = "000000") -> None:
    run.font.name = FONT_NAME
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), FONT_NAME)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def add_table(document: Document, rows: list[list[str]]) -> None:
    table = document.add_table(rows=len(rows), cols=len(rows[0]))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    widths = [Inches(1.8), Inches(4.6)]
    for row_index, values in enumerate(rows):
        for column_index, value in enumerate(values):
            cell = table.cell(row_index, column_index)
            cell.width = widths[column_index]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            if row_index == 0:
                set_cell_shading(cell, "1F4E78")
            elif row_index % 2 == 0:
                set_cell_shading(cell, "EAF2F8")
            paragraph = cell.paragraphs[0]
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if column_index == 0 else WD_ALIGN_PARAGRAPH.LEFT
            run = paragraph.add_run(value)
            style_run(run, 10.5, bold=row_index == 0, color="FFFFFF" if row_index == 0 else "000000")
    document.add_paragraph().paragraph_format.space_after = Pt(0)


def build(output: Path) -> None:
    document = Document()
    section = document.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.85)

    styles = document.styles
    for style_name in ("Normal", "Title", "Heading 1", "Heading 2"):
        style = styles[style_name]
        style.font.name = FONT_NAME
        style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_NAME)
        style.font.color.rgb = RGBColor(0, 0, 0)
    styles["Normal"].font.size = Pt(11)
    styles["Title"].font.size = Pt(24)
    styles["Heading 1"].font.size = Pt(17)
    styles["Heading 2"].font.size = Pt(14)

    title = document.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    style_run(title.add_run("POC 05 统一解析测试"), 24, bold=True)

    intro = document.add_paragraph()
    intro.paragraph_format.space_after = Pt(14)
    style_run(
        intro.add_run("本文件验证 DOCX 标题、章节、显式分页和表格能否保留来源位置。"),
        11,
    )

    heading = document.add_paragraph(style="Heading 1")
    style_run(heading.add_run("项目概况"), 17, bold=True)
    for text in (
        "项目名称：PLM 项目实施辅助工具",
        "项目编号：PLM-2026-005",
        "当前里程碑：需求确认",
    ):
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.space_after = Pt(7)
        style_run(paragraph.add_run(text), 11)

    add_table(
        document,
        [
            ["字段", "值"],
            ["项目编号", "PLM-2026-005"],
            ["里程碑", "需求确认"],
            ["确认方", "客户确认"],
        ],
    )

    document.add_page_break()
    heading = document.add_paragraph(style="Heading 1")
    style_run(heading.add_run("交付清单"), 17, bold=True)
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(12)
    style_run(paragraph.add_run("第二页用于验证显式页码和章节切换。"), 11)
    add_table(
        document,
        [
            ["交付物", "状态"],
            ["解析结果", "待验证"],
            ["来源定位", "待验证"],
        ],
    )

    output.parent.mkdir(parents=True, exist_ok=True)
    document.core_properties.title = "POC 05 统一解析测试"
    document.core_properties.subject = "Document and OCR proof of concept fixture"
    document.save(output)


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()
    build(args.output)
    print(args.output)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
