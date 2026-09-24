from __future__ import annotations

import copy
from dataclasses import dataclass
from pathlib import Path
from typing import Any

from jsonschema import Draft202012Validator, FormatChecker
from openpyxl import load_workbook

from .review_import import ImportIssue, _reviewed_at, _text


CONFIRM_DECISION = "确认本次修正"
RETURN_DECISION = "退回修正"
HEADERS = [
    "CaseId",
    "R7当前问题",
    "R7当前分类",
    "建议改写问题",
    "建议分类",
    "直接依据",
    "建议引用Chunk",
    "打开证据",
    "人工备注",
    "确认状态",
]


@dataclass
class R71ImportResult:
    dataset: dict[str, Any] | None
    issues: list[ImportIssue]
    confirmed: bool
    query_change_count: int = 0
    classification_change_count: int = 0
    citation_change_count: int = 0

    @property
    def ready(self) -> bool:
        return self.confirmed and self.dataset is not None and not self.issues


def _issue(code: str, message: str, *, row: int | None = None, field: str | None = None) -> ImportIssue:
    return ImportIssue(code, message, row=row, field=field)


def import_r7_exception_review(
    workbook_path: Path,
    r7_dataset: dict[str, Any],
    package: dict[str, Any],
    *,
    dataset_id: str,
    schema: dict[str, Any],
    timezone_name: str = "Asia/Shanghai",
) -> R71ImportResult:
    issues: list[ImportIssue] = []
    item = package.get("case") or {}
    if package.get("source_dataset_id") != r7_dataset.get("dataset_id"):
        issues.append(_issue("SOURCE_DATASET_MISMATCH", "R7.1 package does not reference the supplied R7 dataset"))

    case_id = _text(item.get("case_id"))
    source_matches = [case for case in r7_dataset.get("cases") or [] if _text(case.get("case_id")) == case_id]
    if not case_id or len(source_matches) != 1:
        issues.append(_issue("INVALID_EXCEPTION_SCOPE", "R7.1 must reference exactly one existing R7 case"))

    workbook = load_workbook(workbook_path, read_only=True, data_only=False)
    try:
        if "R7.1确认" not in workbook.sheetnames:
            issues.append(_issue("MISSING_SHEET", "R7.1 workbook is missing R7.1确认"))
            return R71ImportResult(None, issues, False)
        sheet = workbook["R7.1确认"]
        if [_text(sheet.cell(8, col).value) for col in range(1, 11)] != HEADERS:
            issues.append(_issue("HEADER_MISMATCH", "R7.1 headers differ from the locked template", row=8))

        decision = _text(sheet["B5"].value)
        if decision not in {"未确认", CONFIRM_DECISION, RETURN_DECISION}:
            issues.append(_issue("INVALID_CONFIRMATION", "R7.1 confirmation has an invalid value", row=5, field="确认决定"))
        confirmed = decision == CONFIRM_DECISION
        if decision == RETURN_DECISION:
            issues.append(_issue("REVIEW_RETURNED", "R7.1 exception was returned for correction", row=5, field="确认决定"))

        locked_values = [
            item.get("case_id"),
            item.get("current_query"),
            item.get("current_classification"),
            item.get("proposed_query"),
            item.get("proposed_classification"),
            item.get("rationale"),
            item.get("chunk_id"),
        ]
        for column, expected in enumerate(locked_values, start=1):
            if _text(sheet.cell(9, column).value) != _text(expected):
                issues.append(
                    _issue(
                        "SOURCE_FIELD_CHANGED",
                        "Immutable R7.1 review field was changed",
                        row=9,
                        field=HEADERS[column - 1],
                    )
                )

        reviewer = _text(sheet["D5"].value)
        reviewed_at = ""
        if confirmed:
            if not reviewer:
                issues.append(_issue("REVIEWER_REQUIRED", "Confirmed R7.1 workbook requires a reviewer", row=5, field="确认人"))
            try:
                reviewed_at = _reviewed_at(sheet["F5"].value, timezone_name)
            except (ValueError, TypeError, KeyError):
                issues.append(_issue("INVALID_REVIEW_TIME", "Confirmed R7.1 workbook requires a valid review date", row=5, field="确认日期"))
        if not confirmed or issues or len(source_matches) != 1:
            return R71ImportResult(None, issues, confirmed)

        source = source_matches[0]
        if source.get("expected_classification") != item.get("current_classification_code"):
            issues.append(_issue("SOURCE_STATE_CHANGED", "R7 source classification no longer matches the R7.1 package"))
            return R71ImportResult(None, issues, confirmed)

        chunk_id = _text(item.get("chunk_id"))
        source_citations = [
            copy.deepcopy(citation)
            for citation in source.get("expected_citations") or []
            if _text(citation.get("chunk_id")) == chunk_id
        ]
        if not source_citations:
            issues.append(_issue("CITATION_NOT_IN_R7", "R7.1 citation is not present in the reviewed R7 source case"))
            return R71ImportResult(None, issues, confirmed)

        dataset = copy.deepcopy(r7_dataset)
        dataset["dataset_id"] = dataset_id
        target = next(case for case in dataset["cases"] if _text(case.get("case_id")) == case_id)
        old_query = target.get("query")
        old_classification = target.get("expected_classification")
        old_chunks = list(target.get("expected_relevant_chunk_ids") or [])
        target["query"] = _text(item.get("proposed_query"))
        target["expected_classification"] = _text(item.get("proposed_classification_code"))
        target["expected_relevant_chunk_ids"] = [chunk_id]
        target["expected_answer_terms"] = list(item.get("proposed_answer_terms") or [])
        target["expected_citations"] = source_citations
        target["review"] = {
            "status": "APPROVED",
            "reviewed_by": reviewer,
            "reviewed_at": reviewed_at,
        }

        validator = Draft202012Validator(schema, format_checker=FormatChecker())
        validation_errors = sorted(validator.iter_errors(dataset), key=lambda error: list(error.absolute_path))
        if validation_errors:
            paths = ["/".join(str(part) for part in error.absolute_path) for error in validation_errors]
            issues.append(_issue("SCHEMA_VALIDATION_FAILED", f"R7.1 Golden Dataset schema validation failed at: {paths}"))
            return R71ImportResult(None, issues, confirmed)

        return R71ImportResult(
            dataset,
            issues,
            confirmed,
            int(target["query"] != old_query),
            int(target["expected_classification"] != old_classification),
            int(target["expected_relevant_chunk_ids"] != old_chunks),
        )
    finally:
        workbook.close()


def build_sanitized_r71_report(
    result: R71ImportResult,
    *,
    generated_at: str,
    output_written: bool,
) -> dict[str, Any]:
    status = "PASS" if result.ready and output_written else ("READY" if result.ready else "FAIL")
    return {
        "schema_version": "poc-03.r7-1-exception-import-result.v1",
        "generated_at": generated_at,
        "status": status,
        "privacy": {
            "workbook_content_committed": False,
            "case_id_committed": False,
            "reviewer_names_committed": False,
            "output_dataset_committed": False,
        },
        "summary": {
            "review_scope_count": 1,
            "resolved_case_count": 1 if result.ready else 0,
            "issue_count": len(result.issues),
            "query_change_count": result.query_change_count,
            "classification_change_count": result.classification_change_count,
            "citation_change_count": result.citation_change_count,
            "output_written": output_written,
        },
        "issues": [issue.sanitized() for issue in result.issues],
        "validation_limit": (
            "R7.1 is an AI-assisted calibration dataset. It cannot close P03-A12/P03-A13; "
            "an independently constructed holdout set remains mandatory."
        ),
        "conclusion": (
            "R7.1 calibration dataset export passed; coverage audit is still required."
            if status == "PASS"
            else "R7.1 dataset export remains blocked."
        ),
    }
