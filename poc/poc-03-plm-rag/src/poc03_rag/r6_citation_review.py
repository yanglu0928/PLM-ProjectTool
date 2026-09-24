from __future__ import annotations

import copy
from collections import Counter
from dataclasses import dataclass
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from .review_import import ImportIssue, _reviewed_at, _split_values, _text


GLOBAL_CONFIRMATION = "确认全部AI建议"
ROW_DECISIONS = {"", "确认AI建议", "修改后确认", "退回"}
HEADERS = [
    "复核编号", "CaseId", "资料类别", "当前分类", "当前问题", "问题原因", "AI建议修订问题",
    "AI建议引用Chunk", "打开证据", "单条决定", "人工修订问题", "人工引用Chunk", "生效状态",
]


@dataclass
class R6ImportResult:
    cases: list[dict[str, Any]]
    issues: list[ImportIssue]
    status_counts: Counter[str]
    confirmed: bool
    scope_count: int
    preserved_count: int

    @property
    def has_errors(self) -> bool:
        return bool(self.issues)

    @property
    def ready(self) -> bool:
        return self.confirmed and not self.has_errors and self.status_counts["PENDING"] == 0


def _changed(row: int, field: str) -> ImportIssue:
    return ImportIssue("SOURCE_FIELD_CHANGED", "Immutable R6 review source field was changed", row=row, field=field)


def import_r6_citation_review(
    workbook_path: Path,
    golden: dict[str, Any],
    package: dict[str, Any],
    *,
    timezone_name: str = "Asia/Shanghai",
) -> R6ImportResult:
    issues: list[ImportIssue] = []
    statuses: Counter[str] = Counter()
    source_cases = list(golden.get("cases") or [])
    source_by_id = {_text(case.get("case_id")): case for case in source_cases}
    package_cases = list(package.get("cases") or [])
    locked_ids = list(package.get("scope_case_ids") or [])
    if locked_ids != [item.get("case_id") for item in package_cases] or len(locked_ids) != 6:
        issues.append(ImportIssue("INVALID_R6_SCOPE", "R6 package scope is not the locked six-case set"))
    if any(case_id not in source_by_id for case_id in locked_ids):
        issues.append(ImportIssue("UNKNOWN_R6_CASE", "R6 package references a case outside the R5 dataset"))

    workbook = load_workbook(workbook_path, read_only=True, data_only=False)
    try:
        if "R6确认" not in workbook.sheetnames:
            issues.append(ImportIssue("MISSING_SHEET", "R6 workbook is missing R6确认"))
            return R6ImportResult([], issues, statuses, False, len(locked_ids), len(source_cases) - len(locked_ids))
        sheet = workbook["R6确认"]
        if [_text(sheet.cell(9, col).value) for col in range(1, 14)] != HEADERS:
            issues.append(ImportIssue("HEADER_MISMATCH", "R6 headers differ from the locked template", row=9))

        global_value = _text(sheet["B5"].value)
        if global_value not in {"未确认", GLOBAL_CONFIRMATION}:
            issues.append(ImportIssue("INVALID_GLOBAL_CONFIRMATION", "R6 batch confirmation has an invalid value", row=5, field="批量确认"))
        global_confirmed = global_value == GLOBAL_CONFIRMATION
        reviewer = _text(sheet["D5"].value)
        reviewed_at = ""

        by_task = {item["task_id"]: item for item in package_cases}
        seen: set[str] = set()
        updates: dict[str, tuple[str, list[str]]] = {}
        any_approved = False
        for row_number, row in enumerate(sheet.iter_rows(min_row=10, max_col=13, values_only=True), start=10):
            task_id = _text(row[0])
            if not task_id:
                continue
            item = by_task.get(task_id)
            if item is None or task_id in seen:
                issues.append(ImportIssue("UNKNOWN_OR_DUPLICATE_TASK", "R6 task is unknown or duplicated", row=row_number))
                continue
            seen.add(task_id)
            locked = [
                item["task_id"], item["case_id"], item["source_type"], item["classification"], item["original_query"],
                item["issue_reason"], item["suggested_query"], "；".join(item["suggested_chunk_ids"]),
            ]
            for column, expected in enumerate(locked, start=1):
                if _text(row[column - 1]) != _text(expected):
                    issues.append(_changed(row_number, HEADERS[column - 1]))

            decision = _text(row[9])
            if decision not in ROW_DECISIONS:
                issues.append(ImportIssue("INVALID_ROW_DECISION", "R6 row decision has an invalid value", row=row_number, field=HEADERS[9]))
                statuses["INVALID"] += 1
                continue
            if decision == "退回":
                statuses["RETURNED"] += 1
                continue
            approved = global_confirmed or decision in {"确认AI建议", "修改后确认"}
            if not approved:
                statuses["PENDING"] += 1
                continue
            any_approved = True
            if decision == "修改后确认":
                query = _text(row[10])
                chunk_ids = _split_values(row[11])
                if not query:
                    issues.append(ImportIssue("MANUAL_QUERY_REQUIRED", "Modified R6 row requires a manual query", row=row_number, field=HEADERS[10]))
                if not chunk_ids:
                    issues.append(ImportIssue("MANUAL_CITATION_REQUIRED", "Modified R6 row requires at least one chunk", row=row_number, field=HEADERS[11]))
            else:
                query = item["suggested_query"]
                chunk_ids = list(item["suggested_chunk_ids"])
            unknown = sorted(set(chunk_ids) - set(item["allowed_chunk_ids"]))
            if unknown:
                issues.append(ImportIssue("CITATION_OUTSIDE_R6_EVIDENCE", "R6 citation is outside the reviewed evidence set", row=row_number, field=HEADERS[11]))
            if query and chunk_ids and not unknown:
                updates[item["case_id"]] = (query, chunk_ids)
                statuses["APPROVED"] += 1

        for task_id in sorted(set(by_task) - seen):
            issues.append(ImportIssue("MISSING_TASK_ROW", "Locked R6 task is missing", candidate_id=task_id))

        confirmed = (global_confirmed or statuses["PENDING"] == 0) and statuses["RETURNED"] == 0 and len(updates) == len(locked_ids)
        if any_approved or global_confirmed:
            if not reviewer:
                issues.append(ImportIssue("REVIEWER_REQUIRED", "Confirmed R6 workbook requires a reviewer", row=5, field="确认人"))
            try:
                reviewed_at = _reviewed_at(sheet["F5"].value, timezone_name)
            except (ValueError, TypeError, KeyError):
                issues.append(ImportIssue("INVALID_REVIEW_TIME", "Confirmed R6 workbook requires a valid review date", row=5, field="确认日期"))

        if not confirmed or issues:
            return R6ImportResult([], issues, statuses, confirmed, len(locked_ids), len(source_cases) - len(locked_ids))

        citation_map: dict[str, dict[str, Any]] = {}
        for item in package_cases:
            for chunk in item["candidate_chunks"]:
                citation_map[chunk["chunk_id"]] = chunk
        output_cases: list[dict[str, Any]] = []
        for source in source_cases:
            case_id = source["case_id"]
            case = copy.deepcopy(source)
            if case_id in updates:
                query, chunk_ids = updates[case_id]
                case["query"] = query
                case["expected_relevant_chunk_ids"] = chunk_ids
                case["expected_citations"] = [
                    {"document_id": citation_map[chunk_id]["document_id"], "chunk_id": chunk_id, "source_locator": locator}
                    for chunk_id in chunk_ids
                    for locator in citation_map[chunk_id]["source_locators"]
                ]
                case["review"] = {"status": "APPROVED", "reviewed_by": reviewer, "reviewed_at": reviewed_at}
            output_cases.append(case)
        return R6ImportResult(output_cases, issues, statuses, True, len(locked_ids), len(source_cases) - len(locked_ids))
    finally:
        workbook.close()


def build_r6_dataset(golden: dict[str, Any], result: R6ImportResult, *, dataset_id: str) -> dict[str, Any]:
    if not result.ready:
        raise ValueError("R6 review is not ready for export")
    dataset = copy.deepcopy(golden)
    dataset["dataset_id"] = dataset_id
    dataset["cases"] = result.cases
    return dataset


def build_sanitized_r6_report(result: R6ImportResult, *, generated_at: str, output_written: bool) -> dict[str, Any]:
    if result.has_errors:
        status = "FAIL"
    elif result.ready:
        status = "PASS" if output_written else "READY"
    else:
        status = "AWAITING_HUMAN_CONFIRMATION"
    return {
        "schema_version": "poc-03.r6-citation-import-result.v1",
        "generated_at": generated_at,
        "status": status,
        "privacy": {"workbook_content_committed": False, "queries_committed": False, "reviewer_names_committed": False, "output_dataset_committed": False},
        "summary": {
            "review_scope_count": result.scope_count,
            "preserved_case_count": result.preserved_count,
            "resolved_case_count": len(result.cases) if result.ready else 0,
            "pending_count": result.status_counts["PENDING"],
            "returned_count": result.status_counts["RETURNED"],
            "issue_count": len(result.issues),
            "output_written": output_written,
        },
        "issues": [issue.sanitized() for issue in result.issues],
        "conclusion": "R6 Golden Dataset export is allowed." if status in {"READY", "PASS"} else "R6 Golden Dataset export remains blocked.",
    }
