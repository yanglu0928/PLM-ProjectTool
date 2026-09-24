from __future__ import annotations

from collections import Counter
from dataclasses import dataclass
from pathlib import Path
from typing import Any

from jsonschema import Draft202012Validator, FormatChecker
from openpyxl import load_workbook

from .review_import import ImportIssue, _reviewed_at, _text


GLOBAL_CONFIRMATION = "确认全部AI建议"
ROW_DECISIONS = {"", "确认AI建议", "修改后确认", "退回"}
CLASSIFICATIONS = {
    "STANDARD_SATISFIED",
    "PARTIALLY_SATISFIED",
    "NON_STANDARD",
    "INSUFFICIENT_INFORMATION",
    "NO_RELIABLE_MATCH",
}
CLASSIFICATION_NAMES = {
    "STANDARD_SATISFIED": "标准满足",
    "PARTIALLY_SATISFIED": "部分满足",
    "NON_STANDARD": "非标准",
    "INSUFFICIENT_INFORMATION": "资料不足",
    "NO_RELIABLE_MATCH": "无可靠匹配",
}
CLASSIFICATION_VALUES = (
    {code: code for code in CLASSIFICATIONS}
    | {name: code for code, name in CLASSIFICATION_NAMES.items()}
    | {
        f"{name}（{code}）": code
        for code, name in CLASSIFICATION_NAMES.items()
    }
)
REVIEW_HEADERS = [
    "复核编号",
    "资料类别",
    "AI建议问题",
    "AI建议分类",
    "AI分类理由",
    "证据位置",
    "打开证据",
    "单条处理",
    "人工最终问题",
    "人工最终分类",
    "人工说明",
    "生效状态",
]
AUDIT_HEADERS = [
    "候选编号",
    "复核编号",
    "资料类别",
    "证据角色",
    "文档编号",
    "ChunkId",
    "正文Hash",
    "来源定位",
    "摘录本地修复",
    "关键词本地修复",
    "来源文件",
    "锁指纹",
]
REQUIRED_SOURCE_TYPES = {
    "STANDARD_CAPABILITY",
    "CONTRACT",
    "TECHNICAL_AGREEMENT",
    "SURVEY",
}


@dataclass
class HoldoutImportResult:
    cases: list[dict[str, Any]]
    issues: list[ImportIssue]
    status_counts: Counter[str]
    scope_count: int
    lock_id: str
    lock_fingerprint: str
    modified_count: int = 0

    @property
    def has_errors(self) -> bool:
        return bool(self.issues)

    @property
    def ready(self) -> bool:
        return (
            not self.has_errors
            and self.scope_count == 50
            and len(self.cases) == self.scope_count
            and self.status_counts["APPROVED"] == self.scope_count
            and self.status_counts["PENDING"] == 0
            and self.status_counts["RETURNED"] == 0
        )


def _issue(
    code: str,
    message: str,
    *,
    row: int | None = None,
    candidate_id: str | None = None,
    field: str | None = None,
) -> ImportIssue:
    return ImportIssue(code, message, row=row, candidate_id=candidate_id, field=field)


def _changed(row: int, field: str) -> ImportIssue:
    return _issue(
        "SOURCE_FIELD_CHANGED",
        "Immutable holdout review source field was changed",
        row=row,
        field=field,
    )


def _expected_audit_row(item: dict[str, Any], lock_fingerprint: str) -> list[str]:
    return [
        _text(item.get("candidate_id")),
        _text(item.get("task_id")),
        _text(item.get("source_type")),
        _text(item.get("evidence_role")),
        _text(item.get("document_id")),
        _text(item.get("chunk_id")),
        _text(item.get("text_sha256")),
        "；".join(_text(value) for value in item.get("source_locators") or []),
        "是" if item.get("evidence_quote_repaired") else "否",
        "是" if item.get("answer_terms_repaired") else "否",
        _text(item.get("source_name")),
        lock_fingerprint,
    ]


def import_holdout_review(
    workbook_path: Path,
    review_package: dict[str, Any],
    holdout_lock: dict[str, Any],
    *,
    timezone_name: str = "Asia/Shanghai",
) -> HoldoutImportResult:
    issues: list[ImportIssue] = []
    statuses: Counter[str] = Counter()
    package_cases = list(review_package.get("cases") or [])
    lock_candidates = list(holdout_lock.get("candidates") or [])
    scope_count = len(package_cases)
    lock_id = _text(holdout_lock.get("lock_id"))
    lock_fingerprint = _text(holdout_lock.get("lock_fingerprint"))

    if (
        review_package.get("schema_version") != "poc-03.holdout-review-package.v1"
        or holdout_lock.get("schema_version") != "poc-03.holdout-lock.v1"
        or review_package.get("lock_fingerprint") != lock_fingerprint
        or review_package.get("scope_count") != 50
        or holdout_lock.get("target_count") != 50
        or scope_count != 50
        or len(lock_candidates) != 50
    ):
        issues.append(
            _issue(
                "INVALID_HOLDOUT_SCOPE",
                "Holdout package and source lock must reference the same locked 50-case scope",
            )
        )

    package_task_ids = [_text(item.get("task_id")) for item in package_cases]
    package_candidate_ids = [_text(item.get("candidate_id")) for item in package_cases]
    lock_candidate_ids = [_text(item.get("candidate_id")) for item in lock_candidates]
    if (
        len(set(package_task_ids)) != len(package_task_ids)
        or len(set(package_candidate_ids)) != len(package_candidate_ids)
        or package_candidate_ids != lock_candidate_ids
    ):
        issues.append(
            _issue(
                "HOLDOUT_ID_MISMATCH",
                "Holdout task and candidate identifiers must be unique and match the source lock order",
            )
        )

    package_by_task = {
        _text(item.get("task_id")): item for item in package_cases
    }
    lock_by_candidate = {
        _text(item.get("candidate_id")): item for item in lock_candidates
    }

    workbook = load_workbook(workbook_path, read_only=True, data_only=False)
    try:
        required_sheets = {"确认清单", "技术底稿"}
        if not required_sheets.issubset(workbook.sheetnames):
            issues.append(
                _issue(
                    "MISSING_SHEET",
                    "Holdout workbook is missing the confirmation or technical audit sheet",
                )
            )
            return HoldoutImportResult(
                [], issues, statuses, scope_count, lock_id, lock_fingerprint
            )

        review_sheet = workbook["确认清单"]
        audit_sheet = workbook["技术底稿"]
        if [
            _text(review_sheet.cell(9, column).value)
            for column in range(1, len(REVIEW_HEADERS) + 1)
        ] != REVIEW_HEADERS:
            issues.append(
                _issue(
                    "HEADER_MISMATCH",
                    "Holdout confirmation headers differ from the locked template",
                    row=9,
                )
            )
        if [
            _text(audit_sheet.cell(6, column).value)
            for column in range(1, len(AUDIT_HEADERS) + 1)
        ] != AUDIT_HEADERS:
            issues.append(
                _issue(
                    "AUDIT_HEADER_MISMATCH",
                    "Holdout technical audit headers differ from the locked template",
                    row=6,
                )
            )

        global_value = _text(review_sheet["B5"].value)
        if global_value not in {"待确认", GLOBAL_CONFIRMATION}:
            issues.append(
                _issue(
                    "INVALID_GLOBAL_CONFIRMATION",
                    "Holdout batch confirmation has an invalid value",
                    row=5,
                    field="批量处理",
                )
            )
        global_confirmed = global_value == GLOBAL_CONFIRMATION
        reviewer = _text(review_sheet["D5"].value)
        reviewed_at = ""

        for index, item in enumerate(package_cases, start=7):
            expected = _expected_audit_row(item, lock_fingerprint)
            actual = [
                _text(audit_sheet.cell(index, column).value)
                for column in range(1, len(AUDIT_HEADERS) + 1)
            ]
            for column, (actual_value, expected_value) in enumerate(
                zip(actual, expected, strict=True), start=1
            ):
                if actual_value != _text(expected_value):
                    issues.append(_changed(index, AUDIT_HEADERS[column - 1]))

        seen: set[str] = set()
        output_cases: list[dict[str, Any]] = []
        modified_count = 0
        any_approved = False
        for row_number, row in enumerate(
            review_sheet.iter_rows(min_row=10, max_row=59, max_col=12, values_only=True),
            start=10,
        ):
            task_id = _text(row[0])
            item = package_by_task.get(task_id)
            if item is None or task_id in seen:
                issues.append(
                    _issue(
                        "UNKNOWN_OR_DUPLICATE_TASK",
                        "Holdout task is unknown or duplicated",
                        row=row_number,
                    )
                )
                continue
            seen.add(task_id)
            locked_values = [
                item.get("task_id"),
                item.get("source_type_display"),
                item.get("suggested_question"),
                item.get("suggested_classification_display"),
                item.get("classification_reason"),
                item.get("location_label"),
            ]
            for column, expected in enumerate(locked_values, start=1):
                if _text(row[column - 1]) != _text(expected):
                    issues.append(_changed(row_number, REVIEW_HEADERS[column - 1]))

            decision = _text(row[7])
            if decision not in ROW_DECISIONS:
                issues.append(
                    _issue(
                        "INVALID_ROW_DECISION",
                        "Holdout row decision has an invalid value",
                        row=row_number,
                        field=REVIEW_HEADERS[7],
                    )
                )
                statuses["INVALID"] += 1
                continue
            if decision == "退回":
                statuses["RETURNED"] += 1
                if len(_text(row[10])) < 2:
                    issues.append(
                        _issue(
                            "RETURN_REASON_REQUIRED",
                            "Returned holdout row requires a reason",
                            row=row_number,
                            field=REVIEW_HEADERS[10],
                        )
                    )
                continue

            approved = global_confirmed or decision in {"确认AI建议", "修改后确认"}
            if not approved:
                statuses["PENDING"] += 1
                continue
            any_approved = True
            if decision == "修改后确认":
                modified_count += 1
                query = _text(row[8])
                classification_value = _text(row[9])
                classification = CLASSIFICATION_VALUES.get(
                    classification_value,
                    CLASSIFICATION_VALUES.get(classification_value.upper(), ""),
                )
                explanation = _text(row[10])
                if len(query) < 10:
                    issues.append(
                        _issue(
                            "MANUAL_QUESTION_REQUIRED",
                            "Modified holdout row requires a substantive final question",
                            row=row_number,
                            field=REVIEW_HEADERS[8],
                        )
                    )
                if classification not in CLASSIFICATIONS:
                    issues.append(
                        _issue(
                            "INVALID_CLASSIFICATION",
                            "Modified holdout row requires one of the five locked classifications",
                            row=row_number,
                            field=REVIEW_HEADERS[9],
                        )
                    )
                if len(explanation) < 2:
                    issues.append(
                        _issue(
                            "MANUAL_EXPLANATION_REQUIRED",
                            "Modified holdout row requires a brief reason",
                            row=row_number,
                            field=REVIEW_HEADERS[10],
                        )
                    )
            else:
                query = _text(item.get("suggested_question"))
                classification = _text(item.get("suggested_classification"))

            candidate_id = _text(item.get("candidate_id"))
            locked = lock_by_candidate.get(candidate_id)
            if locked is None:
                issues.append(
                    _issue(
                        "LOCKED_CANDIDATE_MISSING",
                        "Holdout candidate is missing from the source lock",
                        row=row_number,
                        candidate_id=candidate_id,
                    )
                )
                continue
            package_lock_checks = {
                "source_type": item.get("source_type"),
                "evidence_role": item.get("evidence_role"),
                "document_id": item.get("document_id"),
                "chunk_id": item.get("chunk_id"),
                "text_sha256": item.get("text_sha256"),
                "source_locators": item.get("source_locators") or [],
            }
            if any(
                package_lock_checks[key] != locked.get(key)
                for key in package_lock_checks
            ):
                issues.append(
                    _issue(
                        "PACKAGE_LOCK_MISMATCH",
                        "Holdout review package differs from the immutable source lock",
                        row=row_number,
                        candidate_id=candidate_id,
                    )
                )
                continue
            if classification not in CLASSIFICATIONS:
                issues.append(
                    _issue(
                        "INVALID_CLASSIFICATION",
                        "Holdout row does not resolve to a final business classification",
                        row=row_number,
                        field=REVIEW_HEADERS[9],
                    )
                )
                continue
            source_locators = [
                _text(value) for value in locked.get("source_locators") or [] if _text(value)
            ]
            answer_terms = [
                _text(value) for value in item.get("answer_terms") or [] if _text(value)
            ]
            if not query or not source_locators or not answer_terms:
                issues.append(
                    _issue(
                        "INCOMPLETE_FINAL_CASE",
                        "Approved holdout row is missing its question, answer terms, or citation locator",
                        row=row_number,
                        candidate_id=candidate_id,
                    )
                )
                continue
            output_cases.append(
                {
                    "case_id": f"HO-{len(output_cases) + 1:04d}",
                    "candidate_id": candidate_id,
                    "query": query,
                    "scope": _text(locked.get("scope")),
                    "project_id": locked.get("project_id"),
                    "source_type": _text(locked.get("source_type")),
                    "evidence_role": _text(locked.get("evidence_role")),
                    "expected_classification": classification,
                    "expected_relevant_chunk_ids": [_text(locked.get("chunk_id"))],
                    "expected_answer_terms": list(dict.fromkeys(answer_terms)),
                    "expected_citations": [
                        {
                            "document_id": _text(locked.get("document_id")),
                            "chunk_id": _text(locked.get("chunk_id")),
                            "source_locator": locator,
                        }
                        for locator in source_locators
                    ],
                    "content_sha256": _text(locked.get("text_sha256")),
                    "review": {
                        "status": "APPROVED",
                        "reviewed_by": reviewer,
                        "reviewed_at": "",
                    },
                }
            )
            statuses["APPROVED"] += 1

        for task_id in sorted(set(package_by_task) - seen):
            issues.append(
                _issue(
                    "MISSING_TASK_ROW",
                    "Locked holdout task is missing",
                    candidate_id=task_id,
                )
            )

        if any_approved or global_confirmed:
            if not reviewer:
                issues.append(
                    _issue(
                        "REVIEWER_REQUIRED",
                        "Confirmed holdout workbook requires a reviewer",
                        row=5,
                        field="审核人",
                    )
                )
            try:
                reviewed_at = _reviewed_at(review_sheet["F5"].value, timezone_name)
            except (ValueError, TypeError, KeyError):
                issues.append(
                    _issue(
                        "INVALID_REVIEW_TIME",
                        "Confirmed holdout workbook requires a valid review date",
                        row=5,
                        field="审核日期",
                    )
                )

        if issues or statuses["PENDING"] or statuses["RETURNED"]:
            return HoldoutImportResult(
                [],
                issues,
                statuses,
                scope_count,
                lock_id,
                lock_fingerprint,
                modified_count,
            )
        for case in output_cases:
            case["review"]["reviewed_at"] = reviewed_at
        return HoldoutImportResult(
            output_cases,
            issues,
            statuses,
            scope_count,
            lock_id,
            lock_fingerprint,
            modified_count,
        )
    finally:
        workbook.close()


def build_holdout_dataset(
    result: HoldoutImportResult,
    holdout_lock: dict[str, Any],
    *,
    dataset_id: str,
    schema: dict[str, Any],
) -> dict[str, Any]:
    if not result.ready:
        raise ValueError("Holdout review is not ready for export")
    dataset = {
        "schema_version": "poc-03.holdout.v1",
        "dataset_id": dataset_id,
        "source_lock": {
            "lock_id": result.lock_id,
            "lock_fingerprint": result.lock_fingerprint,
            "source_type_quotas": holdout_lock.get("source_type_quotas") or {},
            "isolation": holdout_lock.get("isolation") or {},
        },
        "cases": result.cases,
    }
    validator = Draft202012Validator(schema, format_checker=FormatChecker())
    errors = sorted(
        validator.iter_errors(dataset), key=lambda item: list(item.absolute_path)
    )
    if errors:
        paths = [
            "/".join(str(part) for part in error.absolute_path) or "<root>"
            for error in errors
        ]
        raise ValueError(f"Holdout Dataset schema validation failed at: {paths}")
    return dataset


def audit_holdout_dataset(dataset: dict[str, Any]) -> dict[str, Any]:
    cases = list(dataset.get("cases") or [])
    source_type_counts = Counter(_text(case.get("source_type")) for case in cases)
    classification_counts = Counter(
        _text(case.get("expected_classification")) for case in cases
    )
    queries = [_text(case.get("query")) for case in cases]
    case_ids = [_text(case.get("case_id")) for case in cases]
    candidate_ids = [_text(case.get("candidate_id")) for case in cases]
    chunk_ids = [
        _text(chunk_id)
        for case in cases
        for chunk_id in case.get("expected_relevant_chunk_ids") or []
    ]
    quotas = {
        _text(key): int(value)
        for key, value in (dataset.get("source_lock", {}).get("source_type_quotas") or {}).items()
    }
    isolation = dataset.get("source_lock", {}).get("isolation") or {}
    citation_lock_aligned = all(
        len(case.get("expected_relevant_chunk_ids") or []) == 1
        and bool(case.get("expected_citations"))
        and all(
            citation.get("chunk_id") == case["expected_relevant_chunk_ids"][0]
            for citation in case.get("expected_citations") or []
        )
        for case in cases
    )
    checks = {
        "exactly_50_cases": len(cases) == 50,
        "unique_case_ids": len(case_ids) == len(set(case_ids)) == 50,
        "unique_candidate_ids": len(candidate_ids) == len(set(candidate_ids)) == 50,
        "distinct_query_per_case": len(queries) == len(set(queries)) == 50,
        "unique_locked_chunk_per_case": len(chunk_ids) == len(set(chunk_ids)) == 50,
        "all_source_types_present": REQUIRED_SOURCE_TYPES.issubset(source_type_counts),
        "source_type_quotas_match": dict(source_type_counts) == quotas,
        "all_final_classifications_present": CLASSIFICATIONS.issubset(
            classification_counts
        ),
        "all_cases_human_approved": all(
            case.get("review", {}).get("status") == "APPROVED" for case in cases
        ),
        "all_cases_project_scoped": all(
            case.get("scope") == "PROJECT" and bool(_text(case.get("project_id")))
            for case in cases
        ),
        "citation_lock_aligned": citation_lock_aligned,
        "holdout_isolation_enforced": all(
            isolation.get(key) is True
            for key in (
                "case_and_query_disjoint",
                "chunk_id_disjoint",
                "source_locator_disjoint",
            )
        ),
    }
    return {
        "schema_version": "poc-03.holdout-coverage-result.v1",
        "status": "PASS" if all(checks.values()) else "FAIL",
        "summary": {
            "case_count": len(cases),
            "unique_query_count": len(set(queries)),
            "unique_candidate_count": len(set(candidate_ids)),
            "unique_chunk_count": len(set(chunk_ids)),
            "source_type_counts": dict(sorted(source_type_counts.items())),
            "classification_counts": dict(sorted(classification_counts.items())),
        },
        "checks": checks,
        "missing_source_types": sorted(REQUIRED_SOURCE_TYPES - set(source_type_counts)),
        "missing_classifications": sorted(CLASSIFICATIONS - set(classification_counts)),
        "privacy": {
            "queries_committed": False,
            "answer_terms_committed": False,
            "reviewer_names_committed": False,
            "dataset_committed": False,
            "source_names_committed": False,
        },
    }


def build_sanitized_holdout_report(
    result: HoldoutImportResult,
    *,
    generated_at: str,
    schema_valid: bool,
    coverage_status: str,
    output_written: bool,
) -> dict[str, Any]:
    if result.has_errors or coverage_status == "FAIL":
        status = "FAIL"
    elif result.ready and schema_valid and output_written and coverage_status == "PASS":
        status = "PASS"
    elif result.ready:
        status = "READY"
    else:
        status = "AWAITING_HUMAN_CONFIRMATION"
    return {
        "schema_version": "poc-03.holdout-import-result.v1",
        "generated_at": generated_at,
        "status": status,
        "privacy": {
            "workbook_content_committed": False,
            "queries_committed": False,
            "answer_terms_committed": False,
            "reviewer_names_committed": False,
            "source_names_committed": False,
            "output_dataset_committed": False,
        },
        "summary": {
            "review_scope_count": result.scope_count,
            "approved_count": result.status_counts["APPROVED"],
            "pending_count": result.status_counts["PENDING"],
            "returned_count": result.status_counts["RETURNED"],
            "modified_count": result.modified_count,
            "issue_count": len(result.issues),
            "schema_valid": schema_valid,
            "coverage_status": coverage_status,
            "output_written": output_written,
        },
        "source_lock": {
            "lock_id": result.lock_id,
            "lock_fingerprint": result.lock_fingerprint,
        },
        "issues": [issue.sanitized() for issue in result.issues],
        "conclusion": (
            "Independent holdout review import, schema validation, and coverage audit passed."
            if status == "PASS"
            else "Independent holdout dataset export remains blocked."
        ),
    }
