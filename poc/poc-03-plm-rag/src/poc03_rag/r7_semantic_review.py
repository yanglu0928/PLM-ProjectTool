from __future__ import annotations

import copy
from collections import Counter
from dataclasses import dataclass
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from .review_import import ImportIssue, _reviewed_at, _split_values, _text


GLOBAL_CONFIRMATION = "确认全部AI建议"
ROW_DECISIONS = {"", "确认AI建议", "修改后确认", "退回"}
CLASSIFICATIONS = {
    "STANDARD_SATISFIED",
    "PARTIALLY_SATISFIED",
    "NON_STANDARD",
    "INSUFFICIENT_INFORMATION",
    "NO_RELIABLE_MATCH",
}
CLASSIFICATION_NAMES = {
    "STANDARD_SATISFIED": "标准满足",
    "PARTIALLY_SATISFIED": "部分满足",
    "NON_STANDARD": "非标准",
    "INSUFFICIENT_INFORMATION": "资料不足",
    "NO_RELIABLE_MATCH": "无可靠匹配",
}
CLASSIFICATION_VALUES = {
    code: code for code in CLASSIFICATIONS
} | {
    name: code for code, name in CLASSIFICATION_NAMES.items()
} | {
    f"{name}（{code}）": code for code, name in CLASSIFICATION_NAMES.items()
}
HEADERS = [
    "复核编号",
    "CaseId",
    "资料类别",
    "R6当前分类",
    "Prompt v2分类",
    "问题诊断",
    "AI建议判定目标",
    "AI建议分类",
    "AI分类理由",
    "AI建议引用Chunk",
    "打开证据",
    "单条决定",
    "人工判定目标",
    "人工最终分类",
    "人工引用Chunk",
    "人工说明",
    "生效状态",
]


@dataclass
class R7ImportResult:
    cases: list[dict[str, Any]]
    issues: list[ImportIssue]
    status_counts: Counter[str]
    confirmed: bool
    scope_count: int
    query_change_count: int = 0
    classification_change_count: int = 0
    citation_change_count: int = 0

    @property
    def has_errors(self) -> bool:
        return bool(self.issues)

    @property
    def ready(self) -> bool:
        return (
            self.confirmed
            and not self.has_errors
            and self.status_counts["PENDING"] == 0
            and self.status_counts["RETURNED"] == 0
            and len(self.cases) == self.scope_count
        )


def _changed(row: int, field: str) -> ImportIssue:
    return ImportIssue(
        "SOURCE_FIELD_CHANGED",
        "Immutable R7 review source field was changed",
        row=row,
        field=field,
    )


def import_r7_semantic_review(
    workbook_path: Path,
    golden: dict[str, Any],
    package: dict[str, Any],
    *,
    timezone_name: str = "Asia/Shanghai",
) -> R7ImportResult:
    issues: list[ImportIssue] = []
    statuses: Counter[str] = Counter()
    source_cases = list(golden.get("cases") or [])
    source_by_id = {_text(case.get("case_id")): case for case in source_cases}
    package_cases = list(package.get("cases") or [])
    locked_ids = list(package.get("scope_case_ids") or [])
    expected_ids = [case.get("case_id") for case in source_cases]
    if (
        len(source_cases) != 120
        or locked_ids != expected_ids
        or locked_ids != [item.get("case_id") for item in package_cases]
    ):
        issues.append(
            ImportIssue(
                "INVALID_R7_SCOPE",
                "R7 package must contain the locked 120-case R6 scope in source order",
            )
        )
    if package.get("source_dataset_id") != golden.get("dataset_id"):
        issues.append(
            ImportIssue(
                "SOURCE_DATASET_MISMATCH",
                "R7 package does not reference the supplied R6 dataset",
            )
        )
    if len(source_by_id) != len(source_cases):
        issues.append(ImportIssue("DUPLICATE_SOURCE_CASE", "R6 dataset contains duplicate CaseId values"))

    workbook = load_workbook(workbook_path, read_only=True, data_only=False)
    try:
        if "R7确认" not in workbook.sheetnames:
            issues.append(ImportIssue("MISSING_SHEET", "R7 workbook is missing R7确认"))
            return R7ImportResult([], issues, statuses, False, len(locked_ids))
        sheet = workbook["R7确认"]
        if [_text(sheet.cell(9, col).value) for col in range(1, 18)] != HEADERS:
            issues.append(ImportIssue("HEADER_MISMATCH", "R7 headers differ from the locked template", row=9))

        global_value = _text(sheet["B5"].value)
        if global_value not in {"未确认", GLOBAL_CONFIRMATION}:
            issues.append(
                ImportIssue(
                    "INVALID_GLOBAL_CONFIRMATION",
                    "R7 batch confirmation has an invalid value",
                    row=5,
                    field="批量确认",
                )
            )
        global_confirmed = global_value == GLOBAL_CONFIRMATION
        reviewer = _text(sheet["D5"].value)
        reviewed_at = ""

        by_task = {item["task_id"]: item for item in package_cases}
        seen: set[str] = set()
        updates: dict[str, tuple[str, str, list[str]]] = {}
        any_approved = False
        for row_number, row in enumerate(
            sheet.iter_rows(min_row=10, max_row=129, max_col=17, values_only=True),
            start=10,
        ):
            task_id = _text(row[0])
            if not task_id:
                continue
            item = by_task.get(task_id)
            if item is None or task_id in seen:
                issues.append(
                    ImportIssue(
                        "UNKNOWN_OR_DUPLICATE_TASK",
                        "R7 task is unknown or duplicated",
                        row=row_number,
                    )
                )
                continue
            seen.add(task_id)
            locked = [
                item["task_id"],
                item["case_id"],
                item["source_type"],
                item.get("current_classification_display", item["current_classification"]),
                item.get("prompt_v2_classification_display", item["prompt_v2_classification"]),
                item["issue_reason"],
                item["suggested_target"],
                item.get("suggested_classification_display", item["suggested_classification"]),
                item["classification_rationale"],
                "；".join(item["suggested_chunk_ids"]),
            ]
            for column, expected in enumerate(locked, start=1):
                if _text(row[column - 1]) != _text(expected):
                    issues.append(_changed(row_number, HEADERS[column - 1]))

            decision = _text(row[11])
            if decision not in ROW_DECISIONS:
                issues.append(
                    ImportIssue(
                        "INVALID_ROW_DECISION",
                        "R7 row decision has an invalid value",
                        row=row_number,
                        field=HEADERS[11],
                    )
                )
                statuses["INVALID"] += 1
                continue
            if decision == "退回":
                statuses["RETURNED"] += 1
                continue
            approved = global_confirmed or decision in {"确认AI建议", "修改后确认"}
            if not approved:
                statuses["PENDING"] += 1
                continue
            any_approved = True
            if decision == "修改后确认":
                target = _text(row[12])
                classification_value = _text(row[13])
                classification = CLASSIFICATION_VALUES.get(
                    classification_value,
                    CLASSIFICATION_VALUES.get(classification_value.upper(), ""),
                )
                chunk_ids = _split_values(row[14])
                explanation = _text(row[15])
                if len(target) < 20:
                    issues.append(
                        ImportIssue(
                            "MANUAL_TARGET_REQUIRED",
                            "Modified R7 row requires a substantive determination target",
                            row=row_number,
                            field=HEADERS[12],
                        )
                    )
                if classification not in CLASSIFICATIONS:
                    issues.append(
                        ImportIssue(
                            "INVALID_CLASSIFICATION",
                            "Modified R7 row requires one of the five locked classifications",
                            row=row_number,
                            field=HEADERS[13],
                        )
                    )
                if not chunk_ids:
                    issues.append(
                        ImportIssue(
                            "MANUAL_CITATION_REQUIRED",
                            "Modified R7 row requires at least one reviewed citation",
                            row=row_number,
                            field=HEADERS[14],
                        )
                    )
                if len(explanation) < 6:
                    issues.append(
                        ImportIssue(
                            "MANUAL_EXPLANATION_REQUIRED",
                            "Modified R7 row requires a brief reason",
                            row=row_number,
                            field=HEADERS[15],
                        )
                    )
            else:
                target = item["suggested_target"]
                classification = item["suggested_classification"]
                chunk_ids = list(item["suggested_chunk_ids"])

            unknown = sorted(set(chunk_ids) - set(item["allowed_chunk_ids"]))
            if unknown:
                issues.append(
                    ImportIssue(
                        "CITATION_OUTSIDE_R7_EVIDENCE",
                        "R7 citation is outside the displayed evidence set",
                        row=row_number,
                        field=HEADERS[14],
                    )
                )
            if target and classification in CLASSIFICATIONS and chunk_ids and not unknown:
                updates[item["case_id"]] = (target, classification, chunk_ids)
                statuses["APPROVED"] += 1

        for task_id in sorted(set(by_task) - seen):
            issues.append(
                ImportIssue(
                    "MISSING_TASK_ROW",
                    "Locked R7 task is missing",
                    candidate_id=task_id,
                )
            )

        confirmed = (
            (global_confirmed or statuses["PENDING"] == 0)
            and statuses["RETURNED"] == 0
            and len(updates) == len(locked_ids)
        )
        if any_approved or global_confirmed:
            if not reviewer:
                issues.append(
                    ImportIssue(
                        "REVIEWER_REQUIRED",
                        "Confirmed R7 workbook requires a reviewer",
                        row=5,
                        field="确认人",
                    )
                )
            try:
                reviewed_at = _reviewed_at(sheet["F5"].value, timezone_name)
            except (ValueError, TypeError, KeyError):
                issues.append(
                    ImportIssue(
                        "INVALID_REVIEW_TIME",
                        "Confirmed R7 workbook requires a valid review date",
                        row=5,
                        field="确认日期",
                    )
                )

        if not confirmed or issues:
            return R7ImportResult([], issues, statuses, confirmed, len(locked_ids))

        citation_map: dict[str, dict[str, Any]] = {}
        for item in package_cases:
            for chunk in item["candidate_chunks"]:
                citation_map[chunk["chunk_id"]] = chunk

        output_cases: list[dict[str, Any]] = []
        query_changes = 0
        classification_changes = 0
        citation_changes = 0
        for source in source_cases:
            case = copy.deepcopy(source)
            target, classification, chunk_ids = updates[source["case_id"]]
            query_changes += target != source.get("query")
            classification_changes += classification != source.get("expected_classification")
            citation_changes += chunk_ids != list(source.get("expected_relevant_chunk_ids") or [])
            case["query"] = target
            case["expected_classification"] = classification
            case["expected_relevant_chunk_ids"] = chunk_ids
            case["expected_citations"] = [
                {
                    "document_id": citation_map[chunk_id]["document_id"],
                    "chunk_id": chunk_id,
                    "source_locator": locator,
                }
                for chunk_id in chunk_ids
                for locator in citation_map[chunk_id]["source_locators"]
            ]
            case["review"] = {
                "status": "APPROVED",
                "reviewed_by": reviewer,
                "reviewed_at": reviewed_at,
            }
            output_cases.append(case)
        return R7ImportResult(
            output_cases,
            issues,
            statuses,
            True,
            len(locked_ids),
            query_changes,
            classification_changes,
            citation_changes,
        )
    finally:
        workbook.close()


def build_r7_dataset(
    golden: dict[str, Any],
    result: R7ImportResult,
    *,
    dataset_id: str,
) -> dict[str, Any]:
    if not result.ready:
        raise ValueError("R7 review is not ready for export")
    dataset = copy.deepcopy(golden)
    dataset["dataset_id"] = dataset_id
    dataset["cases"] = result.cases
    return dataset


def build_sanitized_r7_report(
    result: R7ImportResult,
    *,
    generated_at: str,
    output_written: bool,
) -> dict[str, Any]:
    if result.has_errors:
        status = "FAIL"
    elif result.ready:
        status = "PASS" if output_written else "READY"
    else:
        status = "AWAITING_HUMAN_CONFIRMATION"
    return {
        "schema_version": "poc-03.r7-semantic-import-result.v1",
        "generated_at": generated_at,
        "status": status,
        "privacy": {
            "workbook_content_committed": False,
            "queries_committed": False,
            "reviewer_names_committed": False,
            "output_dataset_committed": False,
        },
        "summary": {
            "review_scope_count": result.scope_count,
            "resolved_case_count": len(result.cases) if result.ready else 0,
            "pending_count": result.status_counts["PENDING"],
            "returned_count": result.status_counts["RETURNED"],
            "issue_count": len(result.issues),
            "query_change_count": result.query_change_count,
            "classification_change_count": result.classification_change_count,
            "citation_change_count": result.citation_change_count,
            "output_written": output_written,
        },
        "validation_limit": (
            "R7 is an AI-assisted calibration dataset. Results measured on the same 120 cases "
            "cannot close P03-A12/P03-A13; an independently constructed holdout set is required."
        ),
        "issues": [issue.sanitized() for issue in result.issues],
        "conclusion": (
            "R7 calibration dataset export is allowed; independent holdout remains mandatory."
            if status in {"READY", "PASS"}
            else "R7 dataset export remains blocked."
        ),
    }
