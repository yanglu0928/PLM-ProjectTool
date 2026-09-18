from __future__ import annotations

from collections import Counter
from pathlib import Path
from typing import Any, Iterable

from openpyxl import load_workbook

from .review_import import (
    CLASSIFICATIONS,
    SOURCE_TYPES,
    ImportIssue,
    ReviewImportResult,
    _candidate_case_id,
    _confirmed_locators,
    _expected_source,
    _reviewed_at,
    _split_values,
    _text,
)


ACTION_SHEET = "确认清单"
ACTION_HEADERS = [
    "待办编号",
    "资料类别",
    "风险",
    "AI分析问题",
    "AI建议结论",
    "打开证据",
    "处理结果",
    "修改后结论 / 补充说明",
    "审核人",
    "审核日期",
    "当前状态",
]
DECISION_STATUS = {
    "": "PENDING",
    "同意AI建议": "APPROVED",
    "修改后确认": "PENDING",
    "退回重新分析": "RETURNED",
    "暂不处理": "PENDING",
}


def import_action_list(
    workbook_path: Path,
    candidates: Iterable[dict[str, Any]],
    tasks: Iterable[dict[str, Any]],
    *,
    timezone_name: str = "Asia/Shanghai",
) -> ReviewImportResult:
    candidate_by_id = {
        _text(item.get("candidate_id")): item for item in candidates
    }
    task_by_id = {_text(item.get("task_id")): item for item in tasks}
    issues: list[ImportIssue] = []
    if not candidate_by_id or "" in candidate_by_id:
        issues.append(
            ImportIssue(
                code="INVALID_CANDIDATE_SOURCE",
                message="Candidate source contains a missing candidate_id",
            )
        )
    if not task_by_id or "" in task_by_id:
        issues.append(
            ImportIssue(
                code="INVALID_TASK_SOURCE",
                message="Task source contains a missing task_id",
            )
        )

    workbook = load_workbook(workbook_path, read_only=True, data_only=False)
    try:
        if ACTION_SHEET not in workbook.sheetnames:
            issues.append(
                ImportIssue(
                    code="MISSING_SHEET",
                    message=f"Workbook is missing required sheet: {ACTION_SHEET}",
                )
            )
            return ReviewImportResult([], issues, Counter(), 0)
        sheet = workbook[ACTION_SHEET]
        actual_headers = [_text(sheet.cell(8, column).value) for column in range(1, 12)]
        if actual_headers != ACTION_HEADERS:
            issues.append(
                ImportIssue(
                    code="HEADER_MISMATCH",
                    message="Action-list headers do not match the locked R4 template",
                    row=8,
                )
            )
            return ReviewImportResult([], issues, Counter(), 0)

        cases: list[dict[str, Any]] = []
        status_counts: Counter[str] = Counter()
        seen_task_ids: set[str] = set()
        row_count = 0
        for row_number, row in enumerate(
            sheet.iter_rows(min_row=9, max_row=128, max_col=11, values_only=True),
            start=9,
        ):
            task_id = _text(row[0])
            if not task_id:
                if any(_text(value) for value in row):
                    issues.append(
                        ImportIssue(
                            code="MISSING_TASK_ID",
                            message="Non-empty action row has no task_id",
                            row=row_number,
                        )
                    )
                continue
            row_count += 1
            if task_id in seen_task_ids:
                issues.append(
                    ImportIssue(
                        code="DUPLICATE_TASK_ROW",
                        message="Workbook contains a duplicate task row",
                        row=row_number,
                    )
                )
                continue
            seen_task_ids.add(task_id)
            task = task_by_id.get(task_id)
            if task is None:
                issues.append(
                    ImportIssue(
                        code="UNKNOWN_TASK",
                        message="Workbook task_id does not exist in the locked task source",
                        row=row_number,
                    )
                )
                continue

            candidate_id = _text(task.get("candidate_id"))
            candidate = candidate_by_id.get(candidate_id)
            if candidate is None:
                issues.append(
                    ImportIssue(
                        code="UNKNOWN_CANDIDATE",
                        message="Task candidate_id does not exist in source candidates",
                        row=row_number,
                        candidate_id=candidate_id or None,
                    )
                )
                continue
            expected = _expected_source(candidate)
            source_checks = {
                "资料类别": (_text(row[1]), expected["source_corpus"]),
                "AI分析问题": (_text(row[3]), _text(task.get("ai_finding"))),
                "任务资料类别": (_text(task.get("source_corpus")), expected["source_corpus"]),
                "文档编号": (_text(task.get("document_id")), expected["document_id"]),
                "Chunk编号": (_text(task.get("chunk_id")), expected["chunk_id"]),
                "正文SHA-256": (
                    _text(task.get("candidate_content_sha256")),
                    expected["candidate_content_sha256"],
                ),
            }
            row_source_issues = [
                ImportIssue(
                    code="SOURCE_FIELD_CHANGED",
                    message="Immutable action-list source field was changed",
                    row=row_number,
                    candidate_id=candidate_id,
                    field=field,
                )
                for field, (actual, locked) in source_checks.items()
                if actual != locked
            ]
            issues.extend(row_source_issues)

            decision = _text(row[6])
            if decision not in DECISION_STATUS:
                issues.append(
                    ImportIssue(
                        code="INVALID_HUMAN_DECISION",
                        message="Action-list decision is not an allowed value",
                        row=row_number,
                        candidate_id=candidate_id,
                        field="处理结果",
                    )
                )
                status_counts["INVALID"] += 1
                continue
            review_status = DECISION_STATUS[decision]
            status_counts[review_status] += 1
            if decision == "修改后确认":
                continue
            if review_status != "APPROVED":
                continue

            query = _text(task.get("ai_finding"))
            source_type = _text(task.get("source_type")).upper()
            classification = _text(task.get("classification")).upper()
            answer_terms = _split_values(task.get("answer_terms"))
            confirmed_locators = _confirmed_locators(
                task.get("confirmed_locator"), expected["source_locators"]
            )
            reviewed_by = _text(row[8])
            approved_issues: list[ImportIssue] = []
            required = {
                "AI分析问题": query,
                "来源类型": source_type,
                "建议分类": classification,
                "答案术语": answer_terms,
                "引用定位": confirmed_locators,
                "审核人": reviewed_by,
                "审核日期": row[9],
            }
            for field, value in required.items():
                if value is None or value == "" or value == []:
                    approved_issues.append(
                        ImportIssue(
                            code="APPROVED_FIELD_MISSING",
                            message="Approved action is missing a required field",
                            row=row_number,
                            candidate_id=candidate_id,
                            field=field,
                        )
                    )
            if query and len(query) < 2:
                approved_issues.append(
                    ImportIssue(
                        code="QUERY_TOO_SHORT",
                        message="Approved query must contain at least two characters",
                        row=row_number,
                        candidate_id=candidate_id,
                        field="AI分析问题",
                    )
                )
            if source_type and source_type not in SOURCE_TYPES:
                approved_issues.append(
                    ImportIssue(
                        code="INVALID_SOURCE_TYPE",
                        message="Approved action has an invalid source type",
                        row=row_number,
                        candidate_id=candidate_id,
                        field="来源类型",
                    )
                )
            if classification and classification not in CLASSIFICATIONS:
                approved_issues.append(
                    ImportIssue(
                        code="INVALID_CLASSIFICATION",
                        message="Approved action has an invalid classification",
                        row=row_number,
                        candidate_id=candidate_id,
                        field="建议分类",
                    )
                )
            if any(locator not in expected["source_locators"] for locator in confirmed_locators):
                approved_issues.append(
                    ImportIssue(
                        code="CITATION_OUTSIDE_SOURCE",
                        message="Confirmed citation is outside the candidate source locators",
                        row=row_number,
                        candidate_id=candidate_id,
                        field="引用定位",
                    )
                )
            try:
                reviewed_at = _reviewed_at(row[9], timezone_name)
            except (ValueError, TypeError, KeyError) as exc:
                reviewed_at = ""
                approved_issues.append(
                    ImportIssue(
                        code="INVALID_REVIEW_TIME",
                        message=f"Review time is invalid: {type(exc).__name__}",
                        row=row_number,
                        candidate_id=candidate_id,
                        field="审核日期",
                    )
                )
            issues.extend(approved_issues)
            if row_source_issues or approved_issues:
                continue

            cases.append(
                {
                    "case_id": _candidate_case_id(candidate_id),
                    "query": query,
                    "scope": expected["scope"],
                    "project_id": expected["project_id"] or None,
                    "source_type": source_type,
                    "expected_classification": classification,
                    "expected_relevant_chunk_ids": expected["relevant_chunk_ids"],
                    "expected_answer_terms": answer_terms,
                    "expected_citations": [
                        {
                            "document_id": expected["document_id"],
                            "chunk_id": expected["chunk_id"],
                            "source_locator": locator,
                        }
                        for locator in confirmed_locators
                    ],
                    "review": {
                        "status": "APPROVED",
                        "reviewed_by": reviewed_by,
                        "reviewed_at": reviewed_at,
                    },
                }
            )

        for task_id in sorted(set(task_by_id) - seen_task_ids):
            issues.append(
                ImportIssue(
                    code="MISSING_TASK_ROW",
                    message="Locked task is missing from the action-list workbook",
                    candidate_id=_text(task_by_id[task_id].get("candidate_id")) or None,
                )
            )
        return ReviewImportResult(cases, issues, status_counts, row_count)
    finally:
        workbook.close()
