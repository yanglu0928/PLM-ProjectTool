from __future__ import annotations

import json
import re
from collections import Counter
from dataclasses import dataclass
from datetime import UTC, date, datetime, time, timedelta, timezone
from pathlib import Path
from typing import Any, Iterable
from zoneinfo import ZoneInfo

from jsonschema import Draft202012Validator, FormatChecker
from openpyxl import load_workbook


REVIEW_SHEET = "候选评审"
HEADERS = [
    "候选编号",
    "资料来源",
    "范围",
    "ProjectId",
    "文档编号",
    "Chunk 编号",
    "建议引用定位",
    "候选正文",
    "正文 SHA-256",
    "人工查询",
    "来源类型",
    "预期分类",
    "答案关键术语",
    "确认引用定位",
    "审核状态",
    "审核人",
    "审核时间",
    "审核备注",
    "完备性",
]
SOURCE_TYPES = {
    "STANDARD_CAPABILITY",
    "CONTRACT",
    "TECHNICAL_AGREEMENT",
    "SURVEY",
}
CLASSIFICATIONS = {
    "STANDARD_SATISFIED",
    "PARTIALLY_SATISFIED",
    "NON_STANDARD",
    "INSUFFICIENT_INFORMATION",
    "NO_RELIABLE_MATCH",
    "HUMAN_CONFIRMATION_REQUIRED",
}
REVIEW_STATUSES = {"PENDING", "APPROVED", "RETURNED"}


@dataclass(frozen=True)
class ImportIssue:
    code: str
    message: str
    row: int | None = None
    candidate_id: str | None = None
    field: str | None = None

    def sanitized(self) -> dict[str, Any]:
        return {
            key: value
            for key, value in {
                "code": self.code,
                "message": self.message,
                "row": self.row,
                "candidate_id": self.candidate_id,
                "field": self.field,
            }.items()
            if value is not None
        }


@dataclass
class ReviewImportResult:
    cases: list[dict[str, Any]]
    issues: list[ImportIssue]
    status_counts: Counter[str]
    row_count: int

    @property
    def has_errors(self) -> bool:
        return bool(self.issues)


def load_candidate_records(path: Path) -> list[dict[str, Any]]:
    records: list[dict[str, Any]] = []
    with path.open("r", encoding="utf-8") as stream:
        for line_number, raw_line in enumerate(stream, start=1):
            line = raw_line.strip()
            if not line:
                continue
            record = json.loads(line)
            if not isinstance(record, dict):
                raise ValueError(f"Candidate line {line_number} is not an object")
            records.append(record)
    return records


def _text(value: Any) -> str:
    return "" if value is None else str(value).strip()


def _split_values(value: Any) -> list[str]:
    parts = re.split(r"[；;\n\r]+", _text(value))
    return list(dict.fromkeys(part.strip() for part in parts if part.strip()))


def _candidate_case_id(candidate_id: str) -> str:
    match = re.fullmatch(r"GD-C-(\d{4})", candidate_id)
    if not match:
        raise ValueError("candidate_id must match GD-C-0000")
    return f"GD-{match.group(1)}"


def _reviewed_at(value: Any, timezone_name: str) -> str:
    try:
        local_timezone = ZoneInfo(timezone_name)
    except KeyError:
        if timezone_name != "Asia/Shanghai":
            raise
        local_timezone = timezone(timedelta(hours=8), name="Asia/Shanghai")
    if isinstance(value, datetime):
        parsed = value
    elif isinstance(value, date):
        parsed = datetime.combine(value, time.min)
    else:
        raw = _text(value)
        if not raw:
            raise ValueError("reviewed_at is required")
        dotted_date = re.fullmatch(r"(\d{4})[./](\d{1,2})[./](\d{1,2})", raw)
        if dotted_date:
            year, month, day = (int(part) for part in dotted_date.groups())
            parsed = datetime(year, month, day)
        else:
            parsed = datetime.fromisoformat(raw.replace("Z", "+00:00"))
    if parsed.tzinfo is None:
        parsed = parsed.replace(tzinfo=local_timezone)
    return parsed.astimezone(UTC).isoformat().replace("+00:00", "Z")


def _expected_source(candidate: dict[str, Any]) -> dict[str, Any]:
    citations = candidate.get("expected_citations") or []
    if len(citations) != 1:
        raise ValueError("candidate must have exactly one source citation")
    citation = citations[0]
    chunk_ids = candidate.get("expected_relevant_chunk_ids") or []
    if len(chunk_ids) != 1:
        raise ValueError("candidate must have exactly one relevant chunk")
    candidate_id = _text(candidate.get("candidate_id"))
    _candidate_case_id(candidate_id)
    return {
        "candidate_id": candidate_id,
        "source_corpus": _text(candidate.get("source_corpus")),
        "scope": _text(candidate.get("scope")),
        "project_id": _text(candidate.get("project_id")),
        "document_id": _text(citation.get("document_id")),
        "chunk_id": _text(citation.get("chunk_id")),
        "source_locators": [_text(item) for item in citation.get("source_locators") or []],
        "candidate_content": _text(candidate.get("candidate_content")),
        "candidate_content_sha256": _text(candidate.get("candidate_content_sha256")),
        "relevant_chunk_ids": [_text(item) for item in chunk_ids],
    }


def _confirmed_locators(value: Any, source_locators: list[str]) -> list[str]:
    text = _text(value)
    if text in {"确认", "全部确认", "确认全部建议定位"}:
        return list(source_locators)
    return _split_values(value)


def _source_checks(
    row_values: tuple[Any, ...],
    expected: dict[str, Any],
    *,
    row_number: int,
) -> list[ImportIssue]:
    candidate_id = expected["candidate_id"]
    checks = {
        "资料来源": (_text(row_values[1]), expected["source_corpus"]),
        "范围": (_text(row_values[2]), expected["scope"]),
        "ProjectId": (_text(row_values[3]), expected["project_id"]),
        "文档编号": (_text(row_values[4]), expected["document_id"]),
        "Chunk 编号": (_text(row_values[5]), expected["chunk_id"]),
        "候选正文": (_text(row_values[7]), expected["candidate_content"]),
        "正文 SHA-256": (_text(row_values[8]), expected["candidate_content_sha256"]),
    }
    issues = [
        ImportIssue(
            code="SOURCE_FIELD_CHANGED",
            message="Immutable candidate source field was changed",
            row=row_number,
            candidate_id=candidate_id,
            field=field,
        )
        for field, (actual, source) in checks.items()
        if actual != source
    ]
    if _split_values(row_values[6]) != expected["source_locators"]:
        issues.append(
            ImportIssue(
                code="SOURCE_FIELD_CHANGED",
                message="Immutable candidate source locator was changed",
                row=row_number,
                candidate_id=candidate_id,
                field="建议引用定位",
            )
        )
    return issues


def import_review_workbook(
    workbook_path: Path,
    candidates: Iterable[dict[str, Any]],
    *,
    timezone_name: str = "Asia/Shanghai",
) -> ReviewImportResult:
    candidate_list = list(candidates)
    candidate_by_id: dict[str, dict[str, Any]] = {}
    issues: list[ImportIssue] = []
    for candidate in candidate_list:
        candidate_id = _text(candidate.get("candidate_id"))
        if not candidate_id or candidate_id in candidate_by_id:
            issues.append(
                ImportIssue(
                    code="INVALID_CANDIDATE_SOURCE",
                    message="Candidate source contains a missing or duplicate candidate_id",
                    candidate_id=candidate_id or None,
                )
            )
            continue
        candidate_by_id[candidate_id] = candidate

    workbook = load_workbook(workbook_path, read_only=True, data_only=False)
    try:
        if REVIEW_SHEET not in workbook.sheetnames:
            issues.append(
                ImportIssue(
                    code="MISSING_SHEET",
                    message=f"Workbook is missing required sheet: {REVIEW_SHEET}",
                )
            )
            return ReviewImportResult([], issues, Counter(), 0)
        sheet = workbook[REVIEW_SHEET]
        actual_headers = [_text(sheet.cell(5, column).value) for column in range(1, 20)]
        if actual_headers != HEADERS:
            issues.append(
                ImportIssue(
                    code="HEADER_MISMATCH",
                    message="Review workbook headers do not match the locked template",
                    row=5,
                )
            )
            return ReviewImportResult([], issues, Counter(), 0)

        seen_ids: set[str] = set()
        cases: list[dict[str, Any]] = []
        status_counts: Counter[str] = Counter()
        row_count = 0
        for row_number, row in enumerate(
            sheet.iter_rows(min_row=6, max_col=19, values_only=True), start=6
        ):
            candidate_id = _text(row[0])
            if not candidate_id:
                if any(_text(value) for value in row):
                    issues.append(
                        ImportIssue(
                            code="MISSING_CANDIDATE_ID",
                            message="Non-empty review row has no candidate_id",
                            row=row_number,
                        )
                    )
                continue
            row_count += 1
            if candidate_id in seen_ids:
                issues.append(
                    ImportIssue(
                        code="DUPLICATE_CANDIDATE_ROW",
                        message="Workbook contains a duplicate candidate row",
                        row=row_number,
                        candidate_id=candidate_id,
                    )
                )
                continue
            seen_ids.add(candidate_id)
            candidate = candidate_by_id.get(candidate_id)
            if candidate is None:
                issues.append(
                    ImportIssue(
                        code="UNKNOWN_CANDIDATE",
                        message="Workbook candidate_id does not exist in source candidates",
                        row=row_number,
                        candidate_id=candidate_id,
                    )
                )
                continue
            try:
                expected = _expected_source(candidate)
            except ValueError as exc:
                issues.append(
                    ImportIssue(
                        code="INVALID_CANDIDATE_SOURCE",
                        message=str(exc),
                        row=row_number,
                        candidate_id=candidate_id,
                    )
                )
                continue
            row_issues = _source_checks(row, expected, row_number=row_number)
            issues.extend(row_issues)

            review_status = _text(row[14]).upper()
            status_counts[review_status or "BLANK"] += 1
            if review_status not in REVIEW_STATUSES:
                issues.append(
                    ImportIssue(
                        code="INVALID_REVIEW_STATUS",
                        message="Review status is not an allowed value",
                        row=row_number,
                        candidate_id=candidate_id,
                        field="审核状态",
                    )
                )
                continue
            if review_status != "APPROVED":
                continue

            query = _text(row[9])
            source_type = _text(row[10]).upper()
            classification = _text(row[11]).upper()
            answer_terms = _split_values(row[12])
            confirmed_locators = _confirmed_locators(row[13], expected["source_locators"])
            reviewed_by = _text(row[15])
            approved_issues: list[ImportIssue] = []
            required = {
                "人工查询": query,
                "来源类型": source_type,
                "预期分类": classification,
                "答案关键术语": answer_terms,
                "确认引用定位": confirmed_locators,
                "审核人": reviewed_by,
                "审核时间": row[16],
            }
            for field, value in required.items():
                if value is None or value == "" or value == []:
                    approved_issues.append(
                        ImportIssue(
                            code="APPROVED_FIELD_MISSING",
                            message="Approved row is missing a required review field",
                            row=row_number,
                            candidate_id=candidate_id,
                            field=field,
                        )
                    )
            if query and len(query) < 2:
                approved_issues.append(
                    ImportIssue(
                        code="QUERY_TOO_SHORT",
                        message="Approved query must contain at least two characters",
                        row=row_number,
                        candidate_id=candidate_id,
                        field="人工查询",
                    )
                )
            if source_type and source_type not in SOURCE_TYPES:
                approved_issues.append(
                    ImportIssue(
                        code="INVALID_SOURCE_TYPE",
                        message="Approved row has an invalid source type",
                        row=row_number,
                        candidate_id=candidate_id,
                        field="来源类型",
                    )
                )
            if classification and classification not in CLASSIFICATIONS:
                approved_issues.append(
                    ImportIssue(
                        code="INVALID_CLASSIFICATION",
                        message="Approved row has an invalid classification",
                        row=row_number,
                        candidate_id=candidate_id,
                        field="预期分类",
                    )
                )
            unknown_locators = [
                locator
                for locator in confirmed_locators
                if locator not in expected["source_locators"]
            ]
            if unknown_locators:
                approved_issues.append(
                    ImportIssue(
                        code="CITATION_OUTSIDE_SOURCE",
                        message="Confirmed citation is not present in the candidate source locators",
                        row=row_number,
                        candidate_id=candidate_id,
                        field="确认引用定位",
                    )
                )
            try:
                reviewed_at = _reviewed_at(row[16], timezone_name)
            except (ValueError, TypeError, KeyError) as exc:
                reviewed_at = ""
                if row[16] not in (None, ""):
                    approved_issues.append(
                        ImportIssue(
                            code="INVALID_REVIEW_TIME",
                            message=f"Review time is invalid: {type(exc).__name__}",
                            row=row_number,
                            candidate_id=candidate_id,
                            field="审核时间",
                        )
                    )
            issues.extend(approved_issues)
            if row_issues or approved_issues:
                continue

            cases.append(
                {
                    "case_id": _candidate_case_id(candidate_id),
                    "query": query,
                    "scope": expected["scope"],
                    "project_id": expected["project_id"] or None,
                    "source_type": source_type,
                    "expected_classification": classification,
                    "expected_relevant_chunk_ids": expected["relevant_chunk_ids"],
                    "expected_answer_terms": answer_terms,
                    "expected_citations": [
                        {
                            "document_id": expected["document_id"],
                            "chunk_id": expected["chunk_id"],
                            "source_locator": locator,
                        }
                        for locator in confirmed_locators
                    ],
                    "review": {
                        "status": "APPROVED",
                        "reviewed_by": reviewed_by,
                        "reviewed_at": reviewed_at,
                    },
                }
            )

        missing_rows = sorted(set(candidate_by_id) - seen_ids)
        for candidate_id in missing_rows:
            issues.append(
                ImportIssue(
                    code="MISSING_CANDIDATE_ROW",
                    message="Source candidate is missing from the review workbook",
                    candidate_id=candidate_id,
                )
            )
        return ReviewImportResult(cases, issues, status_counts, row_count)
    finally:
        workbook.close()


def build_golden_dataset(
    result: ReviewImportResult,
    *,
    dataset_id: str,
    provider: str,
    model: str,
    dimension: int,
    index_version: str,
    schema: dict[str, Any],
) -> dict[str, Any]:
    if result.has_errors:
        raise ValueError("Review workbook contains validation errors")
    if not 100 <= len(result.cases) <= 200:
        raise ValueError("Approved case count must be between 100 and 200")
    dataset = {
        "schema_version": "poc-03.golden.v1",
        "dataset_id": dataset_id,
        "embedding_index": {
            "provider": provider,
            "model": model,
            "dimension": dimension,
            "index_version": index_version,
        },
        "cases": result.cases,
    }
    validator = Draft202012Validator(schema, format_checker=FormatChecker())
    errors = sorted(validator.iter_errors(dataset), key=lambda item: list(item.absolute_path))
    if errors:
        paths = ["/".join(str(part) for part in error.absolute_path) for error in errors]
        raise ValueError(f"Golden Dataset schema validation failed at: {paths}")
    return dataset


def build_sanitized_import_report(
    result: ReviewImportResult,
    *,
    generated_at: str,
    output_written: bool,
) -> dict[str, Any]:
    if result.has_errors:
        status = "FAIL"
    elif 100 <= len(result.cases) <= 200:
        status = "PASS" if output_written else "READY"
    else:
        status = "INCOMPLETE"
    return {
        "schema_version": "poc-03.review-import-result.v1",
        "generated_at": generated_at,
        "status": status,
        "privacy": {
            "workbook_content_committed": False,
            "queries_committed": False,
            "answer_terms_committed": False,
            "reviewer_names_committed": False,
            "output_dataset_committed": False,
        },
        "summary": {
            "workbook_candidate_row_count": result.row_count,
            "approved_case_count": len(result.cases),
            "review_status_counts": dict(sorted(result.status_counts.items())),
            "issue_count": len(result.issues),
            "output_written": output_written,
        },
        "issues": [issue.sanitized() for issue in result.issues],
        "conclusion": (
            "Formal Golden Dataset export is allowed."
            if status in {"READY", "PASS"}
            else "Formal Golden Dataset export remains blocked."
        ),
    }
