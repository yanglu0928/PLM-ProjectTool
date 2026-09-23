from __future__ import annotations

import copy
from collections import Counter
from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from .review_import import ImportIssue, _reviewed_at, _text


CLASSIFICATIONS = {
    "STANDARD_SATISFIED": "标准满足",
    "PARTIALLY_SATISFIED": "部分满足",
    "NON_STANDARD": "非标准",
    "INSUFFICIENT_INFORMATION": "资料不足",
    "NO_RELIABLE_MATCH": "无可靠匹配",
    "HUMAN_CONFIRMATION_REQUIRED": "需人工确认",
}
NAME_TO_CODE = {name: code for code, name in CLASSIFICATIONS.items()}
GLOBAL_CONFIRMATION = "确认使用AI建议批量规则"
GROUP_DECISIONS = {"", "采用本次AI", "沿用R4", "逐条确认", "暂不确认"}

RULE_HEADERS = [
    "分组编号",
    "R4分类",
    "本次AI分类",
    "条数",
    "AI建议处理",
    "人工批量规则（可覆盖）",
    "生效规则",
    "说明",
]
CONFLICT_HEADERS = [
    "待办编号",
    "资料类别",
    "风险",
    "AI分析问题",
    "R4分类",
    "本次AI分类",
    "分组",
    "打开证据",
    "单条最终分类",
    "例外说明",
    "生效最终分类",
    "状态",
]

NO_RELIABLE_MATCH_MARKERS = (
    "暂不作为独立需求项",
    "无法形成有效合同结论",
    "更像缩写、编号或OCR片段",
)
INSUFFICIENT_INFORMATION_MARKERS = (
    "证据片段不足",
    "摘录不足",
    "不足以完整判断",
    "单凭该片段无法",
    "尚不能",
)


@dataclass
class R5ImportResult:
    cases: list[dict[str, Any]]
    issues: list[ImportIssue]
    status_counts: Counter[str]
    confirmed: bool
    total_count: int
    explicit_count: int
    conflict_count: int

    @property
    def has_errors(self) -> bool:
        return bool(self.issues)

    @property
    def ready(self) -> bool:
        return (
            self.confirmed
            and not self.has_errors
            and len(self.cases) == self.total_count
            and self.status_counts["PENDING"] == 0
        )


def parse_explicit_classification(note: Any) -> str | None:
    text = _text(note)
    if any(marker in text for marker in NO_RELIABLE_MATCH_MARKERS):
        return "NO_RELIABLE_MATCH"
    for code, name in CLASSIFICATIONS.items():
        if f"“{name}”" in text or f"按“{name}”" in text or f"按{name}处理" in text:
            return code
    if any(marker in text for marker in INSUFFICIENT_INFORMATION_MARKERS):
        return "INSUFFICIENT_INFORMATION"
    return None


def _prediction_map(quality: dict[str, Any]) -> dict[str, str]:
    predictions = quality.get("predictions") or {}
    return {
        _text(case_id): _text(item.get("classification")).upper()
        for case_id, item in predictions.items()
    }


def _load_r4_rows(path: Path) -> list[tuple[Any, ...]]:
    workbook = load_workbook(path, read_only=True, data_only=False)
    try:
        if "确认清单" not in workbook.sheetnames:
            raise ValueError("R4 workbook is missing 确认清单")
        return list(
            workbook["确认清单"].iter_rows(
                min_row=9, max_col=11, values_only=True
            )
        )
    finally:
        workbook.close()


def _locked_records(
    r4_workbook_path: Path,
    golden: dict[str, Any],
    quality: dict[str, Any],
) -> tuple[list[dict[str, Any]], list[ImportIssue]]:
    issues: list[ImportIssue] = []
    cases = list(golden.get("cases") or [])
    r4_rows = [row for row in _load_r4_rows(r4_workbook_path) if _text(row[0])]
    predictions = _prediction_map(quality)
    if len(cases) != len(r4_rows):
        issues.append(
            ImportIssue(
                code="LOCKED_INPUT_COUNT_MISMATCH",
                message="R4 workbook and Golden Dataset row counts differ",
            )
        )
        return [], issues

    records: list[dict[str, Any]] = []
    for index, (case, row) in enumerate(zip(cases, r4_rows, strict=True), start=1):
        case_id = _text(case.get("case_id"))
        task_id = _text(row[0])
        expected = _text(case.get("expected_classification")).upper()
        ai_code = predictions.get(case_id, "")
        if task_id[-4:] != case_id[-4:]:
            issues.append(
                ImportIssue(
                    code="LOCKED_INPUT_ALIGNMENT_ERROR",
                    message="R4 task and Golden Dataset case are not aligned",
                    row=index + 8,
                )
            )
        if expected not in CLASSIFICATIONS or ai_code not in CLASSIFICATIONS:
            issues.append(
                ImportIssue(
                    code="INVALID_LOCKED_CLASSIFICATION",
                    message="Locked input contains an unsupported classification",
                    row=index + 8,
                )
            )
        explicit = parse_explicit_classification(row[7])
        if explicit and explicit != expected:
            issues.append(
                ImportIssue(
                    code="EXPLICIT_CONCLUSION_MISMATCH",
                    message="R4 explicit conclusion differs from its exported label",
                    row=index + 8,
                )
            )
        records.append(
            {
                "case": case,
                "case_id": case_id,
                "task_id": task_id,
                "source_type": _text(case.get("source_type")),
                "risk": _text(row[2]),
                "query": _text(case.get("query")),
                "r4_code": expected,
                "ai_code": ai_code,
                "explicit_code": explicit,
            }
        )
    return records, issues


def _assign_groups(records: list[dict[str, Any]]) -> tuple[list[dict[str, Any]], dict[str, str]]:
    grouped: dict[tuple[str, str], list[dict[str, Any]]] = {}
    for record in records:
        if record["explicit_code"]:
            continue
        grouped.setdefault((record["r4_code"], record["ai_code"]), []).append(record)
    groups: list[dict[str, Any]] = []
    group_by_case: dict[str, str] = {}
    for index, ((r4_code, ai_code), group_records) in enumerate(sorted(grouped.items()), start=1):
        group_id = f"GRP-{index:02d}"
        recommendation = "采用本次AI" if r4_code == "HUMAN_CONFIRMATION_REQUIRED" else "沿用R4"
        groups.append(
            {
                "group_id": group_id,
                "r4_code": r4_code,
                "ai_code": ai_code,
                "count": len(group_records),
                "recommendation": recommendation,
            }
        )
        for record in group_records:
            group_by_case[record["case_id"]] = group_id
    return groups, group_by_case


def _changed_issue(row: int, field: str) -> ImportIssue:
    return ImportIssue(
        code="SOURCE_FIELD_CHANGED",
        message="Immutable R5 review source field was changed",
        row=row,
        field=field,
    )


def import_r5_label_review(
    workbook_path: Path,
    r4_workbook_path: Path,
    golden: dict[str, Any],
    quality: dict[str, Any],
    *,
    timezone_name: str = "Asia/Shanghai",
) -> R5ImportResult:
    records, issues = _locked_records(r4_workbook_path, golden, quality)
    groups, group_by_case = _assign_groups(records)
    explicit_count = sum(bool(record["explicit_code"]) for record in records)
    conflict_records = [record for record in records if not record["explicit_code"]]
    statuses: Counter[str] = Counter()
    final_by_case: dict[str, str] = {
        record["case_id"]: record["r4_code"]
        for record in records
        if record["explicit_code"]
    }

    workbook = load_workbook(workbook_path, read_only=True, data_only=False)
    try:
        required_sheets = {"批量规则", "冲突确认", "明确结论", "技术底稿"}
        missing = required_sheets - set(workbook.sheetnames)
        if missing:
            issues.append(
                ImportIssue(
                    code="MISSING_SHEET",
                    message=f"R5 workbook is missing required sheets: {sorted(missing)}",
                )
            )
            return R5ImportResult([], issues, statuses, False, len(records), explicit_count, len(conflict_records))

        rules_sheet = workbook["批量规则"]
        conflicts_sheet = workbook["冲突确认"]
        if [_text(rules_sheet.cell(9, col).value) for col in range(1, 9)] != RULE_HEADERS:
            issues.append(ImportIssue("HEADER_MISMATCH", "Batch-rule headers differ from R5 template", row=9))
        if [_text(conflicts_sheet.cell(8, col).value) for col in range(1, 13)] != CONFLICT_HEADERS:
            issues.append(ImportIssue("HEADER_MISMATCH", "Conflict headers differ from R5 template", row=8))

        global_value = _text(rules_sheet["B5"].value)
        confirmed = global_value == GLOBAL_CONFIRMATION
        if global_value not in {"未确认", GLOBAL_CONFIRMATION}:
            issues.append(ImportIssue("INVALID_GLOBAL_CONFIRMATION", "Batch confirmation has an invalid value", row=5, field="批量确认"))
        reviewer = _text(rules_sheet["D5"].value)
        reviewed_value = rules_sheet["F5"].value
        reviewed_at = ""
        if confirmed:
            if not reviewer:
                issues.append(ImportIssue("REVIEWER_REQUIRED", "Confirmed R5 workbook requires a reviewer", row=5, field="确认人"))
            try:
                reviewed_at = _reviewed_at(reviewed_value, timezone_name)
            except (ValueError, TypeError, KeyError):
                issues.append(ImportIssue("INVALID_REVIEW_TIME", "Confirmed R5 workbook requires a valid review date", row=5, field="确认日期"))

        group_decisions: dict[str, str] = {}
        for offset, group in enumerate(groups, start=10):
            row = [rules_sheet.cell(offset, col).value for col in range(1, 9)]
            locked = [
                group["group_id"],
                CLASSIFICATIONS[group["r4_code"]],
                CLASSIFICATIONS[group["ai_code"]],
                group["count"],
                group["recommendation"],
            ]
            for col, (actual, expected) in enumerate(zip(row[:5], locked, strict=True), start=1):
                if _text(actual) != _text(expected):
                    issues.append(_changed_issue(offset, RULE_HEADERS[col - 1]))
            override = _text(row[5])
            if override not in GROUP_DECISIONS:
                issues.append(ImportIssue("INVALID_GROUP_DECISION", "Group override has an invalid value", row=offset, field=RULE_HEADERS[5]))
                override = "暂不确认"
            group_decisions[group["group_id"]] = override or (group["recommendation"] if confirmed else "待确认")

        if len(conflict_records) != sum(1 for row in conflicts_sheet.iter_rows(min_row=9, max_col=1, values_only=True) if _text(row[0])):
            issues.append(ImportIssue("CONFLICT_ROW_COUNT_MISMATCH", "Conflict rows differ from the locked R5 scope"))

        seen: set[str] = set()
        by_task = {record["task_id"]: record for record in conflict_records}
        for row_number, row in enumerate(conflicts_sheet.iter_rows(min_row=9, max_col=12, values_only=True), start=9):
            task_id = _text(row[0])
            if not task_id:
                continue
            record = by_task.get(task_id)
            if record is None or task_id in seen:
                issues.append(ImportIssue("UNKNOWN_OR_DUPLICATE_TASK", "Conflict row task is unknown or duplicated", row=row_number))
                continue
            seen.add(task_id)
            group_id = group_by_case[record["case_id"]]
            locked = {
                1: record["source_type"],
                2: record["risk"],
                3: record["query"],
                4: CLASSIFICATIONS[record["r4_code"]],
                5: CLASSIFICATIONS[record["ai_code"]],
                6: group_id,
            }
            for column, expected in locked.items():
                if _text(row[column]) != _text(expected):
                    issues.append(_changed_issue(row_number, CONFLICT_HEADERS[column]))

            manual_name = _text(row[8])
            if manual_name:
                if manual_name not in NAME_TO_CODE:
                    issues.append(ImportIssue("INVALID_ROW_CLASSIFICATION", "Row override has an invalid classification", row=row_number, field=CONFLICT_HEADERS[8]))
                    statuses["INVALID"] += 1
                    continue
                final_code = NAME_TO_CODE[manual_name]
            else:
                decision = group_decisions.get(group_id, "待确认")
                if decision == "采用本次AI":
                    final_code = record["ai_code"]
                elif decision == "沿用R4":
                    final_code = record["r4_code"]
                else:
                    statuses["PENDING"] += 1
                    continue
            final_by_case[record["case_id"]] = final_code
            statuses["APPROVED"] += 1

        for missing_task in sorted(set(by_task) - seen):
            issues.append(ImportIssue("MISSING_TASK_ROW", "Locked R5 conflict task is missing", candidate_id=missing_task))

        if not confirmed:
            statuses["PENDING"] = max(statuses["PENDING"], len(conflict_records))
            return R5ImportResult([], issues, statuses, False, len(records), explicit_count, len(conflict_records))

        cases: list[dict[str, Any]] = []
        for record in records:
            final_code = final_by_case.get(record["case_id"])
            if not final_code:
                continue
            case = copy.deepcopy(record["case"])
            case["expected_classification"] = final_code
            case["review"] = {
                "status": "APPROVED",
                "reviewed_by": reviewer,
                "reviewed_at": reviewed_at,
            }
            cases.append(case)
        statuses["EXPLICIT_R4"] = explicit_count
        return R5ImportResult(cases, issues, statuses, True, len(records), explicit_count, len(conflict_records))
    finally:
        workbook.close()


def build_r5_dataset(
    golden: dict[str, Any],
    result: R5ImportResult,
    *,
    dataset_id: str,
) -> dict[str, Any]:
    if not result.ready:
        raise ValueError("R5 review is not ready for Golden Dataset export")
    dataset = copy.deepcopy(golden)
    dataset["dataset_id"] = dataset_id
    dataset["cases"] = result.cases
    return dataset


def build_sanitized_r5_report(
    result: R5ImportResult,
    *,
    generated_at: str,
    output_written: bool,
) -> dict[str, Any]:
    if result.has_errors:
        status = "FAIL"
    elif result.ready:
        status = "PASS" if output_written else "READY"
    else:
        status = "AWAITING_HUMAN_CONFIRMATION"
    return {
        "schema_version": "poc-03.r5-label-import-result.v1",
        "generated_at": generated_at,
        "status": status,
        "privacy": {
            "workbook_content_committed": False,
            "queries_committed": False,
            "reviewer_names_committed": False,
            "output_dataset_committed": False,
        },
        "summary": {
            "total_case_count": result.total_count,
            "r4_explicit_conclusion_count": result.explicit_count,
            "conflict_review_count": result.conflict_count,
            "resolved_case_count": len(result.cases),
            "batch_confirmation_received": result.confirmed,
            "pending_count": result.status_counts["PENDING"],
            "issue_count": len(result.issues),
            "output_written": output_written,
        },
        "issues": [issue.sanitized() for issue in result.issues],
        "conclusion": (
            "R5 Golden Dataset export is allowed."
            if status in {"READY", "PASS"}
            else "R5 Golden Dataset export remains blocked."
        ),
    }
