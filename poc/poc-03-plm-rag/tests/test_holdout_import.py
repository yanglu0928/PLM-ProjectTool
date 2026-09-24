from __future__ import annotations

import copy
import json
import sys
import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook, load_workbook


POC_DIR = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(POC_DIR / "src"))

from poc03_rag.holdout_import import (  # noqa: E402
    AUDIT_HEADERS,
    CLASSIFICATION_NAMES,
    GLOBAL_CONFIRMATION,
    REVIEW_HEADERS,
    audit_holdout_dataset,
    build_holdout_dataset,
    import_holdout_review,
)


SOURCE_TYPE_NAMES = {
    "STANDARD_CAPABILITY": "标准能力",
    "CONTRACT": "合同",
    "TECHNICAL_AGREEMENT": "技术协议",
    "SURVEY": "调研记录",
}
SOURCE_TYPES = (
    ["STANDARD_CAPABILITY"] * 33
    + ["CONTRACT"] * 7
    + ["TECHNICAL_AGREEMENT"] * 8
    + ["SURVEY"] * 2
)
CLASSIFICATIONS = [
    "STANDARD_SATISFIED",
    "PARTIALLY_SATISFIED",
    "NON_STANDARD",
    "INSUFFICIENT_INFORMATION",
    "NO_RELIABLE_MATCH",
]


def make_inputs() -> tuple[dict, dict]:
    package_cases = []
    lock_candidates = []
    lock_fingerprint = "a" * 64
    for index, source_type in enumerate(SOURCE_TYPES, start=1):
        candidate_id = f"HO-C-{index:04d}"
        task_id = f"HO-R-{index:04d}"
        classification = CLASSIFICATIONS[(index - 1) % len(CLASSIFICATIONS)]
        locator = f"word/paragraph/{index}"
        chunk_id = f"DOC-{index:04d}-C-{index:04d}"
        sha256 = f"{index:064x}"
        lock_candidates.append(
            {
                "candidate_id": candidate_id,
                "status": "LOCKED_PENDING_HUMAN_REVIEW",
                "scope": "PROJECT",
                "project_id": "PROJECT-HOLDOUT",
                "source_type": source_type,
                "evidence_role": "ACTUAL_RECORD",
                "document_id": f"DOC-{index:04d}",
                "chunk_id": chunk_id,
                "candidate_content": f"锁定候选内容 {index}",
                "text_sha256": sha256,
                "source_locators": [locator],
                "pages": [],
                "sections": [],
            }
        )
        package_cases.append(
            {
                "task_id": task_id,
                "candidate_id": candidate_id,
                "source_type": source_type,
                "source_type_display": SOURCE_TYPE_NAMES[source_type],
                "evidence_role": "ACTUAL_RECORD",
                "suggested_question": f"请判断事项 {index} 的满足程度、边界条件和证据是否充分。",
                "suggested_classification": classification,
                "suggested_classification_display": CLASSIFICATION_NAMES[classification],
                "classification_reason": "根据锁定证据形成的待确认建议。",
                "answer_terms": [f"术语{index}"],
                "evidence_quote": f"证据摘录 {index}",
                "evidence_quote_repaired": False,
                "answer_terms_repaired": False,
                "chunk_id": chunk_id,
                "text_sha256": sha256,
                "document_id": f"DOC-{index:04d}",
                "source_locators": [locator],
                "pages": [],
                "sections": [],
                "location_label": f"段落 {index}",
                "source_name": f"source-{index}.docx",
                "original_url": "",
                "evidence_link": f"evidence.html#case-{index}",
                "candidate_content": f"锁定候选内容 {index}",
            }
        )
    lock = {
        "schema_version": "poc-03.holdout-lock.v1",
        "status": "LOCKED_PENDING_HUMAN_REVIEW",
        "lock_id": "poc-03-holdout-test",
        "lock_fingerprint": lock_fingerprint,
        "project_id": "PROJECT-HOLDOUT",
        "target_count": 50,
        "source_type_quotas": {
            "STANDARD_CAPABILITY": 33,
            "CONTRACT": 7,
            "TECHNICAL_AGREEMENT": 8,
            "SURVEY": 2,
        },
        "isolation": {
            "case_and_query_disjoint": True,
            "chunk_id_disjoint": True,
            "source_locator_disjoint": True,
            "document_disjoint": False,
            "document_disjoint_reason": "Legacy sources use unseen chunks and locators.",
        },
        "candidates": lock_candidates,
    }
    package = {
        "schema_version": "poc-03.holdout-review-package.v1",
        "status": "AWAITING_HUMAN_CONFIRMATION",
        "lock_fingerprint": lock_fingerprint,
        "scope_count": 50,
        "cases": package_cases,
    }
    return package, lock


def write_workbook(
    path: Path,
    package: dict,
    *,
    confirmed: bool,
    modified: bool = False,
) -> None:
    workbook = Workbook()
    review = workbook.active
    review.title = "确认清单"
    workbook.create_sheet("分类说明")
    audit = workbook.create_sheet("技术底稿")
    review["B5"] = GLOBAL_CONFIRMATION if confirmed else "待确认"
    review["D5"] = "审核人" if confirmed else ""
    review["F5"] = "2026-09-21" if confirmed else ""
    for column, header in enumerate(REVIEW_HEADERS, start=1):
        review.cell(9, column, header)
    for row_number, item in enumerate(package["cases"], start=10):
        values = [
            item["task_id"],
            item["source_type_display"],
            item["suggested_question"],
            item["suggested_classification_display"],
            item["classification_reason"],
            item["location_label"],
            "打开证据",
            "",
            "",
            "",
            "",
            "已确认" if confirmed else "待确认",
        ]
        if modified and row_number == 10:
            values[7] = "修改后确认"
            values[8] = "人工确认后的最终问题，要求判断覆盖范围、满足程度和证据边界。"
            values[9] = "部分满足"
            values[10] = "人工核对后修正。"
        for column, value in enumerate(values, start=1):
            review.cell(row_number, column, value)
    for column, header in enumerate(AUDIT_HEADERS, start=1):
        audit.cell(6, column, header)
    for row_number, item in enumerate(package["cases"], start=7):
        values = [
            item["candidate_id"],
            item["task_id"],
            item["source_type"],
            item["evidence_role"],
            item["document_id"],
            item["chunk_id"],
            item["text_sha256"],
            "；".join(item["source_locators"]),
            "否",
            "否",
            item["source_name"],
            package["lock_fingerprint"],
        ]
        for column, value in enumerate(values, start=1):
            audit.cell(row_number, column, value)
    workbook.save(path)


class HoldoutImportTests(unittest.TestCase):
    def schema(self) -> dict:
        return json.loads(
            (POC_DIR / "schema" / "holdout-dataset.schema.json").read_text(
                encoding="utf-8"
            )
        )

    def test_global_confirmation_builds_schema_valid_holdout(self) -> None:
        package, lock = make_inputs()
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "holdout.xlsx"
            write_workbook(path, package, confirmed=True)
            result = import_holdout_review(path, package, lock)
        self.assertTrue(result.ready)
        dataset = build_holdout_dataset(
            result,
            lock,
            dataset_id="holdout-test",
            schema=self.schema(),
        )
        coverage = audit_holdout_dataset(dataset)
        self.assertEqual(len(dataset["cases"]), 50)
        self.assertEqual(coverage["status"], "PASS")
        self.assertEqual(coverage["summary"]["unique_query_count"], 50)

    def test_manual_override_takes_precedence(self) -> None:
        package, lock = make_inputs()
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "holdout.xlsx"
            write_workbook(path, package, confirmed=True, modified=True)
            result = import_holdout_review(path, package, lock)
        self.assertTrue(result.ready)
        self.assertEqual(result.modified_count, 1)
        self.assertEqual(
            result.cases[0]["expected_classification"], "PARTIALLY_SATISFIED"
        )
        self.assertTrue(result.cases[0]["query"].startswith("人工确认后"))

    def test_locked_source_change_is_rejected(self) -> None:
        package, lock = make_inputs()
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "holdout.xlsx"
            write_workbook(path, package, confirmed=True)
            workbook = load_workbook(path)
            workbook["确认清单"]["C10"] = "被篡改的问题"
            workbook.save(path)
            result = import_holdout_review(path, package, lock)
        self.assertTrue(result.has_errors)
        self.assertIn("SOURCE_FIELD_CHANGED", {issue.code for issue in result.issues})

    def test_missing_reviewer_blocks_export(self) -> None:
        package, lock = make_inputs()
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "holdout.xlsx"
            write_workbook(path, package, confirmed=True)
            workbook = load_workbook(path)
            workbook["确认清单"]["D5"] = ""
            workbook.save(path)
            result = import_holdout_review(path, package, lock)
        self.assertFalse(result.ready)
        self.assertIn("REVIEWER_REQUIRED", {issue.code for issue in result.issues})

    def test_lock_fingerprint_mismatch_is_rejected(self) -> None:
        package, lock = make_inputs()
        changed_lock = copy.deepcopy(lock)
        changed_lock["lock_fingerprint"] = "b" * 64
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "holdout.xlsx"
            write_workbook(path, package, confirmed=True)
            result = import_holdout_review(path, package, changed_lock)
        self.assertTrue(result.has_errors)
        self.assertIn("INVALID_HOLDOUT_SCOPE", {issue.code for issue in result.issues})


if __name__ == "__main__":
    unittest.main()
