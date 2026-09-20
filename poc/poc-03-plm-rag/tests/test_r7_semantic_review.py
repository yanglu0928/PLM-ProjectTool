from __future__ import annotations

import copy
import sys
import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook, load_workbook


POC_DIR = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(POC_DIR / "src"))

from poc03_rag.r7_semantic_review import (  # noqa: E402
    GLOBAL_CONFIRMATION,
    HEADERS,
    build_r7_dataset,
    import_r7_semantic_review,
)


def make_inputs() -> tuple[dict, dict]:
    cases = []
    package_cases = []
    for index in range(1, 121):
        case_id = f"GD-{index:04d}"
        current_chunk = f"DOC-C-{index:04d}"
        suggested_chunk = f"ALT-C-{index:04d}"
        cases.append(
            {
                "case_id": case_id,
                "query": f"原问题 {index}",
                "scope": "PROJECT",
                "project_id": "PROJECT-A",
                "source_type": "STANDARD_CAPABILITY",
                "expected_classification": "INSUFFICIENT_INFORMATION",
                "expected_relevant_chunk_ids": [current_chunk],
                "expected_answer_terms": ["术语"],
                "expected_citations": [
                    {
                        "document_id": "DOC",
                        "chunk_id": current_chunk,
                        "source_locator": f"word/paragraph/{index}",
                    }
                ],
                "review": {
                    "status": "APPROVED",
                    "reviewed_by": "旧审核人",
                    "reviewed_at": "2026-09-18T00:00:00Z",
                },
            }
        )
        package_cases.append(
            {
                "task_id": f"R7-{index:04d}",
                "case_id": case_id,
                "source_type": "STANDARD_CAPABILITY",
                "current_classification": "INSUFFICIENT_INFORMATION",
                "prompt_v2_classification": "STANDARD_SATISFIED",
                "issue_reason": "分类与引用均需复核",
                "suggested_target": f"围绕事项 {index}，判断标准能力满足、部分满足、非标准、资料不足或无可靠匹配，并确认引用。",
                "suggested_classification": "STANDARD_SATISFIED",
                "classification_rationale": "证据被模型判断为可直接证明标准满足，仍需人工确认。",
                "suggested_chunk_ids": [suggested_chunk],
                "current_chunk_ids": [current_chunk],
                "allowed_chunk_ids": [suggested_chunk, current_chunk],
                "candidate_chunks": [
                    {
                        "chunk_id": suggested_chunk,
                        "document_id": "ALT",
                        "source_locators": [f"word/paragraph/{index + 1000}"],
                    },
                    {
                        "chunk_id": current_chunk,
                        "document_id": "DOC",
                        "source_locators": [f"word/paragraph/{index}"],
                    },
                ],
            }
        )
    golden = {"dataset_id": "r6", "cases": cases}
    package = {
        "source_dataset_id": "r6",
        "scope_case_ids": [case["case_id"] for case in cases],
        "cases": package_cases,
    }
    return golden, package


def write_workbook(
    path: Path,
    package: dict,
    *,
    confirmed: bool,
    modified: bool = False,
) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "R7确认"
    sheet["A5"] = "批量确认"
    sheet["B5"] = GLOBAL_CONFIRMATION if confirmed else "未确认"
    sheet["D5"] = "审核人" if confirmed else ""
    sheet["F5"] = "2026-09-20" if confirmed else ""
    for col, header in enumerate(HEADERS, start=1):
        sheet.cell(9, col, header)
    for row_number, item in enumerate(package["cases"], start=10):
        values = [
            item["task_id"],
            item["case_id"],
            item["source_type"],
            item["current_classification"],
            item["prompt_v2_classification"],
            item["issue_reason"],
            item["suggested_target"],
            item["suggested_classification"],
            item["classification_rationale"],
            "；".join(item["suggested_chunk_ids"]),
            "打开证据",
            "",
            "",
            "",
            "",
            "",
            "待确认",
        ]
        if modified and row_number == 10:
            values[11] = "修改后确认"
            values[12] = "人工补充了完整的判定目标，需要判断满足程度并确认功能范围、条件、限制和引用。"
            values[13] = "PARTIALLY_SATISFIED"
            values[14] = item["current_chunk_ids"][0]
            values[15] = "人工核对后认为仅部分覆盖。"
        for col, value in enumerate(values, start=1):
            sheet.cell(row_number, col, value)
    workbook.save(path)


class R7SemanticReviewTests(unittest.TestCase):
    def test_unconfirmed_workbook_stays_pending(self) -> None:
        golden, package = make_inputs()
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "r7.xlsx"
            write_workbook(path, package, confirmed=False)
            result = import_r7_semantic_review(path, golden, package)
        self.assertFalse(result.ready)
        self.assertEqual(result.status_counts["PENDING"], 120)
        self.assertEqual(result.cases, [])

    def test_global_confirmation_builds_calibration_dataset(self) -> None:
        golden, package = make_inputs()
        original = copy.deepcopy(golden)
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "r7.xlsx"
            write_workbook(path, package, confirmed=True)
            result = import_r7_semantic_review(path, golden, package)
        self.assertTrue(result.ready)
        dataset = build_r7_dataset(golden, result, dataset_id="r7")
        self.assertEqual(len(dataset["cases"]), 120)
        self.assertEqual(result.query_change_count, 120)
        self.assertEqual(result.classification_change_count, 120)
        self.assertEqual(result.citation_change_count, 120)
        self.assertEqual(original["dataset_id"], "r6")
        self.assertEqual(dataset["dataset_id"], "r7")

    def test_manual_override_takes_precedence_over_global_confirmation(self) -> None:
        golden, package = make_inputs()
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "r7.xlsx"
            write_workbook(path, package, confirmed=True, modified=True)
            result = import_r7_semantic_review(path, golden, package)
        self.assertTrue(result.ready)
        changed = result.cases[0]
        self.assertEqual(changed["expected_classification"], "PARTIALLY_SATISFIED")
        self.assertEqual(changed["expected_relevant_chunk_ids"], ["DOC-C-0001"])

    def test_locked_source_change_is_rejected(self) -> None:
        golden, package = make_inputs()
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "r7.xlsx"
            write_workbook(path, package, confirmed=True)
            workbook = load_workbook(path)
            workbook["R7确认"]["G10"] = "被篡改的AI建议"
            workbook.save(path)
            result = import_r7_semantic_review(path, golden, package)
        self.assertTrue(result.has_errors)
        self.assertIn("SOURCE_FIELD_CHANGED", {issue.code for issue in result.issues})

    def test_manual_citation_must_be_displayed_evidence(self) -> None:
        golden, package = make_inputs()
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "r7.xlsx"
            write_workbook(path, package, confirmed=True, modified=True)
            workbook = load_workbook(path)
            workbook["R7确认"]["O10"] = "OUTSIDE-C-9999"
            workbook.save(path)
            result = import_r7_semantic_review(path, golden, package)
        self.assertTrue(result.has_errors)
        self.assertIn("CITATION_OUTSIDE_R7_EVIDENCE", {issue.code for issue in result.issues})


if __name__ == "__main__":
    unittest.main()
