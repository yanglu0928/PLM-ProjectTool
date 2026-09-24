from __future__ import annotations

import copy
import sys
import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook


POC_DIR = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(POC_DIR / "src"))

from poc03_rag.r6_citation_review import (  # noqa: E402
    GLOBAL_CONFIRMATION,
    HEADERS,
    build_r6_dataset,
    import_r6_citation_review,
)


def make_inputs() -> tuple[dict, dict]:
    cases = []
    package_cases = []
    for index in range(1, 121):
        case_id = f"GD-{index:04d}"
        cases.append({
            "case_id": case_id,
            "query": f"原问题 {index}",
            "scope": "PROJECT",
            "project_id": "PROJECT-A",
            "source_type": "STANDARD_CAPABILITY",
            "expected_classification": "STANDARD_SATISFIED",
            "expected_relevant_chunk_ids": [f"DOC-C-{index:04d}"],
            "expected_answer_terms": ["术语"],
            "expected_citations": [{"document_id": "DOC", "chunk_id": f"DOC-C-{index:04d}", "source_locator": f"word/paragraph/{index}"}],
            "review": {"status": "APPROVED", "reviewed_by": "旧审核人", "reviewed_at": "2026-09-18T00:00:00Z"},
        })
    scope = ["GD-0016", "GD-0018", "GD-0020", "GD-0023", "GD-0075", "GD-0091"]
    for order, case_id in enumerate(scope, start=1):
        index = int(case_id[-4:])
        package_cases.append({
            "task_id": f"R6-{order:02d}",
            "case_id": case_id,
            "source_type": "STANDARD_CAPABILITY",
            "classification": "STANDARD_SATISFIED",
            "original_query": f"原问题 {index}",
            "issue_reason": "问题过泛",
            "suggested_query": f"明确问题 {index}",
            "suggested_chunk_ids": [f"DOC-C-{index:04d}"],
            "allowed_chunk_ids": [f"DOC-C-{index:04d}", f"ALT-C-{index:04d}"],
            "candidate_chunks": [
                {"chunk_id": f"DOC-C-{index:04d}", "document_id": "DOC", "source_locators": [f"word/paragraph/{index}"]},
                {"chunk_id": f"ALT-C-{index:04d}", "document_id": "ALT", "source_locators": [f"word/paragraph/{index + 1}"]},
            ],
        })
    return {"dataset_id": "r5", "cases": cases}, {"scope_case_ids": scope, "cases": package_cases}


def write_workbook(path: Path, package: dict, *, confirmed: bool, modified: bool = False) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "R6确认"
    sheet["A5"] = "批量确认"
    sheet["B5"] = GLOBAL_CONFIRMATION if confirmed else "未确认"
    sheet["D5"] = "审核人" if confirmed else ""
    sheet["F5"] = "2026-09-18" if confirmed else ""
    for col, header in enumerate(HEADERS, start=1):
        sheet.cell(9, col, header)
    for row_number, item in enumerate(package["cases"], start=10):
        values = [
            item["task_id"], item["case_id"], item["source_type"], item["classification"], item["original_query"],
            item["issue_reason"], item["suggested_query"], "；".join(item["suggested_chunk_ids"]), "打开证据", "", "", "", "待确认",
        ]
        if modified and row_number == 10:
            values[9] = "修改后确认"
            values[10] = "人工明确问题"
            values[11] = item["allowed_chunk_ids"][1]
        for col, value in enumerate(values, start=1):
            sheet.cell(row_number, col, value)
    workbook.save(path)


class R6CitationReviewTests(unittest.TestCase):
    def test_unconfirmed_workbook_stays_pending(self) -> None:
        golden, package = make_inputs()
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "r6.xlsx"
            write_workbook(path, package, confirmed=False)
            result = import_r6_citation_review(path, golden, package)
        self.assertFalse(result.ready)
        self.assertEqual(result.status_counts["PENDING"], 6)
        self.assertEqual(result.cases, [])

    def test_global_confirmation_preserves_114_cases(self) -> None:
        golden, package = make_inputs()
        original = copy.deepcopy(golden)
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "r6.xlsx"
            write_workbook(path, package, confirmed=True)
            result = import_r6_citation_review(path, golden, package)
        self.assertTrue(result.ready)
        dataset = build_r6_dataset(golden, result, dataset_id="r6")
        self.assertEqual(len(dataset["cases"]), 120)
        scope = set(package["scope_case_ids"])
        for before, after in zip(original["cases"], dataset["cases"], strict=True):
            if before["case_id"] not in scope:
                self.assertEqual(before, after)
            else:
                self.assertEqual(before["expected_classification"], after["expected_classification"])
                self.assertNotEqual(before["query"], after["query"])

    def test_manual_override_must_use_reviewed_chunk(self) -> None:
        golden, package = make_inputs()
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "r6.xlsx"
            write_workbook(path, package, confirmed=True, modified=True)
            result = import_r6_citation_review(path, golden, package)
        self.assertTrue(result.ready)
        changed = next(case for case in result.cases if case["case_id"] == "GD-0016")
        self.assertEqual(changed["query"], "人工明确问题")
        self.assertEqual(changed["expected_relevant_chunk_ids"], ["ALT-C-0016"])

    def test_locked_field_change_is_rejected(self) -> None:
        golden, package = make_inputs()
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "r6.xlsx"
            write_workbook(path, package, confirmed=True)
            from openpyxl import load_workbook
            workbook = load_workbook(path)
            workbook["R6确认"]["E10"] = "被篡改"
            workbook.save(path)
            result = import_r6_citation_review(path, golden, package)
        self.assertTrue(result.has_errors)
        self.assertIn("SOURCE_FIELD_CHANGED", {issue.code for issue in result.issues})


if __name__ == "__main__":
    unittest.main()
