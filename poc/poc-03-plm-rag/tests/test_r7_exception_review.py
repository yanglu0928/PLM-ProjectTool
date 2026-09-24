from __future__ import annotations

import copy
import sys
import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook, load_workbook


POC_DIR = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(POC_DIR / "src"))

from poc03_rag.r7_exception_review import (  # noqa: E402
    CONFIRM_DECISION,
    HEADERS,
    import_r7_exception_review,
)


def make_inputs() -> tuple[dict, dict, dict]:
    case = {
        "case_id": "GD-0001",
        "query": "R7 query",
        "scope": "PROJECT",
        "project_id": "PROJECT-A",
        "source_type": "CONTRACT",
        "expected_classification": "INSUFFICIENT_INFORMATION",
        "expected_relevant_chunk_ids": ["DOC-C-0001"],
        "expected_answer_terms": ["旧术语"],
        "expected_citations": [
            {"document_id": "DOC", "chunk_id": "DOC-C-0001", "source_locator": "word/paragraph/1"}
        ],
        "review": {"status": "APPROVED", "reviewed_by": "旧审核人", "reviewed_at": "2026-09-20T00:00:00Z"},
    }
    dataset = {"dataset_id": "r7", "cases": [case]}
    package = {
        "source_dataset_id": "r7",
        "case": {
            "case_id": "GD-0001",
            "current_query": "R7 query",
            "current_classification": "资料不足（INSUFFICIENT_INFORMATION）",
            "current_classification_code": "INSUFFICIENT_INFORMATION",
            "proposed_query": "合同明确要求交付二次开发源代码；该事项是否属于非标准开发？",
            "proposed_classification": "非标准（NON_STANDARD）",
            "proposed_classification_code": "NON_STANDARD",
            "proposed_answer_terms": ["二次开发", "源代码", "交付"],
            "rationale": "存在直接二次开发证据。",
            "chunk_id": "DOC-C-0001",
        },
    }
    schema = {
        "type": "object",
        "required": ["dataset_id", "cases"],
        "properties": {
            "dataset_id": {"const": "r7-1"},
            "cases": {"type": "array", "minItems": 1},
        },
    }
    return dataset, package, schema


def write_workbook(path: Path, package: dict, *, confirmed: bool = True) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "R7.1确认"
    sheet["B5"] = CONFIRM_DECISION if confirmed else "未确认"
    sheet["D5"] = "审核人" if confirmed else ""
    sheet["F5"] = "2026-09-20" if confirmed else ""
    for column, header in enumerate(HEADERS, start=1):
        sheet.cell(8, column, header)
    item = package["case"]
    values = [
        item["case_id"],
        item["current_query"],
        item["current_classification"],
        item["proposed_query"],
        item["proposed_classification"],
        item["rationale"],
        item["chunk_id"],
        "打开证据",
        "",
        "已确认" if confirmed else "待确认",
    ]
    for column, value in enumerate(values, start=1):
        sheet.cell(9, column, value)
    workbook.save(path)


class R7ExceptionReviewTests(unittest.TestCase):
    def test_confirmed_exception_builds_new_dataset_without_mutating_r7(self) -> None:
        dataset, package, schema = make_inputs()
        original = copy.deepcopy(dataset)
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "r7-1.xlsx"
            write_workbook(path, package)
            result = import_r7_exception_review(path, dataset, package, dataset_id="r7-1", schema=schema)
        self.assertTrue(result.ready)
        self.assertEqual(result.dataset["dataset_id"], "r7-1")
        self.assertEqual(result.dataset["cases"][0]["expected_classification"], "NON_STANDARD")
        self.assertEqual(result.dataset["cases"][0]["expected_answer_terms"], ["二次开发", "源代码", "交付"])
        self.assertEqual(dataset, original)

    def test_unconfirmed_exception_does_not_export(self) -> None:
        dataset, package, schema = make_inputs()
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "r7-1.xlsx"
            write_workbook(path, package, confirmed=False)
            result = import_r7_exception_review(path, dataset, package, dataset_id="r7-1", schema=schema)
        self.assertFalse(result.ready)
        self.assertIsNone(result.dataset)

    def test_changed_locked_proposal_is_rejected(self) -> None:
        dataset, package, schema = make_inputs()
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "r7-1.xlsx"
            write_workbook(path, package)
            workbook = load_workbook(path)
            workbook["R7.1确认"]["D9"] = "被篡改的建议"
            workbook.save(path)
            result = import_r7_exception_review(path, dataset, package, dataset_id="r7-1", schema=schema)
        self.assertFalse(result.ready)
        self.assertIn("SOURCE_FIELD_CHANGED", {issue.code for issue in result.issues})

    def test_missing_reviewer_is_rejected(self) -> None:
        dataset, package, schema = make_inputs()
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "r7-1.xlsx"
            write_workbook(path, package)
            workbook = load_workbook(path)
            workbook["R7.1确认"]["D5"] = ""
            workbook.save(path)
            result = import_r7_exception_review(path, dataset, package, dataset_id="r7-1", schema=schema)
        self.assertFalse(result.ready)
        self.assertIn("REVIEWER_REQUIRED", {issue.code for issue in result.issues})


if __name__ == "__main__":
    unittest.main()
