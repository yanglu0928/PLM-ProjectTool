from __future__ import annotations

import sys
import tempfile
import unittest
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook


POC_DIR = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(POC_DIR / "src"))

from poc03_rag.r5_label_review import (  # noqa: E402
    CONFLICT_HEADERS,
    GLOBAL_CONFIRMATION,
    RULE_HEADERS,
    build_sanitized_r5_report,
    import_r5_label_review,
)


def make_case(index: int, classification: str, source_type: str = "CONTRACT") -> dict:
    return {
        "case_id": f"GD-{index:04d}",
        "query": f"业务问题 {index}",
        "scope": "PROJECT",
        "project_id": "PROJECT-A",
        "source_type": source_type,
        "expected_classification": classification,
        "expected_relevant_chunk_ids": [f"DOC-C-{index:04d}"],
        "expected_answer_terms": ["术语"],
        "expected_citations": [
            {
                "document_id": "DOC",
                "chunk_id": f"DOC-C-{index:04d}",
                "source_locator": f"word/paragraph/{index}",
            }
        ],
        "review": {
            "status": "APPROVED",
            "reviewed_by": "旧审核人",
            "reviewed_at": "2026-09-18T00:00:00Z",
        },
    }


def write_r4(path: Path, cases: list[dict]) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "确认清单"
    for _ in range(8):
        sheet.append([])
    notes = {
        1: "人工复核：证据片段不足。结论：按“资料不足”处理。",
        2: "",
        3: "",
    }
    for index, case in enumerate(cases, start=1):
        sheet.append(
            [
                f"CONF-{index:04d}",
                case["source_type"],
                "中",
                case["query"],
                "AI建议",
                "打开证据",
                "修改后确认",
                notes[index],
                "杨璐",
                datetime(2026, 9, 18),
                "已完成",
            ]
        )
    workbook.save(path)
    workbook.close()


def write_r5(path: Path, cases: list[dict], *, confirmed: bool) -> None:
    workbook = Workbook()
    rules = workbook.active
    rules.title = "批量规则"
    conflicts = workbook.create_sheet("冲突确认")
    workbook.create_sheet("明确结论")
    workbook.create_sheet("技术底稿")
    rules["B5"] = GLOBAL_CONFIRMATION if confirmed else "未确认"
    rules["D5"] = "杨璐"
    rules["F5"] = datetime(2026, 9, 18)
    for column, value in enumerate(RULE_HEADERS, start=1):
        rules.cell(9, column).value = value
    group_rows = [
        ["GRP-01", "需人工确认", "标准满足", 1, "采用本次AI", "", "", ""],
        ["GRP-02", "非标准", "部分满足", 1, "沿用R4", "", "", ""],
    ]
    for row_index, values in enumerate(group_rows, start=10):
        for column, value in enumerate(values, start=1):
            rules.cell(row_index, column).value = value
    for column, value in enumerate(CONFLICT_HEADERS, start=1):
        conflicts.cell(8, column).value = value
    conflict_rows = [
        ["CONF-0002", "CONTRACT", "中", cases[1]["query"], "需人工确认", "标准满足", "GRP-01", "打开证据", "", "", "", ""],
        ["CONF-0003", "CONTRACT", "中", cases[2]["query"], "非标准", "部分满足", "GRP-02", "打开证据", "", "", "", ""],
    ]
    for row_index, values in enumerate(conflict_rows, start=9):
        for column, value in enumerate(values, start=1):
            conflicts.cell(row_index, column).value = value
    workbook.save(path)
    workbook.close()


class R5LabelReviewTests(unittest.TestCase):
    def setUp(self) -> None:
        self.temp_dir = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp_dir.cleanup)
        self.root = Path(self.temp_dir.name)
        self.cases = [
            make_case(1, "INSUFFICIENT_INFORMATION"),
            make_case(2, "HUMAN_CONFIRMATION_REQUIRED"),
            make_case(3, "NON_STANDARD"),
        ]
        self.golden = {
            "schema_version": "poc-03.golden.v1",
            "dataset_id": "r4",
            "embedding_index": {},
            "cases": self.cases,
        }
        self.quality = {
            "predictions": {
                "GD-0001": {"classification": "STANDARD_SATISFIED"},
                "GD-0002": {"classification": "STANDARD_SATISFIED"},
                "GD-0003": {"classification": "PARTIALLY_SATISFIED"},
            }
        }
        self.r4_path = self.root / "r4.xlsx"
        self.r5_path = self.root / "r5.xlsx"
        write_r4(self.r4_path, self.cases)

    def import_result(self):
        return import_r5_label_review(
            self.r5_path,
            self.r4_path,
            self.golden,
            self.quality,
        )

    def test_unconfirmed_package_remains_pending(self) -> None:
        write_r5(self.r5_path, self.cases, confirmed=False)
        result = self.import_result()
        self.assertFalse(result.has_errors)
        self.assertFalse(result.ready)
        self.assertEqual([], result.cases)
        self.assertEqual(2, result.status_counts["PENDING"])

    def test_confirmed_recommendations_resolve_all_cases(self) -> None:
        write_r5(self.r5_path, self.cases, confirmed=True)
        result = self.import_result()
        self.assertTrue(result.ready)
        self.assertEqual(3, len(result.cases))
        labels = {case["case_id"]: case["expected_classification"] for case in result.cases}
        self.assertEqual("INSUFFICIENT_INFORMATION", labels["GD-0001"])
        self.assertEqual("STANDARD_SATISFIED", labels["GD-0002"])
        self.assertEqual("NON_STANDARD", labels["GD-0003"])

    def test_group_and_row_override_precedence(self) -> None:
        write_r5(self.r5_path, self.cases, confirmed=True)
        workbook = load_workbook(self.r5_path)
        workbook["批量规则"]["F10"] = "沿用R4"
        workbook["冲突确认"]["I9"] = "部分满足"
        workbook.save(self.r5_path)
        workbook.close()
        result = self.import_result()
        labels = {case["case_id"]: case["expected_classification"] for case in result.cases}
        self.assertTrue(result.ready)
        self.assertEqual("PARTIALLY_SATISFIED", labels["GD-0002"])

    def test_tampered_query_and_group_are_rejected(self) -> None:
        write_r5(self.r5_path, self.cases, confirmed=True)
        workbook = load_workbook(self.r5_path)
        workbook["冲突确认"]["D9"] = "被修改的问题"
        workbook["冲突确认"]["G10"] = "GRP-99"
        workbook.save(self.r5_path)
        workbook.close()
        result = self.import_result()
        self.assertFalse(result.ready)
        self.assertEqual(2, sum(issue.code == "SOURCE_FIELD_CHANGED" for issue in result.issues))

    def test_sanitized_report_contains_no_queries_or_reviewer(self) -> None:
        write_r5(self.r5_path, self.cases, confirmed=False)
        result = self.import_result()
        report = build_sanitized_r5_report(result, generated_at="2026-09-18T00:00:00Z", output_written=False)
        serialized = str(report)
        self.assertEqual("AWAITING_HUMAN_CONFIRMATION", report["status"])
        self.assertNotIn("业务问题", serialized)
        self.assertNotIn("杨璐", serialized)


if __name__ == "__main__":
    unittest.main()
