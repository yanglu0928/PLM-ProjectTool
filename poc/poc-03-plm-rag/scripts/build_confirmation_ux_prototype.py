from __future__ import annotations

import argparse
import json
import sys
from datetime import UTC, date, datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


SCRIPT_DIR = Path(__file__).resolve().parent
POC_DIR = SCRIPT_DIR.parent
sys.path.insert(0, str(POC_DIR / "src"))

from poc03_rag.confirmation_prototype import (  # noqa: E402
    build_confirmation_tasks,
    build_evidence_navigator_html,
    build_sanitized_confirmation_report,
)
from poc03_rag.review_import import load_candidate_records  # noqa: E402


def _value(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    return str(value).strip()


def _review_rows(workbook_path: Path) -> list[dict[str, Any]]:
    workbook = load_workbook(workbook_path, read_only=True, data_only=False)
    try:
        sheet = workbook["候选评审"]
        rows: list[dict[str, Any]] = []
        for values in sheet.iter_rows(min_row=6, max_col=19, values_only=True):
            if not _value(values[0]):
                continue
            rows.append(
                {
                    "candidate_id": _value(values[0]),
                    "query": _value(values[9]),
                    "source_type": _value(values[10]),
                    "classification": _value(values[11]),
                    "answer_terms": _value(values[12]),
                    "confirmed_locator": _value(values[13]),
                    "review_status": _value(values[14]),
                    "reviewed_by": _value(values[15]),
                    "reviewed_at": _value(values[16]),
                    "review_note": _value(values[17]),
                }
            )
        return rows
    finally:
        workbook.close()


def _parsed_document(document_id: str, contract_dir: Path, solution_dir: Path) -> dict[str, Any]:
    corpus, parsed_id = document_id.split("-", 1)
    root = contract_dir if corpus == "CONTRACT" else solution_dir
    path = root / f"{parsed_id}.parsed.json"
    if not path.is_file():
        return {}
    return json.loads(path.read_text(encoding="utf-8"))


def _find_original(source_root: Path, file_name: str) -> Path | None:
    if not source_root.is_dir() or not file_name:
        return None
    matches = [path for path in source_root.rglob(file_name) if path.is_file()]
    return matches[0] if len(matches) == 1 else None


def _original_url(path: Path | None, locators: list[str]) -> str:
    if path is None:
        return ""
    uri = path.resolve().as_uri()
    if path.suffix.lower() == ".pdf":
        for locator in locators:
            parts = locator.split("/")
            if len(parts) >= 3 and parts[0] == "pdf" and parts[1] == "pages":
                return f"{uri}#page={parts[2]}"
    return uri


def main() -> int:
    parser = argparse.ArgumentParser(description="Build a local POC-03 confirmation UX prototype.")
    parser.add_argument("--workbook", type=Path, required=True)
    parser.add_argument("--candidates", type=Path, required=True)
    parser.add_argument("--contract-parsed-dir", type=Path, required=True)
    parser.add_argument("--solution-parsed-dir", type=Path, required=True)
    parser.add_argument("--contract-source-root", type=Path, required=True)
    parser.add_argument("--solution-source-root", type=Path, required=True)
    parser.add_argument("--task-output", type=Path, required=True)
    parser.add_argument("--navigator-output", type=Path, required=True)
    parser.add_argument("--sanitized-report", type=Path, required=True)
    args = parser.parse_args()

    candidates = load_candidate_records(args.candidates)
    review_rows = _review_rows(args.workbook)
    tasks = build_confirmation_tasks(candidates, review_rows)
    review_by_id = {row["candidate_id"]: row for row in review_rows}
    navigator_entries: list[dict[str, Any]] = []
    for candidate in candidates:
        candidate_id = candidate["candidate_id"]
        citation = (candidate.get("expected_citations") or [{}])[0]
        document_id = citation.get("document_id", "")
        locators = citation.get("source_locators") or []
        parsed = _parsed_document(
            document_id,
            args.contract_parsed_dir,
            args.solution_parsed_dir,
        )
        file_name = str((parsed.get("source") or {}).get("file_name") or "")
        source_root = (
            args.contract_source_root
            if document_id.startswith("CONTRACT-")
            else args.solution_source_root
        )
        original = _find_original(source_root, file_name)
        navigator_entries.append(
            {
                "candidate_id": candidate_id,
                "query": review_by_id[candidate_id]["query"],
                "source_name": file_name,
                "document_id": document_id,
                "content": candidate.get("candidate_content", ""),
                "locators": locators,
                "original_url": _original_url(original, locators),
            }
        )

    args.task_output.parent.mkdir(parents=True, exist_ok=True)
    args.task_output.write_text(
        json.dumps(tasks, ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8",
    )
    args.navigator_output.parent.mkdir(parents=True, exist_ok=True)
    args.navigator_output.write_text(
        build_evidence_navigator_html(navigator_entries),
        encoding="utf-8",
    )
    report = build_sanitized_confirmation_report(
        tasks,
        generated_at=datetime.now(UTC).isoformat(),
    )
    args.sanitized_report.parent.mkdir(parents=True, exist_ok=True)
    args.sanitized_report.write_text(
        json.dumps(report, ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8",
    )
    print(json.dumps(report, ensure_ascii=False, indent=2))
    return 0 if report["status"] == "PASS_FOR_UX_REVIEW" else 1


if __name__ == "__main__":
    raise SystemExit(main())
