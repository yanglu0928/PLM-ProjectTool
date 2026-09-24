from __future__ import annotations

import argparse
import json
import sys
from collections import Counter
from datetime import UTC, datetime
from pathlib import Path

from openpyxl import load_workbook


SCRIPT_DIR = Path(__file__).resolve().parent
POC_DIR = SCRIPT_DIR.parent
sys.path.insert(0, str(POC_DIR / "src"))

from poc03_rag.action_list_import import import_action_list  # noqa: E402
from poc03_rag.review_import import (  # noqa: E402
    build_golden_dataset,
    build_sanitized_import_report,
    load_candidate_records,
)


def _decision_counts(workbook_path: Path) -> dict[str, int]:
    workbook = load_workbook(workbook_path, read_only=True, data_only=False)
    try:
        sheet = workbook["确认清单"]
        counts: Counter[str] = Counter()
        for row in sheet.iter_rows(min_row=9, max_row=128, max_col=7, values_only=True):
            if not str(row[0] or "").strip():
                continue
            decision = str(row[6] or "").strip() or "BLANK"
            counts[decision] += 1
        return dict(sorted(counts.items()))
    finally:
        workbook.close()


def main() -> int:
    parser = argparse.ArgumentParser(
        description="Validate a local R4 action list and export approved cases."
    )
    parser.add_argument("--workbook", type=Path, required=True)
    parser.add_argument("--tasks", type=Path, required=True)
    parser.add_argument("--candidates", type=Path, required=True)
    parser.add_argument("--report", type=Path, required=True)
    parser.add_argument(
        "--schema",
        type=Path,
        default=POC_DIR / "schema" / "golden-dataset.schema.json",
    )
    parser.add_argument("--output", type=Path)
    parser.add_argument("--dataset-id")
    parser.add_argument("--embedding-provider")
    parser.add_argument("--embedding-model")
    parser.add_argument("--embedding-dimension", type=int)
    parser.add_argument("--index-version")
    parser.add_argument("--timezone", default="Asia/Shanghai")
    parser.add_argument("--validate-only", action="store_true")
    args = parser.parse_args()

    candidates = load_candidate_records(args.candidates)
    tasks = json.loads(args.tasks.read_text(encoding="utf-8"))
    result = import_action_list(
        args.workbook,
        candidates,
        tasks,
        timezone_name=args.timezone,
    )
    output_written = False
    if not args.validate_only and not result.has_errors and 100 <= len(result.cases) <= 200:
        required = {
            "--output": args.output,
            "--dataset-id": args.dataset_id,
            "--embedding-provider": args.embedding_provider,
            "--embedding-model": args.embedding_model,
            "--embedding-dimension": args.embedding_dimension,
            "--index-version": args.index_version,
        }
        missing = [name for name, value in required.items() if value in (None, "")]
        if missing:
            raise ValueError(f"Export requires: {', '.join(missing)}")
        schema = json.loads(args.schema.read_text(encoding="utf-8"))
        dataset = build_golden_dataset(
            result,
            dataset_id=args.dataset_id,
            provider=args.embedding_provider,
            model=args.embedding_model,
            dimension=args.embedding_dimension,
            index_version=args.index_version,
            schema=schema,
        )
        args.output.parent.mkdir(parents=True, exist_ok=True)
        args.output.write_text(
            json.dumps(dataset, ensure_ascii=False, indent=2) + "\n",
            encoding="utf-8",
        )
        output_written = True

    report = build_sanitized_import_report(
        result,
        generated_at=datetime.now(UTC).isoformat(),
        output_written=output_written,
    )
    report["schema_version"] = "poc-03.action-list-import-result.v1"
    report["summary"]["minimum_required_approved_count"] = 100
    report["summary"]["approved_count_gap"] = max(0, 100 - len(result.cases))
    report["summary"]["human_decision_counts"] = _decision_counts(args.workbook)
    args.report.parent.mkdir(parents=True, exist_ok=True)
    args.report.write_text(
        json.dumps(report, ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8",
    )
    print(json.dumps(report, ensure_ascii=False, indent=2))
    if result.has_errors:
        return 1
    if args.validate_only:
        return 0
    return 0 if output_written else 2


if __name__ == "__main__":
    raise SystemExit(main())
