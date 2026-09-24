from __future__ import annotations

import argparse
import json
import sys
from datetime import UTC, date, datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


SCRIPT_DIR = Path(__file__).resolve().parent
POC_DIR = SCRIPT_DIR.parent
sys.path.insert(0, str(POC_DIR / "src"))

from poc03_rag.review_import import load_candidate_records  # noqa: E402
from poc03_rag.source_type_audit import (  # noqa: E402
    audit_candidate_source_types,
    build_sanitized_source_type_report,
)


def _value(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    return str(value).strip()


def _review_rows(workbook_path: Path) -> list[dict[str, str]]:
    workbook = load_workbook(workbook_path, read_only=True, data_only=False)
    try:
        sheet = workbook["候选评审"]
        return [
            {
                "candidate_id": _value(values[0]),
                "source_type": _value(values[10]),
            }
            for values in sheet.iter_rows(min_row=6, max_col=19, values_only=True)
            if _value(values[0])
        ]
    finally:
        workbook.close()


def _documents(parsed_dirs: list[Path]) -> list[dict[str, str]]:
    documents: list[dict[str, str]] = []
    for parsed_dir in parsed_dirs:
        corpus = "CONTRACT" if "contract" in parsed_dir.as_posix().lower() else "SOLUTION"
        for path in sorted(parsed_dir.glob("*.parsed.json")):
            parsed = json.loads(path.read_text(encoding="utf-8"))
            parsed_id = path.name.removesuffix(".parsed.json")
            documents.append(
                {
                    "document_id": f"{corpus}-{parsed_id}",
                    "file_name": str((parsed.get("source") or {}).get("file_name") or ""),
                }
            )
    return documents


def main() -> int:
    parser = argparse.ArgumentParser(description="Audit POC-03 source-type eligibility.")
    parser.add_argument("--workbook", type=Path, required=True)
    parser.add_argument("--candidates", type=Path, required=True)
    parser.add_argument("--parsed-dir", type=Path, action="append", required=True)
    parser.add_argument("--local-mapping", type=Path, required=True)
    parser.add_argument("--sanitized-report", type=Path, required=True)
    args = parser.parse_args()

    rows = audit_candidate_source_types(
        load_candidate_records(args.candidates),
        _review_rows(args.workbook),
        _documents(args.parsed_dir),
    )
    args.local_mapping.parent.mkdir(parents=True, exist_ok=True)
    args.local_mapping.write_text(
        json.dumps(rows, ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8",
    )
    report = build_sanitized_source_type_report(
        rows,
        generated_at=datetime.now(UTC).isoformat(),
    )
    args.sanitized_report.parent.mkdir(parents=True, exist_ok=True)
    args.sanitized_report.write_text(
        json.dumps(report, ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8",
    )
    print(json.dumps(report, ensure_ascii=False, indent=2))
    return 0 if report["status"] == "PASS" else 1


if __name__ == "__main__":
    raise SystemExit(main())
