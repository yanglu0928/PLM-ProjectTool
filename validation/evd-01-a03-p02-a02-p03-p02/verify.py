"""Synthetic on-disk Office sources, independently reopened at Parser locators."""

from __future__ import annotations

import hashlib
import tempfile
import uuid
from pathlib import Path

from docx import Document
from openpyxl import Workbook, load_workbook
from pptx import Presentation
from pptx.util import Inches

from plm_assistant.modules.document.application.read_documents import DocumentReadQuery
from plm_assistant.modules.document.application.read_parse_result import VerifiedParseResult
from plm_assistant.modules.evidence.application.parsed_node_proof import ParsedNodeEvidenceProofService
from plm_assistant.modules.parser.application.extract_office import extract_office
from plm_assistant.modules.parser.application.prepare_input import VerifiedParserInput
from plm_assistant.modules.parser.application.profile_selection import (
    ParserInputVersion, choose_parser_profile,
)


MIMES = {
    "DOCX": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "PPTX": "application/vnd.openxmlformats-officedocument.presentationml.presentation",
    "XLSX": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
}


class ResultPort:
    def __init__(self, value: VerifiedParseResult) -> None:
        self.value = value

    def read(self, *_args, **_kwargs) -> VerifiedParseResult:
        return self.value


def parse_file(path: Path, profile: str):
    raw = path.read_bytes()
    source = ParserInputVersion(uuid.uuid4(), hashlib.sha256(raw).digest(),
                                len(raw), MIMES[profile])
    with path.open("rb") as stream:
        prepared = VerifiedParserInput(choose_parser_profile(source), uuid.uuid4(), 1, 1, stream)
        parsed = extract_office(prepared)
    assert parsed.parser_profile == profile and parsed.source_sha256 == source.content_sha256
    return source, parsed


def prove(source: ParserInputVersion, parsed) -> None:
    content = parsed.canonical_bytes()
    record_id, document_id = uuid.uuid4(), uuid.uuid4()
    result = VerifiedParseResult(
        record_id, source.document_version_id, uuid.uuid4(),
        parsed.parser_profile, parsed.parser_version, source.content_sha256,
        hashlib.sha256(content).digest(), content,
    )
    service = ParsedNodeEvidenceProofService(results=ResultPort(result))
    query = DocumentReadQuery(b"s" * 32, uuid.uuid4(), "PROJECT", uuid.uuid4())
    for node in parsed.nodes:
        locator = node.position.to_locator()
        direct = service.prove(query, document_id=document_id,
                               document_version_id=source.document_version_id,
                               parse_record_id=record_id, locator=locator)
        assert (direct.node_id == node.node_id
                and direct.content_fingerprint == hashlib.sha256(node.text.encode()).digest())
        structured = {"locator_type": "STRUCTURED_NODE",
                      "parse_record_id": str(record_id), "node_id": node.node_id,
                      "source_locator": locator}
        assert service.prove(query, document_id=document_id,
                             document_version_id=source.document_version_id,
                             parse_record_id=record_id, locator=structured).node_id == node.node_id


def docx(path: Path) -> None:
    document = Document()
    document.add_paragraph("需求甲")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "接口"
    table.cell(0, 1).text = "确认"
    document.add_paragraph("结论乙")
    document.save(path)
    source, parsed = parse_file(path, "DOCX")
    reopened = Document(path)
    assert len(parsed.nodes) == 4
    for node in parsed.nodes:
        locator = node.position.to_locator()
        if locator["locator_type"] == "PARAGRAPH":
            assert reopened.paragraphs[locator["paragraph_index"] - 1].text == node.text
        else:
            index = int(locator["table_anchor"].split("/")[-1]) - 1
            assert reopened.tables[index].cell(locator["row_no"] - 1,
                                               locator["column_no"] - 1).text == node.text
    prove(source, parsed)


def pptx(path: Path) -> None:
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(3), Inches(1))
    box.text = "进度甲"
    table = slide.shapes.add_table(1, 1, Inches(1), Inches(2), Inches(3), Inches(1))
    table.table.cell(0, 0).text = "里程碑乙"
    presentation.save(path)
    source, parsed = parse_file(path, "PPTX")
    reopened = Presentation(path)
    assert len(parsed.nodes) == 2
    for node in parsed.nodes:
        locator = node.position.to_locator()
        slide = reopened.slides[locator.get("slide_no", 1) - 1]
        shape_id = (locator["shape_id"] if locator["locator_type"] == "SLIDE_SHAPE"
                    else locator["table_anchor"].split("/")[-1])
        matched = [shape for shape in slide.shapes if str(shape.shape_id) == shape_id]
        assert len(matched) == 1
        shape = matched[0]
        if locator["locator_type"] == "SLIDE_SHAPE":
            assert shape.text == node.text
        else:
            assert shape.table.cell(locator["row_no"] - 1,
                                    locator["column_no"] - 1).text == node.text
    prove(source, parsed)


def xlsx(path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "业务调研"
    sheet["B2"] = "项目甲"
    sheet["C3"] = "=1+1"
    workbook.save(path)
    workbook.close()
    source, parsed = parse_file(path, "XLSX")
    reopened = load_workbook(path, read_only=True, data_only=False, keep_links=False)
    try:
        assert len(parsed.nodes) == 2
        for node in parsed.nodes:
            locator = node.position.to_locator()
            assert locator["start_cell"] == locator["end_cell"]
            assert str(reopened[locator["sheet_name"]][locator["start_cell"]].value) == node.text
        assert parsed.nodes[1].text == "=1+1"
    finally:
        reopened.close()
    prove(source, parsed)


def main() -> None:
    with tempfile.TemporaryDirectory(prefix="plm-evd-office-") as scratch:
        root = Path(scratch)
        docx(root / "synthetic.docx")
        pptx(root / "synthetic.pptx")
        xlsx(root / "synthetic.xlsx")
    print("PASS: real synthetic DOCX/PPTX/XLSX files reopened at paragraph/table/shape/cell "
          "locators and matched by Evidence proof")


if __name__ == "__main__":
    main()
