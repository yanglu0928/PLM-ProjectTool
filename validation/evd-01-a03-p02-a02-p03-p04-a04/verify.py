"""On-disk DOCX heading -> fixed Parser result -> SECTION Evidence proof."""

from __future__ import annotations

import hashlib
import tempfile
import uuid
from pathlib import Path

from docx import Document

from plm_assistant.modules.document.application.read_documents import DocumentReadQuery
from plm_assistant.modules.document.application.read_parse_result import VerifiedParseResult
from plm_assistant.modules.evidence.application.parsed_node_proof import (
    EvidenceNodeProofError, ParsedNodeEvidenceProofService,
)
from plm_assistant.modules.parser.application.extract_office import extract_office
from plm_assistant.modules.parser.application.prepare_input import VerifiedParserInput
from plm_assistant.modules.parser.application.profile_selection import (
    ParserInputVersion, choose_parser_profile,
)


MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


class ResultPort:
    def __init__(self, result: VerifiedParseResult) -> None:
        self.result = result

    def read(self, *_args, **_kwargs) -> VerifiedParseResult:
        return self.result


def verify(path: Path) -> None:
    document = Document()
    document.add_paragraph("范围", style="Heading 1")
    document.add_paragraph("范围")
    document.add_paragraph("范围", style="Heading 2")
    document.save(path)
    raw = path.read_bytes()
    source = ParserInputVersion(uuid.uuid4(), hashlib.sha256(raw).digest(), len(raw), MIME)
    with path.open("rb") as stream:
        parsed = extract_office(VerifiedParserInput(choose_parser_profile(source),
                                                     uuid.uuid4(), 1, 1, stream))
    content = parsed.canonical_bytes()
    record_id, document_id = uuid.uuid4(), uuid.uuid4()
    result = VerifiedParseResult(
        record_id, source.document_version_id, uuid.uuid4(), parsed.parser_profile,
        parsed.parser_version, source.content_sha256,
        hashlib.sha256(content).digest(), content,
    )
    port = ResultPort(result)
    proof = ParsedNodeEvidenceProofService(results=port)
    query = DocumentReadQuery(b"s" * 32, uuid.uuid4(), "PROJECT", uuid.uuid4())
    sections = [node for node in parsed.nodes if node.kind == "DOCX_SECTION"]
    assert len(sections) == 2
    reopened = Document(path)
    for node in sections:
        locator = node.position.to_locator()
        _, _, level, index = locator["section_path"].split("/")
        paragraph = reopened.paragraphs[int(index) - 1]
        assert paragraph.text == node.text
        assert paragraph.style.style_id == f"Heading{level}"
        direct = proof.prove(query, document_id=document_id,
                             document_version_id=source.document_version_id,
                             parse_record_id=record_id, locator=locator)
        assert direct.node_id == node.node_id
        assert direct.content_fingerprint == hashlib.sha256(node.text.encode()).digest()
        structured = {"locator_type": "STRUCTURED_NODE",
                      "parse_record_id": str(record_id), "node_id": node.node_id,
                      "source_locator": locator}
        assert proof.prove(query, document_id=document_id,
                           document_version_id=source.document_version_id,
                           parse_record_id=record_id, locator=structured).node_id == node.node_id
    ordinary = {"locator_type": "SECTION", "section_path": "word/heading/1/2"}
    try:
        proof.prove(query, document_id=document_id,
                    document_version_id=source.document_version_id,
                    parse_record_id=record_id, locator=ordinary)
    except EvidenceNodeProofError:
        pass
    else:
        raise AssertionError("ordinary paragraph accepted as a section")


def main() -> None:
    with tempfile.TemporaryDirectory(prefix="plm-evd-section-proof-") as scratch:
        verify(Path(scratch) / "synthetic.docx")
    print("PASS: on-disk DOCX headings became fixed SECTION proofs; ordinary text rejected")


if __name__ == "__main__":
    main()
