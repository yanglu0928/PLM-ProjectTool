"""Replay DOCX v2 heading anchors against a synthetic on-disk source."""

from __future__ import annotations

import hashlib
import tempfile
import uuid
from dataclasses import replace
from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE

from plm_assistant.modules.parser.application.extract_office import extract_office
from plm_assistant.modules.parser.application.prepare_input import VerifiedParserInput
from plm_assistant.modules.parser.application.profile_selection import (
    ParserInputVersion, choose_parser_profile,
)
from plm_assistant.modules.parser.application.structured_result import ParserResultError


MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


def verify(path: Path) -> None:
    document = Document()
    document.add_paragraph("项目范围", style="Heading 1")
    document.add_paragraph("普通正文")
    document.add_paragraph("", style="Heading 2")
    custom = document.styles.add_style("CustomHeading", WD_STYLE_TYPE.PARAGRAPH)
    document.add_paragraph("自定义标题", style=custom)
    document.add_paragraph("项目范围", style="Heading 2")
    document.save(path)
    raw = path.read_bytes()
    source = ParserInputVersion(uuid.uuid4(), hashlib.sha256(raw).digest(), len(raw), MIME)
    assert choose_parser_profile(source).parser_version == "2"
    with path.open("rb") as stream:
        parsed = extract_office(VerifiedParserInput(choose_parser_profile(source),
                                                     uuid.uuid4(), 1, 1, stream))
    assert parsed.parser_version == "2"
    reopened = Document(path)
    sections = [node for node in parsed.nodes if node.kind == "DOCX_SECTION"]
    assert len(sections) == 2
    assert [node.node_id for node in sections] == ["heading:1", "heading:5"]
    for node in sections:
        locator = node.position.to_locator()
        _, _, level, index = locator["section_path"].split("/")
        paragraph = reopened.paragraphs[int(index) - 1]
        assert paragraph.style.style_id == f"Heading{level}"
        assert paragraph.text == node.text == "项目范围"
    assert all(node.text != "自定义标题" for node in sections)
    paragraphs = tuple(node for node in parsed.nodes if node.kind == "DOCX_PARAGRAPH")
    assert len(paragraphs) == 4
    assert replace(parsed, parser_version="1", nodes=paragraphs).parser_version == "1"
    try:
        replace(parsed, parser_version="1")
    except ParserResultError:
        pass
    else:
        raise AssertionError("v1 accepted a v2 section node")


def main() -> None:
    with tempfile.TemporaryDirectory(prefix="plm-evd-docx-headings-") as scratch:
        verify(Path(scratch) / "synthetic.docx")
    print("PASS: on-disk DOCX v2 heading anchors independently reopened; v1 node guard intact")


if __name__ == "__main__":
    main()
