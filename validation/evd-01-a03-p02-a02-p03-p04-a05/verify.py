"""Isolated PG18: immutable DOCX v1 and v2 results coexist on one version."""

from __future__ import annotations

import hashlib
import io
import shutil
import subprocess
import tempfile
import uuid
from dataclasses import replace
from importlib.util import module_from_spec, spec_from_file_location
from pathlib import Path

import psycopg
from alembic import command
from docx import Document
from psycopg.types.json import Jsonb
from sqlalchemy.engine import URL

from plm_assistant.modules.document.application.prepare_download import (
    DownloadError, VerifiedDownload,
)
from plm_assistant.modules.document.application.read_documents import DocumentReadQuery
from plm_assistant.modules.document.application.read_parse_result import (
    DocumentParseResultReadService, ParseResultReadError,
)
from plm_assistant.modules.document.infrastructure.parse_result_read_repository import (
    SqlAlchemyParseResultReadRepository,
)
from plm_assistant.modules.document.infrastructure.parse_result_storage import LocalParseResultStorage
from plm_assistant.modules.evidence.application.parsed_node_proof import (
    EvidenceNodeProofError, ParsedNodeEvidenceProofService,
)
from plm_assistant.modules.parser.application.extract_office import extract_office
from plm_assistant.modules.parser.application.prepare_input import VerifiedParserInput
from plm_assistant.modules.parser.application.profile_selection import (
    ParserInputVersion, choose_parser_profile,
)
from plm_assistant.modules.platform.infrastructure.database import create_database_runtime
from plm_assistant.modules.platform.infrastructure.migration import create_migration_config


ROOT = Path(__file__).resolve().parents[2]
PREVIOUS = ROOT / "validation/evd-01-a03-p02-a02-p01/verify.py"
SPEC = spec_from_file_location("evd_pg_helpers", PREVIOUS)
assert SPEC is not None and SPEC.loader is not None
HELPER = module_from_spec(SPEC)
SPEC.loader.exec_module(HELPER)
MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


class Snapshots:
    def __init__(self, version: uuid.UUID, project: uuid.UUID, content: bytes) -> None:
        self.version, self.project, self.content = version, project, content
        self.allowed = True

    def prepare(self, query: DocumentReadQuery, _document_id: uuid.UUID,
                version_id: uuid.UUID) -> VerifiedDownload:
        if (not self.allowed or query.scope != "PROJECT"
                or query.project_id != self.project or version_id != self.version):
            raise DownloadError("RESOURCE_NOT_FOUND")
        return VerifiedDownload(version_id, len(self.content), MIME,
                                hashlib.sha256(self.content).digest(), io.BytesIO(self.content))


def fixture_docx() -> bytes:
    document = Document()
    document.add_paragraph("范围", style="Heading 1")
    document.add_paragraph("正文")
    stream = io.BytesIO()
    document.save(stream)
    return stream.getvalue()


def seed_result(db: psycopg.Connection, storage: LocalParseResultStorage, *,
                actor: uuid.UUID, project: uuid.UUID, document: uuid.UUID,
                version: uuid.UUID, parsed) -> uuid.UUID:
    job = db.execute(
        "INSERT INTO plm.job_jobs(owner_module,job_type,scope,project_id,actor_ref,"
        "trace_id,payload_refs,idempotency_key,max_attempts) VALUES "
        "('document','DOCUMENT_PARSE','PROJECT',%s,%s,%s,%s,%s,3) RETURNING job_id",
        (project, actor, str(uuid.uuid4()),
         Jsonb({"document_id": str(document), "document_version_id": str(version)}),
         f"evd-docx-v{parsed.parser_version}"),
    ).fetchone()[0]
    record = db.execute(
        "INSERT INTO plm.doc_parse_records(document_version_id,scope,project_id,"
        "parser_profile,parser_version,job_ref,attempt_no) VALUES "
        "(%s,'PROJECT',%s,'DOCX',%s,%s,1) RETURNING parse_record_id",
        (version, project, parsed.parser_version, job),
    ).fetchone()[0]
    db.execute("UPDATE plm.doc_parse_records SET parse_state='RUNNING',"
               "started_at=statement_timestamp(),lock_version=1 WHERE parse_record_id=%s",
               (record,))
    result_id = uuid.uuid4()
    stored = storage.write_once(scope="PROJECT", project_id=project,
                                result_ref_id=result_id, content=parsed.canonical_bytes())
    db.execute(
        "INSERT INTO plm.doc_parse_result_refs(parse_result_ref_id,parse_record_id,"
        "storage_locator,result_schema_version,sha256,size_bytes) VALUES (%s,%s,%s,1,%s,%s)",
        (result_id, record, stored.storage_locator, stored.sha256, stored.size_bytes),
    )
    db.execute("UPDATE plm.doc_parse_records SET parse_state='SUCCEEDED',"
               "completed_at=statement_timestamp(),result_ref=%s,result_sha256=%s,"
               "retryable=false,lock_version=2 WHERE parse_record_id=%s",
               (result_id, stored.sha256, record))
    return record


def verify(port: int, scratch: Path) -> None:
    url = URL.create("postgresql+psycopg", username=HELPER.USER, host="127.0.0.1",
                     port=port, database="postgres")
    command.upgrade(create_migration_config(url), "head")
    raw = fixture_docx()
    source_sha = hashlib.sha256(raw).digest()
    data_root = scratch / "private-results"
    data_root.mkdir()
    storage = LocalParseResultStorage(data_root)
    with psycopg.connect(host="127.0.0.1", port=port, user=HELPER.USER,
                         dbname="postgres", autocommit=True) as db:
        actor = db.execute(
            "INSERT INTO plm.auth_users(username_display,username_normalized) "
            "VALUES ('Evidence Section','evidence section') RETURNING user_id"
        ).fetchone()[0]
        project = db.execute(
            "INSERT INTO plm.prj_projects(project_code,project_code_normalized,name,created_by) "
            "VALUES ('EVDS','evds','Evidence section',%s) RETURNING project_id",
            (actor,),
        ).fetchone()[0]
        document = db.execute(
            "INSERT INTO plm.doc_documents(scope,project_id,document_category,title,"
            "original_display_name,created_by) VALUES "
            "('PROJECT',%s,'PROJECT_RECORD','Section','synthetic.docx',%s) "
            "RETURNING document_id", (project, actor),
        ).fetchone()[0]
        file_id = db.execute(
            "INSERT INTO plm.doc_file_objects(scope,project_id,storage_class,storage_locator,"
            "original_name_metadata,created_by,file_state,sha256,size_bytes,detected_mime,available_at) "
            "VALUES ('PROJECT',%s,'PERSISTENT','synthetic/only','synthetic.docx',%s,"
            "'AVAILABLE',%s,%s,%s,statement_timestamp()) RETURNING file_object_id",
            (project, actor, source_sha, len(raw), MIME),
        ).fetchone()[0]
        version = db.execute(
            "INSERT INTO plm.doc_document_versions(document_id,scope,project_id,version_no,"
            "file_object_id,content_sha256,size_bytes,detected_mime,source_metadata,created_by) "
            "VALUES (%s,'PROJECT',%s,1,%s,%s,%s,%s,%s,%s) "
            "RETURNING document_version_id",
            (document, project, file_id, source_sha, len(raw), MIME, Jsonb({}), actor),
        ).fetchone()[0]
        db.execute("UPDATE plm.doc_documents SET latest_version_ref=%s,effective_version_ref=%s "
                   "WHERE document_id=%s", (version, version, document))
        source = ParserInputVersion(version, source_sha, len(raw), MIME)
        with io.BytesIO(raw) as stream:
            parsed_v2 = extract_office(VerifiedParserInput(choose_parser_profile(source),
                                                            uuid.uuid4(), 1, 1, stream))
        parsed_v1 = replace(parsed_v2, parser_version="1", nodes=tuple(
            node for node in parsed_v2.nodes if node.kind != "DOCX_SECTION"))
        record_v1 = seed_result(db, storage, actor=actor, project=project,
                                document=document, version=version, parsed=parsed_v1)
        record_v2 = seed_result(db, storage, actor=actor, project=project,
                                document=document, version=version, parsed=parsed_v2)
        assert record_v1 != record_v2
        rows = db.execute(
            "SELECT parser_version,parse_state FROM plm.doc_parse_records "
            "WHERE document_version_id=%s ORDER BY parser_version", (version,)
        ).fetchall()
        assert rows == [("1", "SUCCEEDED"), ("2", "SUCCEEDED")], rows

    runtime = create_database_runtime(url)
    try:
        snapshots = Snapshots(version, project, raw)
        reader = DocumentParseResultReadService(
            documents=snapshots, metadata=SqlAlchemyParseResultReadRepository(),
            storage=storage, unit_of_work=runtime.unit_of_work,
        )
        proof = ParsedNodeEvidenceProofService(results=reader)
        query = DocumentReadQuery(b"s" * 32, uuid.uuid4(), "PROJECT", project)
        old = reader.read(query, document_id=document,
                          document_version_id=version, parse_record_id=record_v1)
        new = reader.read(query, document_id=document,
                          document_version_id=version, parse_record_id=record_v2)
        assert old.content == parsed_v1.canonical_bytes()
        assert new.content == parsed_v2.canonical_bytes()
        section = next(node for node in parsed_v2.nodes if node.kind == "DOCX_SECTION")
        section_locator = section.position.to_locator()
        assert proof.prove(query, document_id=document,
                           document_version_id=version, parse_record_id=record_v2,
                           locator=section_locator).node_id == section.node_id
        try:
            proof.prove(query, document_id=document,
                        document_version_id=version, parse_record_id=record_v1,
                        locator=section_locator)
        except EvidenceNodeProofError:
            pass
        else:
            raise AssertionError("v1 history accepted a v2 section")
        other = DocumentReadQuery(b"s" * 32, uuid.uuid4(), "PROJECT", uuid.uuid4())
        try:
            reader.read(other, document_id=document,
                        document_version_id=version, parse_record_id=record_v2)
        except ParseResultReadError as error:
            assert error.code == "RESOURCE_NOT_FOUND"
        else:
            raise AssertionError("cross-project result accepted")
        snapshots.allowed = False
        try:
            reader.read(query, document_id=document,
                        document_version_id=version, parse_record_id=record_v2)
        except ParseResultReadError as error:
            assert error.code == "RESOURCE_NOT_FOUND"
        else:
            raise AssertionError("revoked result accepted")
        snapshots.allowed = True
        with psycopg.connect(host="127.0.0.1", port=port, user=HELPER.USER,
                             dbname="postgres") as db:
            locator = db.execute(
                "SELECT r.storage_locator FROM plm.doc_parse_result_refs r "
                "JOIN plm.doc_parse_records p ON p.result_ref=r.parse_result_ref_id "
                "WHERE p.parse_record_id=%s", (record_v2,)
            ).fetchone()[0]
        (data_root / locator).write_bytes(b"tampered")
        try:
            reader.read(query, document_id=document,
                        document_version_id=version, parse_record_id=record_v2)
        except ParseResultReadError as error:
            assert error.code == "FILE_INTEGRITY_MISMATCH"
        else:
            raise AssertionError("tampered private result accepted")
        assert reader.read(query, document_id=document,
                           document_version_id=version,
                           parse_record_id=record_v1).content == parsed_v1.canonical_bytes()
    finally:
        runtime.dispose()
    print("PASS: isolated PG18 DOCX v1/v2 results coexist; v2 SECTION exact, v1 denied, "
          "cross-project/revocation/tamper fail closed")


def main() -> None:
    scratch = Path(tempfile.mkdtemp(prefix="plm-evd-docx-pg-"))
    if not str(scratch).isascii():
        raise RuntimeError("ASCII temporary PostgreSQL path required")
    install = scratch / "pgsql"
    for name in ("bin", "lib", "share"):
        shutil.copytree(HELPER.PG_SOURCE / name, install / name)
    shutil.copy2(HELPER.VECTOR_SOURCE / "vector.dll", install / "lib/vector.dll")
    shutil.copy2(HELPER.VECTOR_SOURCE / "vector.control",
                 install / "share/extension/vector.control")
    for path in (HELPER.VECTOR_SOURCE / "sql").glob("vector--*.sql"):
        shutil.copy2(path, install / "share/extension" / path.name)
    bin_dir = install / "bin"
    data = scratch / "data"
    log = scratch / "postgres.log"
    port = HELPER.free_port()
    started = False
    try:
        HELPER.run(str(bin_dir / "initdb.exe"), "-D", str(data), "-U", HELPER.USER,
                   "-A", "trust", "--no-locale", "-E", "UTF8")
        HELPER.run(str(bin_dir / "pg_ctl.exe"), "-D", str(data), "-l", str(log),
                   "-o", f"-h 127.0.0.1 -p {port}", "-w", "start", detached=True)
        started = True
        verify(port, scratch)
    finally:
        if started:
            HELPER.run(str(bin_dir / "pg_ctl.exe"), "-D", str(data), "-m", "fast", "-w", "stop")
        status = subprocess.run([str(bin_dir / "pg_ctl.exe"), "-D", str(data), "status"],
                                stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL,
                                timeout=15)
        if status.returncode == 0:
            raise RuntimeError(f"isolated PostgreSQL remains running; data preserved at {scratch}")
        if (scratch.resolve().is_relative_to(Path(tempfile.gettempdir()).resolve())
                and scratch.name.startswith("plm-evd-docx-pg-")):
            shutil.rmtree(scratch)


if __name__ == "__main__":
    main()
