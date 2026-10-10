"""Real Office extraction from verified private bytes; no Evidence publication."""

from __future__ import annotations

import hashlib
import hmac
import io
import re
import zipfile

from docx import Document
from docx.oxml.table import CT_Tbl
from docx.oxml.text.paragraph import CT_P
from docx.table import Table
from docx.text.paragraph import Paragraph
from openpyxl import load_workbook
from pptx import Presentation

from .prepare_input import VerifiedParserInput
from .profile_selection import choose_parser_profile
from .structured_result import (
    ParagraphPosition, ParsedNode, ParsedResult, ParserResultError,
    SectionPosition, SheetCellPosition, SlideShapePosition, TableCellPosition,
)


_MAX_ZIP_MEMBERS = 10_000
_MAX_UNCOMPRESSED_BYTES = 300_000_000
_MAX_NODES = 100_000
_MAX_RESULT_CHARS = 32_000_000
_MAX_SHEET_SCAN_CELLS = 1_000_000


def extract_office(prepared: VerifiedParserInput) -> ParsedResult:
    if type(prepared) is not VerifiedParserInput:
        raise ParserResultError("PARSER_INPUT_INVALID")
    plan = prepared.plan
    if (plan != choose_parser_profile(plan.source)
            or plan.parser_profile not in ("DOCX", "PPTX", "XLSX")):
        raise ParserResultError("PARSER_FORMAT_UNSUPPORTED")
    try:
        prepared.stream.seek(0)
        raw = prepared.stream.read(plan.source.size_bytes + 1)
    except (OSError, ValueError, AttributeError):
        raise ParserResultError("PARSER_INPUT_UNAVAILABLE") from None
    if (type(raw) is not bytes or len(raw) != plan.source.size_bytes
            or not hmac.compare_digest(hashlib.sha256(raw).digest(),
                                       plan.source.content_sha256)):
        raise ParserResultError("FILE_INTEGRITY_MISMATCH")
    _check_office_package(raw)
    nodes = _NodeCollector()
    try:
        if plan.parser_profile == "DOCX":
            _docx_nodes(raw, nodes)
        elif plan.parser_profile == "PPTX":
            _pptx_nodes(raw, nodes)
        else:
            _xlsx_nodes(raw, nodes)
    except ParserResultError:
        raise
    except Exception:
        raise ParserResultError("PARSER_OFFICE_INVALID") from None
    return ParsedResult(plan.source.document_version_id, plan.source.content_sha256,
                        plan.parser_profile, plan.parser_version, tuple(nodes.nodes))


def _check_office_package(raw: bytes) -> None:
    try:
        with zipfile.ZipFile(io.BytesIO(raw)) as archive:
            members = archive.infolist()
            if not members or len(members) > _MAX_ZIP_MEMBERS:
                raise ParserResultError("PARSER_OFFICE_LIMIT_EXCEEDED")
            total_size = 0
            names: set[str] = set()
            for member in members:
                parts = member.filename.replace("\\", "/").split("/")
                if (member.flag_bits & 1 or member.filename.startswith("/")
                        or any(part == ".." for part in parts)
                        or ":" in parts[0] or member.filename in names):
                    raise ParserResultError("PARSER_OFFICE_INVALID")
                names.add(member.filename)
                total_size += member.file_size
                if total_size > _MAX_UNCOMPRESSED_BYTES:
                    raise ParserResultError("PARSER_OFFICE_LIMIT_EXCEEDED")
    except (zipfile.BadZipFile, zipfile.LargeZipFile):
        raise ParserResultError("PARSER_OFFICE_INVALID") from None


class _NodeCollector:
    def __init__(self) -> None:
        self.nodes: list[ParsedNode] = []
        self.chars = 0

    def append(self, node: ParsedNode) -> None:
        self.chars += len(node.text)
        if len(self.nodes) >= _MAX_NODES or self.chars > _MAX_RESULT_CHARS:
            raise ParserResultError("PARSER_RESULT_LIMIT_EXCEEDED")
        self.nodes.append(node)


def _docx_nodes(raw: bytes, nodes: _NodeCollector) -> None:
    document = Document(io.BytesIO(raw))
    paragraph_index = table_index = 0
    for child in document.element.body.iterchildren():
        if isinstance(child, CT_P):
            paragraph_index += 1
            paragraph = Paragraph(child, document)
            value = paragraph.text
            if value.strip():
                nodes.append(ParsedNode(f"p:{paragraph_index}", "DOCX_PARAGRAPH",
                                          value, ParagraphPosition(paragraph_index)))
                heading = re.fullmatch(r"Heading([1-9])", paragraph.style.style_id)
                if heading is not None:
                    nodes.append(ParsedNode(
                        f"heading:{paragraph_index}", "DOCX_SECTION", value,
                        SectionPosition(f"word/heading/{heading.group(1)}/{paragraph_index}")))
        elif isinstance(child, CT_Tbl):
            table_index += 1
            table = Table(child, document)
            seen_cells: set[int] = set()
            for row_no, row in enumerate(table.rows, 1):
                for column_no, cell in enumerate(row.cells, 1):
                    cell_identity = id(cell._tc)
                    if cell_identity in seen_cells:
                        continue
                    seen_cells.add(cell_identity)
                    if cell.text.strip():
                        nodes.append(ParsedNode(f"t:{table_index}:r:{row_no}:c:{column_no}",
                                                  "DOCX_TABLE_CELL", cell.text,
                                                  TableCellPosition(f"word/table/{table_index}",
                                                                    row_no, column_no)))


def _pptx_nodes(raw: bytes, nodes: _NodeCollector) -> None:
    presentation = Presentation(io.BytesIO(raw))
    for slide_no, slide in enumerate(presentation.slides, 1):
        for shape in slide.shapes:
            shape_id = str(shape.shape_id)
            if shape.has_table:
                for row_no, row in enumerate(shape.table.rows, 1):
                    for column_no, cell in enumerate(row.cells, 1):
                        if cell.text.strip():
                            nodes.append(ParsedNode(
                                f"s:{slide_no}:h:{shape_id}:r:{row_no}:c:{column_no}",
                                "PPTX_TABLE_CELL", cell.text,
                                TableCellPosition(f"slide/{slide_no}/shape/{shape_id}",
                                                  row_no, column_no)))
            elif shape.has_text_frame and shape.text.strip():
                nodes.append(ParsedNode(f"s:{slide_no}:h:{shape_id}", "PPTX_SHAPE",
                                          shape.text, SlideShapePosition(slide_no, shape_id)))


def _xlsx_nodes(raw: bytes, nodes: _NodeCollector) -> None:
    workbook = load_workbook(io.BytesIO(raw), read_only=True, data_only=False,
                             keep_links=False)
    try:
        for sheet in workbook.worksheets:
            if ((sheet.max_row or 0) * (sheet.max_column or 0)
                    > _MAX_SHEET_SCAN_CELLS):
                raise ParserResultError("PARSER_RESULT_LIMIT_EXCEEDED")
            for row in sheet.iter_rows():
                for cell in row:
                    if cell.value is not None:
                        value = str(cell.value)
                        if value.strip():
                            nodes.append(ParsedNode(
                                f"sheet:{sheet.title}:cell:{cell.coordinate}", "XLSX_CELL",
                                value, SheetCellPosition(sheet.title, cell.coordinate)))
    finally:
        workbook.close()
