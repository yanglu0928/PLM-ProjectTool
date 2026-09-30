from __future__ import annotations

import hashlib
import io
import unittest
import uuid
import zipfile

from docx import Document
from openpyxl import Workbook
from pptx import Presentation
from pptx.util import Inches

from plm_assistant.modules.evidence.domain.locator import validate_evidence_locator
from plm_assistant.modules.parser.application.extract_office import extract_office
from plm_assistant.modules.parser.application.prepare_input import VerifiedParserInput
from plm_assistant.modules.parser.application.profile_selection import (
    ParserInputVersion, choose_parser_profile,
)
from plm_assistant.modules.parser.application.structured_result import ParserResultError


_MIMES = {
    "DOCX": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "PPTX": "application/vnd.openxmlformats-officedocument.presentationml.presentation",
    "XLSX": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
}


class ExtractOfficeTests(unittest.TestCase):
    def prepared(self, raw: bytes, profile: str) -> VerifiedParserInput:
        source = ParserInputVersion(uuid.uuid4(), hashlib.sha256(raw).digest(),
                                    len(raw), _MIMES[profile])
        return VerifiedParserInput(choose_parser_profile(source), uuid.uuid4(), 1, 1,
                                   io.BytesIO(raw))

    def test_docx_body_order_paragraph_and_cell(self) -> None:
        document = Document()
        document.add_paragraph("需求甲")
        table = document.add_table(rows=1, cols=2)
        table.cell(0, 0).text = "接口"
        table.cell(0, 1).text = "确认"
        document.add_paragraph("结论乙")
        stream = io.BytesIO()
        document.save(stream)
        prepared = self.prepared(stream.getvalue(), "DOCX")
        parsed = extract_office(prepared)
        self.assertEqual([node.text for node in parsed.nodes],
                         ["需求甲", "接口", "确认", "结论乙"])
        self.assertEqual([node.kind for node in parsed.nodes],
                         ["DOCX_PARAGRAPH", "DOCX_TABLE_CELL", "DOCX_TABLE_CELL",
                          "DOCX_PARAGRAPH"])
        self.assertEqual(parsed.nodes[1].position.to_locator(),
                         {"locator_type": "TABLE_CELL", "table_anchor": "word/table/1",
                          "row_no": 1, "column_no": 1})
        self._assert_replay(prepared, parsed)

    def test_pptx_slide_shape_and_table_cell(self) -> None:
        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        textbox = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(3), Inches(1))
        textbox.text = "进度甲"
        table_shape = slide.shapes.add_table(1, 1, Inches(1), Inches(2),
                                            Inches(3), Inches(1))
        table_shape.table.cell(0, 0).text = "里程碑乙"
        stream = io.BytesIO()
        presentation.save(stream)
        prepared = self.prepared(stream.getvalue(), "PPTX")
        parsed = extract_office(prepared)
        self.assertEqual([node.text for node in parsed.nodes], ["进度甲", "里程碑乙"])
        self.assertEqual(parsed.nodes[0].position.to_locator()["slide_no"], 1)
        self.assertEqual(parsed.nodes[1].position.to_locator()["table_anchor"],
                         f"slide/1/shape/{table_shape.shape_id}")
        self._assert_replay(prepared, parsed)

    def test_docx_merged_cell_emits_one_top_left_position(self) -> None:
        document = Document()
        table = document.add_table(rows=2, cols=2)
        table.cell(0, 0).merge(table.cell(1, 1)).text = "合并范围"
        stream = io.BytesIO()
        document.save(stream)
        parsed = extract_office(self.prepared(stream.getvalue(), "DOCX"))
        self.assertEqual(len(parsed.nodes), 1)
        self.assertEqual(parsed.nodes[0].position.to_locator()["row_no"], 1)
        self.assertEqual(parsed.nodes[0].position.to_locator()["column_no"], 1)

    def test_xlsx_sheet_cells_formula_is_text_not_execution(self) -> None:
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "业务调研"
        sheet["B2"] = "项目甲"
        sheet["C3"] = "=1+1"
        stream = io.BytesIO()
        workbook.save(stream)
        prepared = self.prepared(stream.getvalue(), "XLSX")
        parsed = extract_office(prepared)
        self.assertEqual([node.text for node in parsed.nodes], ["项目甲", "=1+1"])
        self.assertEqual(parsed.nodes[0].position.to_locator(),
                         {"locator_type": "SHEET_RANGE", "sheet_name": "业务调研",
                          "start_cell": "B2", "end_cell": "B2"})
        self._assert_replay(prepared, parsed)

    def test_corrupt_digest_and_unsafe_zip_fail_closed(self) -> None:
        prepared = self.prepared(b"not a zip", "DOCX")
        with self.assertRaises(ParserResultError) as error:
            extract_office(prepared)
        self.assertEqual(error.exception.code, "PARSER_OFFICE_INVALID")
        stream = io.BytesIO()
        with zipfile.ZipFile(stream, "w") as archive:
            archive.writestr("../evil.xml", "x")
        prepared = self.prepared(stream.getvalue(), "DOCX")
        with self.assertRaises(ParserResultError) as error:
            extract_office(prepared)
        self.assertEqual(error.exception.code, "PARSER_OFFICE_INVALID")
        prepared = self.prepared(stream.getvalue(), "DOCX")
        prepared.stream = io.BytesIO(stream.getvalue() + b"changed")
        with self.assertRaises(ParserResultError) as error:
            extract_office(prepared)
        self.assertEqual(error.exception.code, "FILE_INTEGRITY_MISMATCH")

    def test_sparse_xlsx_huge_declared_range_fails_before_scan(self) -> None:
        workbook = Workbook()
        workbook.active["XFD1048576"] = "unsafe scan"
        stream = io.BytesIO()
        workbook.save(stream)
        with self.assertRaises(ParserResultError) as error:
            extract_office(self.prepared(stream.getvalue(), "XLSX"))
        self.assertEqual(error.exception.code, "PARSER_RESULT_LIMIT_EXCEEDED")

    def _assert_replay(self, prepared, parsed) -> None:
        self.assertEqual(parsed.canonical_bytes(), extract_office(prepared).canonical_bytes())
        for node in parsed.nodes:
            locator = node.position.to_locator()
            self.assertEqual(validate_evidence_locator(locator), locator)


if __name__ == "__main__":
    unittest.main()
